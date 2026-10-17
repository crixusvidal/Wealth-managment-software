"""Figura E6 de ESG: cascada de ahorros ESG (y de menor costo de capital) hasta el precio objetivo, frente al precio de la acción.

Usa la hoja ESG_Ahorros de la variante del modelo (valuacion/modelos/Mallplaza_Valoracion_Consolidada_v3_ESG.xlsx, que es el
modelo v3 más esa hoja). Los volúmenes físicos (MWh, m³, toneladas) son datos de la Memoria Mallplaza 2025; los precios
unitarios y el % del ahorro que queda en Mallplaza NO están en las fuentes: se cargan aquí como supuestos ilustrativos
(celdas amarillas de la hoja, marcadas como supuesto) y el resultado es un escenario, no un dato. Los escenarios de menor
costo de capital (ΔKd y ΔKe) vienen de la Sensibilidad 2 de la misma hoja.

Pasos: carga los supuestos en las celdas B26:B29 y B31, recalcula con LibreOffice, lee los resultados y dibuja la figura
(cascada arriba, tabla abajo; sin título ni fuente en la imagen).

Uso (desde secciones/): python esg/graficos/esg_ahorro.py <modelo_v3_ESG.xlsx> [carpeta_salida]
Salida: <carpeta>/{es,en}_E6_ahorro_esg.png y figuras/datos_esg_ahorro.json
"""
import json
import subprocess
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl
from matplotlib.patches import Rectangle
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))  # paleta institucional
import estilo  # noqa: E402,F401
from estilo import AMARILLO, GRILLA, NEGRO, ROJO, CELESTE, TINTA_SECUNDARIA, num

AZUL = "#14284B"
FIGS = Path(__file__).parent / "figuras"
RECALC = "/root/.claude/skills/synced/508cc6bc-ad76-422f-ada9-ad0fc4279625_0b2a2ef9-bf53-46a5-b8ed-72255cce463a/xlsx/scripts/recalc.py"
# Supuestos ilustrativos (NO son datos de las fuentes): precio unitario de cada ahorro y % que retiene Mallplaza
SUPUESTOS = {"B26": 110, "B27": 110, "B28": 1500, "B29": 10000, "B31": 0.25}     # CLP/kWh, CLP/kWh, CLP/m³, CLP/t, %
ESC_KC = [(-25, "D81"), (-50, "E82"), (-100, "F83")]                               # ΔKd = ΔKe (pb) → celda de Δ precio objetivo
T = {
    "es": {"pasos": ["Objetivo\nbase", "Energía", "Solar", "Agua", "Residuos", "Costo cap.\n−25 pb", "Con ESG"],
           "cols": ["Ahorro ESG (escenario)", "MM CLP/año", "Δ objetivo"], "mm": "CLP por acción",
           "filas": ["Energía: eficiencia (13.785 MWh)", "Energía: autogeneración solar (2.100 MWh)", "Agua: eficiencia (31.276 m³)",
                     "Residuos valorizados (14.988 t)", "Subtotal operativo (25% queda en Mallplaza)",
                     "Costo de capital −25 pb (Kd y Ke)", "Costo de capital −50 pb (Kd y Ke)", "Costo de capital −100 pb (Kd y Ke)"],
           "precio": "Precio {p}", "brecha": "Brecha {b}"},
    "en": {"pasos": ["Base\ntarget", "Energy", "Solar", "Water", "Waste", "Cost of cap.\n−25 bp", "With ESG"],
           "cols": ["ESG savings (scenario)", "CLP MM/yr", "Δ target"], "mm": "CLP per share",
           "filas": ["Energy: efficiency (13,785 MWh)", "Energy: solar self-generation (2,100 MWh)", "Water: efficiency (31,276 m³)",
                     "Recovered waste (14,988 t)", "Operating subtotal (25% stays with Mallplaza)",
                     "Cost of capital −25 bp (Kd and Ke)", "Cost of capital −50 bp (Kd and Ke)", "Cost of capital −100 bp (Kd and Ke)"],
           "precio": "Price {p}", "brecha": "Gap {b}"},
}


def calcular(ruta):
    wb = openpyxl.load_workbook(ruta)
    ws = wb["ESG_Ahorros"]
    for c, v in SUPUESTOS.items():
        ws[c] = v
        ws[c].font = Font(color="0000FF")
        ws["H" + c[1:]] = "SUPUESTO ILUSTRATIVO del equipo (no es dato de las fuentes); reemplazar con el dato real"
    wb.save(ruta)
    subprocess.run([sys.executable, RECALC, str(ruta)], check=True, stdout=subprocess.DEVNULL)
    v = openpyxl.load_workbook(ruta, data_only=True)
    e, po = v["ESG_Ahorros"], v["Precio_Objetivo"]
    factor = e["B49"].value                                   # Δ precio objetivo por MM CLP/año de ahorro
    captura = e["B31"].value
    brutos = [e[f"D{r}"].value for r in (26, 27, 28, 29)]
    delta_op = [b * captura * factor for b in brutos]
    kc = [e[c].value for _, c in ESC_KC]
    assert abs(sum(delta_op) - e["D33"].value) < 1e-9, "las partidas deben sumar el Δ precio objetivo de la hoja"
    return {"base": po["B15"].value, "precio": po["B16"].value, "factor": factor, "captura": captura, "supuestos": SUPUESTOS,
            "brutos": brutos, "capturado": [b * captura for b in brutos], "delta_op": delta_op, "delta_op_total": sum(delta_op),
            "kc": dict(zip([f"{p}" for p, _ in ESC_KC], kc)), "ebitda_ltm_pct": e["D30"].value / v["Supuestos"]["B55"].value if False else None}


def figura(lang, r, salida):
    t = T[lang]
    fig = plt.figure(figsize=(3550 / 1440, 2.5), dpi=400)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0.13, 0.5, 0.85, 0.47])
    pasos = r["delta_op"] + [r["kc"]["-25"]]
    acum = [r["base"]]
    for d in pasos:
        acum.append(acum[-1] + d)
    n = len(acum)
    ax.bar(0, acum[0], width=0.62, color=AZUL, linewidth=0)
    ax.text(0, acum[0] + 40, num(acum[0], lang, 0), ha="center", va="bottom", fontsize=4.6, fontweight="bold", color=AZUL)
    for k in range(1, n):
        d = acum[k] - acum[k - 1]
        ax.bar(k, max(d, 6), bottom=acum[k - 1], width=0.62, color=CELESTE, linewidth=0)
        ax.text(k, acum[k] + 40, "+" + num(d, lang, 1 if d < 20 else 0), ha="center", va="bottom", fontsize=4.6, color=NEGRO)
        ax.plot([k - 0.69, k - 0.31], [acum[k - 1]] * 2, color=NEGRO, linewidth=0.3)
    ax.bar(n, acum[-1], width=0.62, color=NEGRO, linewidth=0)
    ax.text(n, acum[-1] + 40, num(acum[-1], lang, 0), ha="center", va="bottom", fontsize=4.8, fontweight="bold", color=NEGRO)
    ax.axhline(r["precio"], color=ROJO, linewidth=0.8, linestyle=(0, (4, 2)))
    ax.text(-0.45, r["precio"] + 40, t["precio"].format(p=num(r["precio"], lang, 0)), ha="left", va="bottom", fontsize=4.6, fontweight="bold", color=ROJO)
    ax.text(n + 0.4, r["precio"] + 40, t["brecha"].format(b=num(r["precio"] - acum[-1], lang, 0)), ha="right", va="bottom", fontsize=4.4, color=ROJO)
    ax.set_xticks(range(n + 1), t["pasos"], fontsize=4.2, linespacing=1.05)
    ax.set_ylim(1800, 3900)
    ax.set_xlim(-0.55, n + 0.55)
    ax.yaxis.set_major_formatter(lambda y, _: num(y, lang, 0))
    ax.tick_params(labelsize=4.2, length=1.5, colors=TINTA_SECUNDARIA)
    ax.grid(axis="y", color=GRILLA, linewidth=0.4)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_ylabel(t["mm"], fontsize=4.2, color=TINTA_SECUNDARIA, labelpad=2)
    # Tabla
    axt = fig.add_axes([0.03, 0.0, 0.94, 0.36])
    axt.axis("off")
    xs, alin = [0.0, 0.8, 0.99], ["left", "right", "right"]
    axt.add_patch(Rectangle((0, 0.88), 1, 0.12, color=AZUL, transform=axt.transAxes))
    for x, c, a in zip(xs, t["cols"], alin):
        axt.text(x + (0.01 if a == "left" else 0), 0.94, c, ha=a, va="center", fontsize=4.5, fontweight="bold", color="white", transform=axt.transAxes)
    filas = [(num(b, lang, 0), "+" + num(d, lang, 1)) for b, d in zip(r["brutos"], r["delta_op"])]
    filas.append((num(sum(r["capturado"]), lang, 0), "+" + num(r["delta_op_total"], lang, 1)))
    filas += [("—", "+" + num(r["kc"][k], lang, 0)) for k in ("-25", "-50", "-100")]
    for k, (a, b) in enumerate(filas):
        y = 0.8 - k * 0.1075
        if k == 4:
            axt.add_patch(Rectangle((0, y - 0.052), 1, 0.105, color=AMARILLO, linewidth=0, transform=axt.transAxes))
        for x, v, al in zip(xs, (t["filas"][k], a, b), alin):
            axt.text(x + (0.01 if al == "left" else 0), y, v, ha=al, va="center", fontsize=4.4, color=NEGRO,
                     fontweight="bold" if k == 4 else "normal", transform=axt.transAxes)
    axt.set_xlim(0, 1)
    fig.savefig(salida, facecolor="white")
    plt.close(fig)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ruta = Path(sys.argv[1])
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else FIGS
    r = calcular(ruta)
    r.pop("ebitda_ltm_pct")
    (FIGS / "datos_esg_ahorro.json").write_text(json.dumps(r, indent=1, ensure_ascii=False))
    for lang in ("es", "en"):
        figura(lang, r, salida / f"{lang}_E6_ahorro_esg.png")
    print({k: (round(v, 2) if isinstance(v, float) else v) for k, v in r.items() if k in ("base", "delta_op_total", "kc", "factor")}, [round(b) for b in r["brutos"]])
