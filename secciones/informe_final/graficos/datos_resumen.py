"""Datos clave del Resumen de inversión, leídos de los modelos (sin cifras escritas a mano).

Uso (desde secciones/):
    python informe_final/graficos/datos_resumen.py <Valoracion_Consolidada_v4.xlsx> <Figuras_Valuacion_recalc.xlsx>
Salida: informe_final/graficos/datos_resumen.json
"""
import json
import sys
from pathlib import Path

import openpyxl

v2 = openpyxl.load_workbook(sys.argv[1], data_only=True)
fig = openpyxl.load_workbook(sys.argv[2], data_only=True)
s, p, d = v2["Supuestos"], v2["Precio_Objetivo"], v2["Deuda"]
precio, acciones = s["B8"].value, s["B9"].value             # CLP; millones de acciones
semanas = next(r for r in fig["V1_Football"].iter_rows(min_row=5, max_row=10, values_only=True)
               if r[0] and str(r[0]).startswith("Rango 52"))
dividendo = d["D16"].value / acciones                       # dividendo ordinario 2027E por acción (CLP)
datos = {
    "precio_objetivo": p["B15"].value, "precio": precio, "potencial": p["B17"].value,
    "dcf": p["B12"].value, "multiplos": p["B13"].value, "peso_dcf": p["B14"].value,
    "acciones_mm": acciones, "capitalizacion_mm": precio * acciones,
    "rango_52s": [semanas[1], semanas[2]], "dividendo": dividendo, "dividend_yield": dividendo / precio,
    "dfn_mm": s["B10"].value, "ebitda_udm_mm": s["B13"].value,
    "fuente": "Mallplaza_Valoracion_Consolidada_v4.xlsx (Supuestos B8–B13, Precio_Objetivo B12–B17, Deuda D16); "
              "Mallplaza_Figuras_Valuacion_recalc.xlsx (V1_Football, rango 52 semanas)",
}
salida = Path(__file__).parent / "datos_resumen.json"
salida.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
print(salida, datos)
