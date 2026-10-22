import copy,sys
from docx import Document
from docx.shared import Pt,Inches,RGBColor,Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image,ImageFilter
SRC,DST,EQUIPO=sys.argv[1],sys.argv[2],sys.argv[3]
# logo: recorte ajustado de la captura, fondo blanco, ampliado y afinado
import numpy as np
im=Image.open("/tmp/claude-0/-home-user-Wealth-managment-software/4853c681-1399-5e51-b2e7-c4edae34dd33/images/27.png").convert("RGB").crop((104,112,386,178))
a=np.asarray(im).astype(float); claro=(a.min(axis=2)>205); a[claro]=255; a[a.min(axis=2)>190]=255; im=Image.fromarray(a.astype("uint8"))
im=im.resize((im.width*6,im.height*6),Image.LANCZOS).filter(ImageFilter.UnsharpMask(radius=5,percent=120,threshold=3))
a=np.asarray(im).astype(float); a[a.min(axis=2)>200]=255; _im=Image.fromarray(a.astype("uint8")).convert("RGBA"); _d=np.asarray(_im).copy()
_d[...,3]=np.clip((255-_d[...,:3].min(axis=2))*2.2,0,255).astype("uint8"); _d[...,:3]=np.where(_d[...,3:4]<255,_d[...,:3],_d[...,:3]); Image.fromarray(_d).save("logo.png")
doc=Document(SRC); body=doc.element.body
GRIS=RGBColor(0x59,0x59,0x59); AZ="5B7BD5"
nuevos=[]
def para(texto="",tam=10,neg=False,color=None,al=None,antes=0,desp=0,jus=False,interlineado=None):
    p=doc.add_paragraph(); nuevos.append(p._p)
    pf=p.paragraph_format; pf.space_before=Pt(antes); pf.space_after=Pt(desp)
    if interlineado: pf.line_spacing=interlineado
    if al: p.alignment=al
    if jus: p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if texto: run(p,texto,tam,neg,color)
    return p
def run(p,texto,tam,neg=False,color=None):
    r=p.add_run(texto); r.font.size=Pt(tam); r.font.bold=neg; r.font.name="Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"),"Calibri")
    if color: r.font.color.rgb=color
    return r
def regla(p,pos="bottom"):
    ppr=p._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); e=OxmlElement(f"w:{pos}")
    for k,v in (("val","single"),("sz","18"),("space","1"),("color",AZ)): e.set(qn(f"w:{k}"),v)
    b.append(e); ppr.append(b)
C=WD_ALIGN_PARAGRAPH.CENTER
p=para(antes=105,desp=40,al=C); p.add_run().add_picture("logo.png",width=Pt(332))
r=para(desp=20); regla(r,"top"); r.paragraph_format.space_after=Pt(18)
for rr in r.runs: rr.font.size=Pt(2)
para("CFA Institute Research Challenge",17.3,True,GRIS,C,desp=17)
para("Hosted by",13.2,True,GRIS,C,desp=20)
para("CFA Society Chile",17,True,GRIS,C,desp=24)
para(EQUIPO,17,True,GRIS,C,desp=22)
r=para(desp=14); regla(r,"top")
para("The CFA Institute Research Challenge is a global competition that tests the equity research and valuation, investment report writing, and presentation skills of university students. The following report was prepared in compliance with the Official Rules of the CFA Institute Research Challenge, is submitted by a team of university students as part of this annual educational initiative and should not be considered a professional report.",11.4,desp=30,interlineado=1.1)
para("Disclosures:",15.6,True,None,antes=0,desp=6)
def sub(t,txt=None):
    p=para(desp=0,interlineado=1.12); run(p,t,10.2,True)
    if txt:
        q=para(desp=0,interlineado=1.12); run(q,txt,10.2,False)
sub("Ownership and material conflicts of interest:","The author(s), or a member of their household, of this report does not hold a financial interest in the securities of this Company. The author(s), or a member of their household, of this report does not know of the existence of any conflicts of interest that might bias the content or publication of this report.")
sub("Receipt of compensation","Compensation of the author(s) of this report is not based on investment banking revenue.")
p=para(desp=0,interlineado=1.12); run(p,"Position as an officer or a director",10.2,True); run(p," The author(s), or a member of their household, does not serve as an officer, director, or advisory board member of the subject Company.",10.2)
sub("Market making","The author(s) does not act as a market maker in the subject Company's securities.")
sub("Disclaimer","The information set forth herein has been obtained or derived from sources generally available to the public and believed by the author(s) to be reliable, but the author(s) does not make any representation or warranty, express or implied, as to its accuracy or completeness. The information is not intended to be used as the basis of any investment decisions by any person or entity. This information does not constitute investment advice, nor is it an offer or a solicitation of an offer to buy or sell any security.")
# fin de sección: A4 con márgenes normales, sin pie de página
ultimo=nuevos[-1]; ppr=ultimo.find(qn("w:pPr"))
if ppr is None: ppr=OxmlElement("w:pPr"); ultimo.insert(0,ppr)
sp=OxmlElement("w:sectPr")
t=OxmlElement("w:type"); t.set(qn("w:val"),"nextPage"); sp.append(t)
pg=OxmlElement("w:pgSz"); pg.set(qn("w:w"),"11906"); pg.set(qn("w:h"),"16838"); sp.append(pg)
m=OxmlElement("w:pgMar")
for k,v in (("top","720"),("right","580"),("bottom","600"),("left","580"),("header","567"),("footer","567"),("gutter","0")): m.set(qn(f"w:{k}"),v)
sp.append(m)
ppr.append(sp)
# insertar al inicio del cuerpo
primero=body[0]
for el in nuevos: primero.addprevious(el)
# la primera sección original arranca en la página 1 del informe (la portada no cuenta)
for s in body.iter(qn("w:sectPr")):
    if s is sp: continue
    pn=s.find(qn("w:pgNumType"))
    if pn is None: pn=OxmlElement("w:pgNumType"); s.append(pn)
    pn.set(qn("w:start"),"1"); break
doc.save(DST); print("ok")
