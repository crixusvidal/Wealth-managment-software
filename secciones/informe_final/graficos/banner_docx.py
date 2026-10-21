"""Pone al inicio de un Word del informe el banner de portada (logo, banda roja con la recomendación).

Uso: python3 banner_docx.py entrada.docx salida.docx es|en
Las cifras salen de cifras.json (mismas que el banner de la versión Typst).
"""
import copy, json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

RAIZ = Path(__file__).resolve().parents[2]
LOGO = RAIZ / "comun" / "logos" / "mallplaza.png"
ROJO = "E9004B"


def num(v, lang, d=0):
    s = f"{abs(v):,.{d}f}"
    s = s.replace(",", "X").replace(".", "," if lang == "es" else ".").replace("X", "." if lang == "es" else ",")
    return ("−" if v < 0 else "") + s


def celda_fondo(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def run(p, texto, pt, negrita=False, cursiva=False):
    r = p.add_run(texto)
    r.font.size, r.font.bold, r.font.italic = Pt(pt), negrita, cursiva
    r.font.color.rgb = RGBColor(255, 255, 255)
    return r


def agregar(entrada, salida, lang):
    es = lang == "es"
    c = json.loads((RAIZ / "informe_final" / "graficos" / "cifras.json").read_text())
    precio_obj, precio = c["po"], c["precio"]
    pot = precio_obj / precio - 1
    doc = Document(entrada)
    primero = next(p for p in doc.paragraphs if p.text.strip().upper() == "ABSTRACT")
    ancho = 11506
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    logo, banda = t.rows[0].cells
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    w_ = OxmlElement("w:tblW"); w_.set(qn("w:w"), str(ancho)); w_.set(qn("w:type"), "dxa"); tblPr.append(w_)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for g, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), (1300, ancho - 1300)):
        g.set(qn("w:w"), str(w))
    for celda, w in ((logo, 1300), (banda, ancho - 1300)):
        celda.width = Twips(w)
        celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    pl = logo.paragraphs[0]
    pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pl.add_run().add_picture(str(LOGO), height=Inches(0.5))
    celda_fondo(banda, ROJO)
    p = banda.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    run(p, "Mallplaza (Plaza S.A.)  ·  MALLPLAZ CI", 15, negrita=True)
    q = banda.add_paragraph()
    run(q, ("Recomendación: VENDER" if es else "Recommendation: SELL") + "        ", 10.5, negrita=True)
    run(q, f"{'Precio objetivo' if es else 'Target price'}: CLP {num(precio_obj, lang)} ({num(pot * 100, lang, 1)}%)", 10.5, negrita=True)
    r = banda.add_paragraph()
    r.paragraph_format.space_after = Pt(0)
    run(r, ("Sector: Centros comerciales (Chile, Perú, Colombia)" if es else "Sector: Shopping centers (Chile, Peru, Colombia)")
        + f"        {'Precio' if es else 'Price'} (30-{'sep' if es else 'Sep'}-26): CLP {num(precio, lang)}", 8.5, cursiva=True)
    for par in (pl, q):
        par.paragraph_format.space_before = Pt(0)
    t._tbl.getparent().remove(t._tbl)
    primero._p.addprevious(t._tbl)
    doc.save(salida)


if __name__ == "__main__":
    agregar(sys.argv[1], sys.argv[2], sys.argv[3])
