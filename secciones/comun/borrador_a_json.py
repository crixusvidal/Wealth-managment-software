"""Convierte un borrador Word del equipo (formato de dos columnas) en un JSON con el texto tal cual.

El texto no se modifica: solo se quitan espacios de ancho cero y se unen las corridas con el mismo
formato. El JSON lo leen los constructores Word de cada sección (documento/*.js), que aplican el
formato principal del informe y reemplazan las imágenes del borrador por los gráficos del repo.

Estructura de salida:
    {"bloques": [
        {"tipo": "titulo", "texto": "..."},
        {"tipo": "pagina", "filas": [{"izq": [bloque...], "der": [bloque...]}]},   # tabla texto | figuras
        {"tipo": "figuras", "filas": [[celda, celda], ...]},                        # anexos lado a lado
        {"tipo": "tabla", "filas": [[texto, ...], ...], "negrita": [[bool, ...], ...]},
        {"tipo": "caption" | "fuente" | "parrafo", ...}]}
    bloque de celda = {"tipo": "h1", "texto"} | {"tipo": "p", "runs": [[texto, negrita, cursiva]]}
                    | {"tipo": "caption", "n": "Figura 1:", "texto"} | {"tipo": "imagen", "archivo"}
                    | {"tipo": "fuente", "runs": [...]}

Uso:
    python comun/borrador_a_json.py <borrador.docx> <salida.json>
"""
import json
import re
import sys

import docx
from docx.oxml.ns import qn

ZW = re.compile("[​‌‍﻿]")


def corridas(p):
    out = []
    for r in p.runs:
        t = ZW.sub("", r.text)
        if not t:
            continue
        b, i = bool(r.bold), bool(r.italic)
        if out and out[-1][1] == b and out[-1][2] == i:
            out[-1][0] += t
        else:
            out.append([t, b, i])
    return out


def imagen(p, rels):
    blip = p._p.find(".//" + qn("a:blip"))
    if blip is None:
        return None
    return rels[blip.get(qn("r:embed"))].target_ref.split("/")[-1]


def bloque(p, rels):
    img = imagen(p, rels)
    if img:
        return {"tipo": "imagen", "archivo": img}
    rs = corridas(p)
    if not rs:
        return None
    texto = "".join(t for t, _, _ in rs)
    if texto.startswith("Fuente") or texto.startswith("Source"):
        return {"tipo": "fuente", "runs": rs}
    m = re.match(r"^((?:Figura|Figure|Anexo|Appendix)\s+\S+?:)\s*", texto)
    if m and rs[0][1]:
        return {"tipo": "caption", "n": m.group(1), "texto": texto[m.end():]}
    if all(b for _, b, _ in rs) and texto.upper() == texto and len(texto) > 3:
        return {"tipo": "h1", "texto": texto}
    return {"tipo": "p", "runs": rs}


def celda(c, rels):
    return [b for b in (bloque(p, rels) for p in c.paragraphs) if b]


def convertir(ruta):
    d = docx.Document(ruta)
    rels = d.part.rels
    salida = []
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = docx.text.paragraph.Paragraph(el, d)
            rs = corridas(p)
            if not rs:
                continue
            texto = "".join(t for t, _, _ in rs)
            b = bloque(p, rels)
            if b["tipo"] == "h1" or (b["tipo"] == "p" and all(x[1] for x in rs) and len(texto) > 3
                                      and texto.upper() == texto):
                salida.append({"tipo": "titulo", "texto": texto})
            elif b["tipo"] == "p" and rs[0][1] and re.match(r"^(Anexo|Appendix)\s+\S+:", texto):
                m = re.match(r"^((?:Anexo|Appendix)\s+\S+:)\s*", texto)
                salida.append({"tipo": "caption", "n": m.group(1), "texto": texto[m.end():]})
            else:
                salida.append(b if b["tipo"] != "p" else {"tipo": "parrafo", "runs": rs})
        elif el.tag == qn("w:tbl"):
            t = docx.table.Table(el, d)
            anchos = [int(g.get(qn("w:w"))) for g in el.iter(qn("w:gridCol"))]
            maqueta = len(anchos) == 3 and anchos[1] < 400
            filas = [[celda(c, rels) for c in row.cells] for row in t.rows]
            if maqueta and any(any(b["tipo"] == "p" for b in f[0]) for f in filas):
                salida.append({"tipo": "pagina", "filas": [{"izq": f[0], "der": f[2]} for f in filas]})
            elif maqueta:
                salida.append({"tipo": "figuras", "filas": [[f[0], f[2]] for f in filas]})
            else:
                salida.append({"tipo": "tabla",
                               "filas": [[ZW.sub("", c.text).strip() for c in row.cells] for row in t.rows],
                               "negrita": [[bool(c.paragraphs and c.paragraphs[0].runs and c.paragraphs[0].runs[0].bold)
                                            for c in row.cells] for row in t.rows],
                               "anchos": anchos})
    return {"origen": ruta.split("/")[-1], "bloques": salida}


if __name__ == "__main__":
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(convertir(sys.argv[1]), f, ensure_ascii=False, indent=1)
    print(sys.argv[2])
