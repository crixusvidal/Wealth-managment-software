"""Figuras con los datos de mercado extraídos de Bloomberg y el Banco Central (Mallplaza_Datos_Graficos_Informe.xlsx):

  V18  Spread del cap rate implícito vs el bono UF a 10 años (BCU), 4T22–30-sep-26: dos líneas (% real) y barras del spread (pb, eje derecho).
  V19  Precio semanal rebasado a 100 (1-oct-21): Mallplaza, IPSA, Parque Arauco y Cenco Malls, con cuatro hitos anotados en Mallplaza.
  V20  P/NAV ajustado (EPRA NTA) vs pares, con la línea de 1,0x.

Mismo motor y estilo que las demás figuras (paleta institucional, ancho de columna 3550 twips, 400 dpi); sin título ni fuente dentro de la imagen.
Uso (desde secciones/): python valuacion/graficos/datos_mercado.py
Salida: valuacion/graficos/figuras/{es,en}_V18_caprate_spread.png, _V19_precio_base100.png, _V20_pnav.png
"""
import datetime as dt
import sys
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))
import estilo  # noqa: E402,F401
from estilo import AMARILLO, AZUL, CELESTE, GRILLA, MORADO_CENCOSUD, NEGRO, ROJO, TINTA_SECUNDARIA, VERDE_PARQUE_ARAUCO, num  # noqa: E402

LIBRO = Path(__file__).resolve().parents[2] / "valuacion/modelos/Mallplaza_Datos_Graficos_Informe.xlsx"
FIGS = Path(__file__).parent / "figuras"
ANCHO, DPI = 3550 / 1440, 400


def leer():
    wb = openpyxl.load_workbook(LIBRO, data_only=True)
    cr = [r for r in wb["Fig_CapRate"].iter_rows(min_row=3, values_only=True) if isinstance(r[1], dt.datetime) and r[8] is not None]
    pr = [r for r in wb["Fig_Precio"].iter_rows(min_row=3, values_only=True) if isinstance(r[0], dt.datetime) and isinstance(r[6], (int, float))]
    pr = [list(r) for r in pr]
    for i in range(1, len(pr)):                      # huecos de una serie: se arrastra el último valor
        for c in range(6, 10):
            if pr[i][c] is None:
                pr[i][c] = pr[i - 1][c]
    nav = {r[0]: r[1:4] for r in wb["NAV"].iter_rows(values_only=True) if r[0] in ("P/NAV ajustado", "P/BV (= P/NAV contable)")}
    return cr, pr, nav


def ejes(ax, lang, xlabel=None):
    ax.tick_params(labelsize=4.0, length=1.5, colors=TINTA_SECUNDARIA)
    ax.grid(axis="y", color=GRILLA, linewidth=0.4, zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)


def caprate(cr, lang, salida):
    es = lang == "es"
    etq = [r[0] if r[0] != "30-sep-26" or es else "30-Sep-26" for r in cr]
    x = np.arange(len(cr))
    cap = np.array([r[8] * 100 for r in cr]); bcu = np.array([r[9] * 100 for r in cr]); spr = np.array([r[10] for r in cr])
    prom = cr[0][11]
    fig = plt.figure(figsize=(ANCHO, 2.05))
    ax = fig.add_axes([0.09, 0.2, 0.78, 0.64])
    ax2 = ax.twinx()
    ax2.bar(x, spr, color=AMARILLO, width=0.7, zorder=1)
    ax2.axhline(prom, color=ROJO, linewidth=0.7, linestyle=(0, (3, 2)), zorder=2)
    ax2.set_ylim(0, 1100)
    ax2.text(0, prom + 25, (f"Prom. 4T22–3T25: {num(prom, lang)} pb" if es else f"Avg. 4Q22–3Q25: {num(prom, lang)} bp"), fontsize=3.6, color=ROJO, va="bottom")
    ax.set_zorder(ax2.get_zorder() + 1); ax.patch.set_visible(False)
    ax.plot(x, cap, color=AZUL, linewidth=1.0, marker="o", markersize=1.8, zorder=4)
    ax.plot(x, bcu, color=CELESTE, linewidth=1.0, marker="o", markersize=1.8, zorder=4)
    ax.set_ylim(0, 11)
    ax.set_xlim(-0.7, len(x) - 0.1)
    ax.set_xticks(x); ax.set_xticklabels(etq, rotation=45, ha="right")
    ax.yaxis.set_major_formatter(lambda v, _: f"{num(v, lang)}%"); ax.tick_params(labelsize=3.8)
    ax2.tick_params(labelsize=3.8, length=1.5, colors=TINTA_SECUNDARIA)
    ax2.set_ylabel("Spread (pb)" if es else "Spread (bp)", fontsize=4.0, color=TINTA_SECUNDARIA, labelpad=3)
    ejes(ax, lang)
    for s in ("top",): ax2.spines[s].set_visible(False)
    ax.text(x[-1] - 0.15, cap[-1] + 0.55, f"{num(cap[-1], lang, 2)}%", fontsize=4.2, color=AZUL, fontweight="bold", va="bottom", ha="center")
    ax.text(x[-1] - 0.15, bcu[-1] + 0.5, f"{num(bcu[-1], lang, 2)}%", fontsize=4.2, color=AZUL, fontweight="bold", va="bottom", ha="center")
    ax2.text(x[-1] - 0.15, 105, f"{num(spr[-1], lang)} pb" if es else f"{num(spr[-1], lang)} bp", fontsize=4.0, color=NEGRO, fontweight="bold", ha="center", rotation=90, va="center")
    # leyenda
    for i, (txt, col, tipo) in enumerate(((("Cap rate pro forma" if es else "Pro forma cap rate"), AZUL, "l"),
                                          (("BCU 10 años" if es else "10-year BCU"), CELESTE, "l"),
                                          (("Spread (pb, eje derecho)" if es else "Spread (bp, right axis)"), AMARILLO, "b"))):
        xx = 0.10 + i * (0.36 if i == 0 else 0.22) + (0.0 if i == 0 else 0.0)
        xx = (0.10, 0.36, 0.56)[i]
        fig.add_artist(plt.Line2D([xx, xx + 0.03], [0.95, 0.95], color=col, linewidth=2.2 if tipo == "b" else 1.0, transform=fig.transFigure))
        fig.text(xx + 0.04, 0.95, txt, fontsize=3.9, color=NEGRO, va="center")
    fig.savefig(salida, dpi=DPI, facecolor="white")
    plt.close(fig)


def precio(pr, lang, salida):
    es = lang == "es"
    f = [r[0] for r in pr]
    series = [("Mallplaza", 6, ROJO, 1.3, "-"), ("IPSA", 7, NEGRO, 0.7, "-"), ("Parque Arauco", 8, VERDE_PARQUE_ARAUCO, 0.9, "-"), ("Cenco Malls", 9, MORADO_CENCOSUD, 0.9, (0, (2, 1.5)))]
    fig = plt.figure(figsize=(ANCHO, 2.05))
    ax = fig.add_axes([0.10, 0.12, 0.87, 0.73])
    fin = sorted([(pr[-1][c], nombre, col) for nombre, c, col, lw, ls in series])
    ys, prev = [], -1e9
    for v, nombre, col in fin:                       # etiquetas finales sin pisarse
        y = max(v, prev + 24); ys.append((y, v, nombre, col)); prev = y
    for nombre, c, col, lw, ls in series:
        ax.plot(f, [r[c] for r in pr], color=col, linewidth=lw, linestyle=ls, zorder=4 if nombre == "Mallplaza" else 3)
    for y, v, nombre, col in ys:
        ax.text(f[-1] + dt.timedelta(days=22), y, f"{num(v, lang)}", fontsize=3.8, color=col, va="center", fontweight="bold")
    hitos = [(dt.datetime(2025, 10, 10), 2470, "Mín. 10-oct-25\n2.470" if es else "Low 10-Oct-25\n2,470", (dt.datetime(2024, 1, 15), 335)),
             (dt.datetime(2025, 11, 25), None, "MSCI 25-nov-25" if es else "MSCI 25-Nov-25", (dt.datetime(2024, 1, 15), 425)),
             (dt.datetime(2026, 4, 27), 4700, "Máx. 27-abr-26\n4.700" if es else "High 27-Apr-26\n4,700", (dt.datetime(2025, 1, 10), 520)),
             (dt.datetime(2026, 9, 30), 3555, ("30-sep-26\n3.555" if es else "30-Sep-26\n3,555"), (dt.datetime(2026, 4, 20), 340))]
    base0 = pr[0][1]
    for fecha, p, txt, (dx, dy) in hitos:
        i = min(range(len(f)), key=lambda k: abs(f[k] - fecha))
        y = pr[i][6]
        ax.scatter([f[i]], [y], s=9, facecolor="white", edgecolor=ROJO, linewidth=0.8, zorder=6)
        t = txt if es or "Mín" not in txt else txt
        ax.annotate(t, (f[i], y), xytext=(dx, dy), fontsize=3.8, color=NEGRO, ha="left", va="center", linespacing=1.05, arrowprops=dict(arrowstyle="-", color=TINTA_SECUNDARIA, linewidth=0.4))
    ax.axhline(100, color=GRILLA, linewidth=0.6, zorder=1)
    ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=(1, 7))); ax.xaxis.set_major_formatter(lambda v, _: (("ene" if es else "Jan") if mdates.num2date(v).month == 1 else ("jul" if es else "Jul")) + "-" + str(mdates.num2date(v).year)[2:])
    ax.set_xlim(f[0], f[-1] + dt.timedelta(days=110))
    ejes(ax, lang)
    ax.set_ylabel("Base 100 = 1-oct-21" if es else "Base 100 = 1-Oct-21", fontsize=3.8, color=TINTA_SECUNDARIA, labelpad=1)
    for i, (nombre, c, col, lw, ls) in enumerate(series):
        xx = (0.09, 0.27, 0.43, 0.68)[i]
        fig.add_artist(plt.Line2D([xx, xx + 0.035], [0.95, 0.95], color=col, linewidth=lw + 0.4, linestyle=ls, transform=fig.transFigure))
        fig.text(xx + 0.045, 0.95, nombre, fontsize=4.0, color=NEGRO, va="center")
    fig.savefig(salida, dpi=DPI, facecolor="white")
    plt.close(fig)


def pnav(nav, lang, salida):
    es = lang == "es"
    v = nav["P/NAV ajustado"]
    nombres = ["Mallplaza", "Parque Arauco", "Cenco Malls"]
    cols = [AZUL, VERDE_PARQUE_ARAUCO, MORADO_CENCOSUD]
    fig = plt.figure(figsize=(ANCHO, 1.5))
    ax = fig.add_axes([0.26, 0.18, 0.68, 0.74])
    y = np.arange(3)[::-1]
    ax.barh(y, v, color=cols, height=0.55, zorder=3)
    ax.axvline(1.0, color=ROJO, linewidth=0.8, linestyle=(0, (4, 2)), zorder=4)
    for yy, val in zip(y, v):
        ax.text(val + 0.03, yy, f"{num(val, lang, 2)}x", fontsize=4.6, fontweight="bold", color=NEGRO, va="center")
    ax.set_yticks(y); ax.set_yticklabels(nombres, fontsize=4.6, color=NEGRO)
    ax.set_xlim(0, 1.8)
    ax.xaxis.set_major_formatter(lambda t, _: f"{num(t, lang, 1)}x")
    ax.set_ylim(-0.6, 2.75)
    ax.text(1.0, 2.42, "NAV = 1,0x" if es else "NAV = 1.0x", fontsize=3.8, color=ROJO, ha="center", va="bottom")
    ax.tick_params(labelsize=4.0, length=1.5, colors=TINTA_SECUNDARIA)
    ax.grid(axis="x", color=GRILLA, linewidth=0.4, zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_xlabel("Precio / NAV ajustado (EPRA NTA)" if es else "Price / adjusted NAV (EPRA NTA)", fontsize=4.0, color=TINTA_SECUNDARIA, labelpad=2)
    fig.savefig(salida, dpi=DPI, facecolor="white")
    plt.close(fig)


if __name__ == "__main__":
    plt.rcParams.update({"font.family": "DejaVu Sans"})
    cr, pr, nav = leer()
    for lang in ("es", "en"):
        caprate(cr, lang, FIGS / f"{lang}_V18_caprate_spread.png")
        precio(pr, lang, FIGS / f"{lang}_V19_precio_base100.png")
        pnav(nav, lang, FIGS / f"{lang}_V20_pnav.png")
    print("ok", len(cr), len(pr))
