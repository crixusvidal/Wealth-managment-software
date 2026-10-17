"""TIR del proyecto nuevo de Mallplaza (un m² de expansión) frente a la WACC, en nominal CLP y real (UF).

Mide lo que rinde cada m² nuevo (no el negocio completo): el flujo de una cohorte de la hoja Expansion del modelo v3
por m², con el mismo costo, rampa de ingresos y margen que el DCF:
  - Capex (costo por m² del pipeline, indexado al IPC) 50% dos años antes de abrir y 50% el año anterior.
  - Ingresos por m² de la hoja Expansion; rampa 35% / 90% / 100% desde el año de apertura; después solo inflación.
  - EBITDA con el margen de la expansión; impuesto de largo plazo; D&A, mantención (5,5% de ingresos) y capital de trabajo
    (15,6% de la variación de ingresos) como en el valor terminal del DCF.
  - Todo en pesos reales de 2026; la TIR nominal es (1 + real)(1 + π) − 1. Horizonte de 160 años (la cola es despreciable).
Variante «como el RONIC»: sin mantención ni capital de trabajo (el RONIC de la hoja Expansion no los incluye).

Agrega (o reemplaza) la hoja TIR_Proyecto en el modelo v3 con fórmulas vivas, la recalcula con LibreOffice, la verifica con
scipy y vuelca figuras/datos_tir_proyecto.json para documento/valuacion.js.

Uso (desde secciones/): python valuacion/graficos/tir_proyecto.py <Mallplaza_Valoracion_Consolidada_v3.xlsx>
"""
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill
from scipy.optimize import brentq

SALIDA = Path(__file__).resolve().parent / "figuras" / "datos_tir_proyecto.json"
RECALC = "/root/.claude/skills/synced/508cc6bc-ad76-422f-ada9-ad0fc4279625_0b2a2ef9-bf53-46a5-b8ed-72255cce463a/xlsx/scripts/recalc.py"
N = 160                      # años de flujo (2027 en adelante)
F0 = 35                      # primera fila de la tabla de flujos
IT = 50                      # iteraciones de la bisección que resuelve la TIR (LibreOffice no converge con IRR() en 160 períodos)
AMARILLO = PatternFill("solid", fgColor="FFFF00")


def escribir_hoja(ruta):
    wb = openpyxl.load_workbook(ruta)
    if "TIR_Proyecto" in wb.sheetnames:
        del wb["TIR_Proyecto"]
    ws = wb.create_sheet("TIR_Proyecto")
    ws["A1"] = "TIR del proyecto nuevo (un m² de expansión) — NUEVO"
    ws["A1"].font, ws["A1"].fill = Font(bold=True), AMARILLO
    ws["A2"] = ("Flujo por m² de una cohorte (misma mecánica que la hoja Expansion y el valor terminal del DCF), en pesos reales de 2026. "
                "TIR nominal = (1 + real)(1 + π) − 1. Todo con fórmulas vivas sobre Supuestos, Expansion, WACC y DCF.")
    ins = [("Inflación de largo plazo π", "=Supuestos!B34"), ("Tasa de impuesto de largo plazo", "=Supuestos!B59"),
           ("D&A / ingresos", "=Supuestos!B57"), ("Mantención + TI / ingresos", "=Supuestos!B58"),
           ("Capital de trabajo / ingresos", "=Supuestos!B85*Supuestos!B84"), ("Margen EBITDA de la expansión", "=Expansion!C9"),
           ("Costo por m² 2026 (CLP, real)", "=Expansion!C7*1000000"), ("Capex pagado dos años antes de abrir", "=Supuestos!B68"),
           ("Rampa de ingresos: año de apertura", "=Supuestos!B78"), ("Rampa: segundo año", "=Supuestos!B79"), ("Rampa: tercer año en adelante", "=Supuestos!B80")]
    for i, (a, f) in enumerate(ins):
        ws.cell(4 + i, 1, a), ws.cell(4 + i, 2, f)
    ws["A15"], ws["B15"] = "WACC nominal terminal", "=DCF!I25"
    ws["A16"], ws["B16"] = "WACC real terminal", "=WACC!I44"
    ws["A18"], ws["B18"], ws["C18"], ws["D18"] = "Resultado", "Real (UF)", "Nominal CLP", "Brecha vs WACC real / nominal (pb)"
    last = F0 + N - 1
    ws["A19"] = "TIR del proyecto (con mantención y capital de trabajo)"
    ws["B19"] = f"=O{F0 + IT - 1}"
    ws["C19"] = "=(1+B19)*(1+B4)-1"
    ws["D19"], ws["E19"] = "=(B19-B16)*10000", "=(C19-B15)*10000"
    ws["A20"] = "Variante como el RONIC (sin mantención ni capital de trabajo)"
    ws["B20"] = f"=T{F0 + IT - 1}"
    ws["C20"] = "=(1+B20)*(1+B4)-1"
    ws["D20"], ws["E20"] = "=(B20-B16)*10000", "=(C20-B15)*10000"
    ws["A21"], ws["B21"] = "RONIC real de la hoja Expansion (yield × (1 − t))", "=Expansion!C38"
    ws["A22"], ws["B22"] = "Yield on cost de la hoja Expansion (EBITDA maduro / costo)", "=Expansion!C37"
    ws["A24"] = "Ingreso por m² de la expansión (hoja Expansion)"
    for j, (h, f) in enumerate([("Año", None), ("Índice IPC (2026 = 1)", "=Expansion!{c}6"), ("Ingreso por m² nominal (CLP)", "=Expansion!{c}8"),
                                ("Ingreso por m² real 2026 (CLP)", "=C{r}/B{r}")]):
        ws.cell(25, 1 + j, h)
    for k, c in enumerate("CDEFGH"):
        r = 26 + k
        ws.cell(r, 1, 2026 + k), ws.cell(r, 2, f"=Expansion!{c}6"), ws.cell(r, 3, f"=Expansion!{c}8"), ws.cell(r, 4, f"=C{r}/B{r}")
    for j, h in enumerate(["Año", "Factor de rampa", "Ingreso real por m²", "EBITDA", "EBIT", "NOPAT", "Δ capital de trabajo", "Capex", "FCFF real (con mant. y CT)", "FCFF real (como RONIC)"]):
        ws.cell(F0 - 1, 1 + j, h)
    for k in range(N):
        r = F0 + k
        prev = f"C{r - 1}" if k else "0"
        ws.cell(r, 1, 2027 + k)
        ws.cell(r, 2, f"=IF(A{r}<2029,0,IF(A{r}=2029,$B$12,IF(A{r}=2030,$B$13,$B$14)))")
        ws.cell(r, 3, f"=INDEX($D$26:$D$31,MIN(A{r}-2025,6))*B{r}")
        ws.cell(r, 4, f"=C{r}*$B$9")
        ws.cell(r, 5, f"=D{r}-C{r}*$B$6")
        ws.cell(r, 6, f"=E{r}*(1-$B$5)")
        ws.cell(r, 7, f"=$B$8*(C{r}-{prev}/(1+$B$4))")
        ws.cell(r, 8, f"=IF(A{r}=2027,$B$11,IF(A{r}=2028,1-$B$11,0))*$B$10")
        ws.cell(r, 9, f"=F{r}+C{r}*$B$6-C{r}*$B$7-G{r}-H{r}")
        ws.cell(r, 10, f"=F{r}+C{r}*$B$6-H{r}")
    # TIR por bisección (VNA = 0): columnas L:P para el flujo con mantención y CT, Q:U para la variante como el RONIC
    for base_col, flujo, nombre in ((12, "I", "con mant. y CT"), (17, "J", "como el RONIC")):
        for j, h in enumerate(["Iteración", "Mínimo", "Máximo", "Punto medio", "VNA al punto medio"]):
            ws.cell(F0 - 1, base_col + j, f"{h} ({nombre})" if j == 0 else h)
        c0, c1, c2, c3 = (openpyxl.utils.get_column_letter(base_col + k) for k in range(4))
        for i in range(IT):
            r = F0 + i
            ws.cell(r, base_col, i + 1)
            ws.cell(r, base_col + 1, 0.0 if i == 0 else f"=IF({openpyxl.utils.get_column_letter(base_col + 4)}{r - 1}>0,{c3}{r - 1},{c1}{r - 1})")
            ws.cell(r, base_col + 2, 0.2 if i == 0 else f"=IF({openpyxl.utils.get_column_letter(base_col + 4)}{r - 1}>0,{c2}{r - 1},{c3}{r - 1})")
            ws.cell(r, base_col + 3, f"=({c1}{r}+{c2}{r})/2")
            ws.cell(r, base_col + 4, f"=NPV({c3}{r},{flujo}${F0}:{flujo}${last})")
    for c in ("A19", "B19", "C19", "D19", "E19"):
        ws[c].fill = AMARILLO
    wb.save(ruta)


def verificar(wb):
    ex, s, d = wb["Expansion"], wb["Supuestos"], wb["DCF"]
    pi, t, mant, da = s["B34"].value, s["B59"].value, s["B58"].value, s["B57"].value
    kct, ramp = s["B85"].value * s["B84"].value, (s["B78"].value, s["B79"].value, s["B80"].value)
    cost, mg = ex["C7"].value * 1e6, ex["C9"].value
    real = [ex[c + "8"].value / ex[c + "6"].value for c in "CDEFGH"]

    def flujos(mant_, kct_, escala=1.0):
        out, prev = [], 0.0
        for k in range(N):
            y = 2027 + k
            capex = (s["B68"].value if y == 2027 else 1 - s["B68"].value if y == 2028 else 0) * cost
            fac = 0 if y < 2029 else ramp[min(y - 2029, 2)]
            r = real[min(y - 2026, 5)] * fac * escala
            nopat = (r * mg - da * r) * (1 - t)
            out.append(nopat + da * r - mant_ * r - kct_ * (r - prev / (1 + pi)) - capex)
            prev = r
        return np.array(out)

    tir = lambda f: brentq(lambda r: (f / (1 + r) ** np.arange(N)).sum(), -0.05, 0.3)
    base, ronic = tir(flujos(mant, kct)), tir(flujos(0, 0))
    wacc_real = wb["WACC"]["I44"].value
    escala = brentq(lambda e: tir(flujos(mant, kct, e)) - wacc_real, 0.5, 4)
    return base, ronic, escala


def main(ruta):
    escribir_hoja(ruta)
    subprocess.run([sys.executable, RECALC, ruta], check=True, stdout=subprocess.DEVNULL)
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["TIR_Proyecto"]
    base, ronic, escala = verificar(wb)
    assert abs(ws["B19"].value - base) < 1e-6 and abs(ws["B20"].value - ronic) < 1e-6, (ws["B19"].value, base, ws["B20"].value, ronic)
    yield_ = ws["B22"].value
    out = {"real": base, "nom": ws["C19"].value, "brecha_real": ws["D19"].value, "brecha_nom": ws["E19"].value,
           "ronic_variante": {"real": ronic, "nom": ws["C20"].value}, "ronic": ws["B21"].value, "yield": yield_,
           "wacc_real": wb["WACC"]["I44"].value, "wacc_nom": wb["DCF"]["I25"].value, "pi": ws["B4"].value,
           "escala_igualar_wacc": escala, "yield_necesario": yield_ * escala}
    SALIDA.write_text(json.dumps(out, indent=1, ensure_ascii=False))
    print(SALIDA, {k: (round(v * 100, 2) if isinstance(v, float) and abs(v) < 1 else v) for k, v in out.items() if not isinstance(v, dict)})


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
