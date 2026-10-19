"""Datos de los anexos [B], [C], [E] y [F] de Valuación (DCF, supuestos, múltiplos, escenarios).

Vuelca figuras/datos_anexos.json para documento/valuacion.js. Todas las cifras salen de los modelos
(openpyxl, data_only=True); los escenarios bear/bull/tasadores, que no están en el modelo, se calculan
con la réplica del DCF (escenarios_dcf.py, sobre riesgos/graficos/tornado_operacional.py; reproduce el valor del modelo exacto).

Uso:
    python graficos/anexos_valuacion.py <Valoracion_Consolidada_v4.xlsx> <Figuras_Valuacion_recalc.xlsx>
                                        <FA_Exhibits_v7.xlsx> <Tornado_Sensibilidad_Operacional.xlsx>
"""
import copy
import json
import sys
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "riesgos" / "graficos"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from tornado_operacional import leer_modelo, leer_rangos, proyectar  # noqa: E402

COLS = "CDEFGHI"           # 2S26E … 2031E, terminal (hojas DCF, Forecast_Base, Expansion)


def fila(ws, r, cols=COLS):
    return [ws[f"{c}{r}"].value for c in cols]


def anexo_b(wb):
    d, s = wb["DCF"], wb["Sensibilidad"]
    v = lambda c: d[c].value
    acc = v("C48")
    filas = {k: fila(d, r) for k, r in (("ing", 7), ("ebitda", 11), ("margen", 12), ("da", 13), ("ebit", 14),
                                         ("tasa", 15), ("imp", 16), ("nopat", 17), ("capex_base", 19),
                                         ("capex_exp", 20), ("ct", 21), ("fcff", 22), ("wacc", 25), ("fd", 28),
                                         ("vp", 29))}
    return {
        "filas": filas,
        "puente": {"vp_exp": v("C43"), "vp_vt": v("C42"), "ev": v("C44"), "dfn": v("C45"), "nci": v("C46"),
                   "patrimonio": v("C47"), "acciones": acc, "valor": v("C49"),
                   "ps_exp": v("C43") / acc, "ps_vt": v("C42") / acc, "ps_deuda": (v("C45") + v("C46")) / acc},
        "vt": {"wacc_t": v("C38"), "g": v("C39"), "reinv": v("C40"), "vt": v("C41"), "vt_ev": v("C54"),
               "ev_ebitda": v("C55"), "cap_ev": v("C56"), "cap_vt": v("C57"), "wacc_real": v("C58"),
               "m2_maduro": v("C33"), "ebitda_norm": v("C35")},
        "sens": {"g": [s[f"{c}5"].value for c in "BCDE"], "dw": [s[f"A{r}"].value for r in range(6, 11)],
                 "v": [[s[f"{c}{r}"].value for c in "BCDE"] for r in range(6, 11)]},
    }


def anexo_c(wb):
    f, e, d, su, ct = wb["Forecast_Base"], wb["Expansion"], wb["DCF"], wb["Supuestos"], wb["Capital_Trabajo"]
    s = lambda c: su[c].value
    cal = [f[f"{c}5"].value for c in "CDEFGH"]
    return {
        "filas": {"ing_cal": cal, "crec": fila(d, 8, "CDEFGH"), "ing_exp": fila(d, 6, "CDEFGH"),
                  "margen": fila(f, 9, "CDEFGH"), "tasa": fila(f, 11, "CDEFGH"), "capex_usd": fila(f, 12, "CDEFGH"),
                  "usdclp": fila(f, 13, "CDEFGH"), "capex_proy": fila(f, 14, "CDEFGH"), "mant": fila(f, 15, "CDEFGH"),
                  "capex_base": fila(f, 16, "CDEFGH"), "capex_exp": fila(e, 24, "CDEFGH"),
                  "costo_m2": fila(e, 7, "CDEFGH"), "ct": [-x for x in fila(d, 21, "CDEFGH")]},
        "sup": {"ing_1s26": s("B54"), "ebitda_1s26": s("B55"), "ing25": s("B56"), "da": s("B57"), "mant": s("B58"),
                "tax_lp": s("B59"), "gla_anual": s("B66"), "gla_22": s("B63"), "gla_23": s("B64"), "gla_25": s("B65"),
                "primer": s("B67"), "pago_t2": s("B68"), "pipe_usd": s("B69"), "pipe_m2": s("B70"),
                "costo_usd_m2": s("B71"), "mix_cl": s("B75"), "margen_cl": s("B76"), "margen_pe": s("B77"),
                "rampa": [s("B78"), s("B79"), s("B80")], "yield": e["C37"].value, "ronic": e["C38"].value,
                "ct_ing": s("B85"), "g": s("B88"), "dfn_obj": s("B42"), "anio_obj": s("B43"), "spread": s("B45"),
                "payout": s("B48"), "deudores": ct["E6"].value, "anticip": ct["E12"].value, "ctn": ct["E17"].value,
                "ebitda_ltm": s("B13"), "ing_ltm": wb["Comparables"]["J5"].value},
    }


def anexo_e(wb, wf):
    c, po, v7 = wb["Comparables"], wb["Precio_Objetivo"], wf["V7_Multiplos"]
    pares = []
    for r in range(5, 13):
        g = lambda col: c[f"{col}{r}"].value
        pares.append({"nombre": g("A"), "pais": g("C"), "set": g("D"), "chile": g("E"), "mcap": g("F"),
                      "ev": g("I"), "ebitda": g("K"), "m": [g(x) for x in "PQRST"]})
    est = {k: [c[f"{x}{r}"].value for x in "PQRST"] for k, r in (("med_set", 15), ("med_cl", 16), ("pond", 17))}
    metodos = [[po[f"B{r}"].value, po[f"C{r}"].value, po[f"D{r}"].value] for r in range(5, 10)]
    pond = [[v7[f"{x}{r}"].value for x in "ABCDE"] for r in range(26, 31)]
    s = wb["Supuestos"]
    return {"pares": pares, "est": est, "metodos": metodos, "pond": pond,
            "puente": {"mult": po["B5"].value, "ebitda": po["C5"].value, "ev": po["B5"].value * po["C5"].value,
                       "deuda_bbg": s["B14"].value, "nci": c["H5"].value, "acciones": s["B9"].value,
                       "valor": po["D5"].value},
            "peso_cl": s["B93"].value, "peso_dcf": s["B92"].value}


def anexo_f(wb, wf, ruta_fa, ruta_modelo, ruta_tor):
    import escenarios_dcf
    r = escenarios_dcf.calcular(ruta_fa, ruta_modelo, ruta_tor)
    p = r["param"]
    a, w = wb["APV"], wf["WACC_Construccion"]
    t2 = wb["Sensibilidad"]
    return {
        "esc": r["esc"], "efectos": r["efectos"],
        "param": {"dw_bear": p["dw"], "dw_bull": -p["dw"], "g_bull": p["g"], "m_lo": p["m_lo"], "m_hi": p["m_hi"],
                  "margen26": p["margen26"], "tax_bull": p["tax_bull"], "wacc_tas": p["wacc_tas"],
                  "wacc_tas_real": p["wacc_tas_real"], "wacc_t": p["wacc_t"], "g": p["g"], "share_var": p["share_var"],
                  "rangos": p["rangos"], "bajo": p["bajo"], "alto": p["alto"]},
        "apv": {"ku": a["C5"].value, "sin_deuda": a["C22"].value, "escudos": a["C23"].value, "valor": a["C21"].value,
                "dif": wb["DCF"]["C64"].value},
        "rf": {"bono": w["B74"].value, "dcf": w["B77"].value, "po": w["B78"].value, "efecto": w["B76"].value},
        "deuda": {"spread": [t2[f"{c}15"].value for c in "BCD"],   # fila 16 = estructura actual (Supuestos B41)
                  "filas": [[wb["Supuestos"]["B41"].value if r_ == 16 else
                             float(t2[f"A{r_}"].value.rstrip("x").replace(",", "."))]
                            + [t2[f"{c}{r_}"].value for c in "BCD"] for r_ in range(16, 20)]},
        "vf": {"dcf": wb["DCF"]["C61"].value, "wacc": wb["Supuestos"]["B109"].value},
        "implicita": {"wacc": wf["V6_DCFInverso"]["B17"].value, "brecha": wf["V6_DCFInverso"]["B18"].value,
                      "wacc_po": wf["V6_DCFInverso"]["B19"].value},
        "precio": r["precio"], "mult": r["mult"], "peso": r["peso"],
    }


def anexo_nuevo(wb, ruta_fa, ruta_modelo):
    """Secciones nuevas de los Anexos 15 (15D desglose, 15E cuadratura), 16 (16D calendario de deuda) y 18 (18D sensibilidades de la revisión)."""
    d, fb, ex, de, su = wb["DCF"], wb["Forecast_Base"], wb["Expansion"], wb["Deuda"], wb["Supuestos"]
    fx = openpyxl.load_workbook(ruta_fa, data_only=True)["Forecast"]
    cal = lambda r: [fx[f"{c}{r}"].value for c in "DEFGHI"]
    f6 = lambda ws, r: fila(ws, r, "CDEFGH")
    ing, ing_b, ing_e = fila(d, 7), f6(d, 5), f6(d, 6)
    ebitda, ebitda_b, ebitda_e = fila(d, 11), f6(d, 9), f6(d, 10)
    k_ct = su["B84"].value * su["B85"].value
    maduracion = d["C34"].value - fb["H5"].value
    pais = {"ing": {p_: cal(r) for p_, r in (("cl", 91), ("pe", 92), ("co", 93))}, "margen": {p_: cal(r) for p_, r in (("cl", 101), ("pe", 102), ("co", 103))}}
    gla_nuevo = [fx["D41"].value + fx["D42"].value, fx["E41"].value + fx["E42"].value, fx["F41"].value + fx["F42"].value + fx["D45"].value] + [su["B66"].value] * 3
    capex_inc = [a + b for a, b in zip(f6(fb, 14), [-x for x in f6(d, 20)])]
    desglose = {
        "ing_base": ing_b, "ing_exp": ing_e, "ing_pais": pais["ing"], "ebitda_pais": {p_: [a * b for a, b in zip(pais["ing"][p_], pais["margen"][p_])] for p_ in pais["ing"]},
        "margen_pais": {p_: pais["margen"][p_][0] for p_ in pais["margen"]},
        "madura_m2": d["C33"].value, "madura_ing": maduracion, "indexa_ing": d["I7"].value - d["C34"].value, "ing_term": d["I7"].value,
        "ebitda_base": ebitda_b, "ebitda_exp": ebitda_e, "margen_exp": ex["C9"].value, "da": fila(d, 13), "da_term": d["I13"].value, "da_pct": su["B57"].value,
        "ebit": fila(d, 14), "ebit_term": d["I14"].value, "imp": fila(d, 16), "imp_term": d["I16"].value,
        "capex_usd": f6(fb, 12), "usdclp": f6(fb, 13), "capex_pipeline": f6(fb, 14), "capex_mant": f6(fb, 15), "capex_exp": f6(d, 20), "mant_term": d["I19"].value,
        "gla_nuevo": gla_nuevo, "capex_ing": [(-a - b) / c for a, b, c in zip(f6(d, 19), f6(d, 20), ing[:6])] + [-d["I19"].value / d["I7"].value],
        "capex_inc_ing": [x / c for x, c in zip(capex_inc, ing[:6])] + [0.0],
        "dIng_ct": [-x / k_ct for x in fila(d, 21)], "ct_pct": k_ct,
    }
    # 15E: el desglose solo reordena; se recalcula cada línea desde sus componentes y se compara con el DCF (máx. |diferencia| por línea)
    tasa = fila(d, 15)
    comp = {
        "Ingresos totales": [abs(a + b - c) for a, b, c in zip(ing_b, ing_e, ing[:6])],
        "EBITDA total": [abs(a + b - c) for a, b, c in zip(ebitda_b, ebitda_e, ebitda[:6])],
        "D&A": [abs(su["B57"].value * i - da_) for i, da_ in zip(ing, fila(d, 13))],
        "Impuesto operacional": [abs(-(e - su["B57"].value * i) * t - x) for e, i, t, x in zip(ebitda, ing, tasa, fila(d, 16))],
        "NOPAT": [abs((e - su["B57"].value * i) * (1 - t) - x) for e, i, t, x in zip(ebitda, ing, tasa, fila(d, 17))],
        "Capex total (mantención + incremental)": [abs(-(a + b) - (c + e_)) for a, b, c, e_ in zip(f6(fb, 15), capex_inc, f6(d, 19), f6(d, 20))][:0] or
                                                   [abs((m_ + p_ + x_) + (c_ + e_)) for m_, p_, x_, c_, e_ in zip(f6(fb, 15), f6(fb, 14), [-v for v in f6(d, 20)], f6(d, 19), f6(d, 20))],
        "Inversión en capital de trabajo": [abs(-k_ct * dI - x) for dI, x in zip([-x / k_ct for x in fila(d, 21)], fila(d, 21))],
        "FCFF": [abs(n + da_ + cb + ce + cw - fcff) for n, da_, cb, ce, cw, fcff in zip(fila(d, 17), fila(d, 18), fila(d, 19), fila(d, 20), fila(d, 21), fila(d, 22))],
        "VP del FCFF": [abs(a * b - c) for a, b, c in zip(f6(d, 22), f6(d, 28), f6(d, 29))],
    }
    ps = (d["C43"].value + d["C42"].value + d["C45"].value + d["C46"].value) / d["C48"].value
    comp["Valor terminal / VP del VT / EV"] = [abs(d["C43"].value + d["C42"].value - d["C44"].value), abs(d["C41"].value * d["H27"].value - d["C42"].value)]
    cuadratura = {k: max(v) for k, v in comp.items()}
    cuadratura["Valor por acción (CLP)"] = abs(ps - d["C49"].value)
    # 16D: calendario de la deuda (hoja Deuda)
    cd = lambda r: fila(de, r, "CDEFGH")
    deuda = {"apertura": cd(7), "fcff": cd(10), "costo_at": cd(12), "reajuste_at": cd(14), "div": cd(16), "adic": cd(20), "cierre": cd(21), "ratio": cd(22), "bruta": cd(23),
             "ratio0": su["B41"].value, "ratio_obj": su["B42"].value, "spread": su["B45"].value, "costo0": su["B46"].value, "uf": su["B47"].value, "caja": su["B12"].value}
    # 18D: sensibilidades de la revisión (réplica del DCF; mismo método que la hoja Sens_Revision del equipo: validado contra v3, −51 y +2)
    m, _ = leer_modelo(ruta_fa, ruta_modelo)
    base = proyectar(m)[2]
    po = wb["Precio_Objetivo"]
    mult, peso, precio = po["B13"].value, po["B14"].value, po["B16"].value
    x = copy.deepcopy(m)
    esc_t = 0.283 / m["tax_lp"]                 # Chile ~70%, Perú ~20%, Colombia ~10% del EBITDA: 70%·27% + 20%·29,5% + 10%·35% = 28,3%
    x["tax"], x["tax_lp"] = m["tax"] * esc_t, m["tax_lp"] * esc_t
    y = copy.deepcopy(m)
    y["da"] = 0.0088
    cap = d["C56"].value
    gp = (1 - 0.094 / cap) * 376 * su["B102"].value / su["B9"].value          # USD 376 MM a cap rate 9,4% valorado a nuestro cap rate
    sens = [{"clave": "tax", "dcf": proyectar(x)[2], "tasa": 0.283}, {"clave": "da", "dcf": proyectar(y)[2], "tasa": 0.0088}, {"clave": "gp", "dcf": base - gp, "tasa": 0.094}]
    for r in sens:
        r["delta"] = r["dcf"] - base
        r["po"] = peso * r["dcf"] + (1 - peso) * mult
        r["vs"] = r["po"] / precio - 1
    return {"desglose": desglose, "cuadratura": cuadratura, "deuda": deuda, "sens_revision": sens, "dcf": base}


if __name__ == "__main__":
    if len(sys.argv) < 5:
        sys.exit(__doc__)
    ruta_v2, ruta_fig, ruta_fa, ruta_tor = sys.argv[1:5]
    wb = openpyxl.load_workbook(ruta_v2, data_only=True)
    wf = openpyxl.load_workbook(ruta_fig, data_only=True)
    datos = {"B": anexo_b(wb), "C": anexo_c(wb), "E": anexo_e(wb, wf),
             "F": anexo_f(wb, wf, ruta_fa, ruta_v2, ruta_tor), "N": anexo_nuevo(wb, ruta_fa, ruta_v2)}
    salida = Path(__file__).parent / "figuras" / "datos_anexos.json"
    salida.write_text(json.dumps(datos, ensure_ascii=False, indent=1, default=float))
    print(salida)
