"""Sensibilidad del valor por acción a las tasas largas y a la inflación (reales vs nominales).

Idea: una tasa larga nominal se mueve por dos razones y el DCF reacciona distinto a cada una.
  - Tasa REAL (Δr): sube el costo de capital y los flujos no se enteran → pega fuerte al precio.
  - INFLACIÓN esperada (Δπ): sube la tasa nominal (Fisher), pero las rentas reajustables (UF) y el valor terminal
    crecen con ella → el efecto neto es mucho menor. El «mix» de ambas decide cuánto pesa una misma alza de la tasa nominal.

Mecánica (réplica del DCF v2/v3, mismos flujos, mismo puente a patrimonio; con Δ = 0 reproduce CLP 1.856,81):
  Ke_real  = (1 + Ke_nom) / (1 + π) − 1 + λ·Δr          (λ = traspaso de la tasa real al Ke; base 1,0)
  Kd_real  = Kd_real_modelo + Δr
  Ke_nom'  = (1 + Ke_real')(1 + π + Δπ) − 1               (Fisher; igual que la conversión USD→CLP del modelo)
  Kd_nom'  = (1 + Kd_real')(1 + π + Δπ) − 1
  WACC'    = E/V·Ke_nom' + D/V·Kd_nom'·(1 − t)           (pesos D/V por período, hoja WACC)
  FCFF'    = FCFF · [1 − θ + θ·((1 + π + Δπ)/(1 + π))^τ]  (θ = parte del flujo que se reajusta con la inflación)
  g'       = (1 + g)/(1 + π)·(1 + π + Δπ) − 1             (crecimiento real terminal constante)
La DFN al 30-jun-26 y los minoritarios no cambian (el puente a patrimonio es un hecho de hoy).

Empírico (control): regresión semanal del precio de Mallplaza contra la tasa real a 10 años (BCU10), la inflación
implícita (BCP10 − BCU10) y el IPSA, jul-24 a sep-26 (Mallplaza_Tasas_Historico.xlsx).

Salidas (desde secciones/): modelos/Mallplaza_Sensibilidad_Tasas_Inflacion.xlsx (fórmulas vivas), figuras/datos_sens_tasas.json
y {es,en}_V11_sens_tasas_inflacion.png (grilla) y {es,en}_V12_mix_tasas_inflacion.png (mix real/inflación).

Uso: python valuacion/graficos/sensibilidad_tasas.py <modelo_v3.xlsx> <Tasas_Historico.xlsx> [carpeta_figuras]
"""
import json
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.patches import Rectangle
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))  # paleta institucional
import estilo  # noqa: E402,F401
from estilo import AMARILLO, AZUL, CELESTE, GRILLA, NEGRO, ROJO, TINTA_SECUNDARIA, num

THETA, LAMBDA = 0.85, 1.0                      # supuestos de traspaso (editables en la hoja Supuestos)
DR = (-150, -100, -50, 0, 50, 100, 150)         # Δ tasa real, pb
DPI = (-100, -50, 0, 50, 100)                   # Δ inflación esperada, pb
MIX = [k / 10 for k in range(11)]               # parte de la alza de la tasa nominal que viene de inflación
THETAS = (1.0, THETA, 0.70)
TAU = (0.25, 1, 2, 3, 4, 5, 6)                  # años (a mitad de período) desde el 30-jun-26
AZUL_BI = "0000FF"
AMARILLO_XL = PatternFill("solid", fgColor="FFFF00")


# --- Insumos -----------------------------------------------------------------------------------------------
def insumos(v3):
    wb = openpyxl.load_workbook(v3, data_only=True)
    d, w, s = wb["DCF"], wb["WACC"], wb["Supuestos"]
    cols = "CDEFGHI"
    return {
        "pi": s["B34"].value, "g": d["C39"].value, "t": s["B35"].value, "dfn": -d["C45"].value, "min": -d["C46"].value,
        "acc": d["C48"].value, "precio": d["C50"].value, "mult": wb["Precio_Objetivo"]["D5"].value, "peso": wb["Precio_Objetivo"]["B14"].value,
        "fcff": [d[c + "22"].value for c in cols], "frac": [d[c + "26"].value for c in "CDEFGH"],
        "ke": [w[c + "39"].value for c in cols], "kd": [w[c + "40"].value for c in cols], "dv": [w[c + "35"].value for c in cols],
        "dcf": d["C49"].value, "wacc_t": w["I43"].value, "wacc_real_t": w["I44"].value,
    }


def valor(I, dr=0.0, dpi=0.0, theta=THETA, lam=LAMBDA):
    """Valor por acción (CLP) con Δ tasa real y Δ inflación (decimales)."""
    pi, p2 = I["pi"], I["pi"] + dpi
    wacc = []
    for ke, kd, dv in zip(I["ke"], I["kd"], I["dv"]):
        ke2 = (1 + ((1 + ke) / (1 + pi) - 1 + lam * dr)) * (1 + p2) - 1
        kd2 = (1 + kd + dr) * (1 + p2) - 1
        wacc.append((1 - dv) * ke2 + dv * kd2 * (1 - I["t"]))
    flujos = [f * (1 - theta + theta * ((1 + p2) / (1 + pi)) ** t) for f, t in zip(I["fcff"], TAU)]
    cierre, pv = 1.0, 0.0
    for i in range(6):
        pv += flujos[i] * cierre * (1 + wacc[i]) ** (-I["frac"][i] / 2)
        cierre *= (1 + wacc[i]) ** (-I["frac"][i])
    g2 = (1 + I["g"]) / (1 + pi) * (1 + p2) - 1
    vt = flujos[6] / (wacc[6] - g2)
    return (pv + vt * cierre - I["dfn"] - I["min"]) / I["acc"]


def escenarios(I):
    """Filas del cálculo: (Δr pb, Δπ pb, θ, etiqueta). La primera es la base."""
    f = [(0, 0, None, "Base")]
    f += [(dr, dp, None, "Grilla") for dp in DPI for dr in DR]
    for di in (100, -100):
        for th in THETAS:
            for m in MIX:
                f.append(((1 - m) * di, m * di, th, f"Mix Δi={di:+d} θ={th}"))
    return f


# --- Empírico ----------------------------------------------------------------------------------------------
def semanal(historico):
    ws = openpyxl.load_workbook(historico, data_only=True)["Datos"]
    filas = [(r[0], r[1], r[2], r[3], r[4]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and None not in r[1:5]]
    semana = {}
    for f in filas:                                          # último dato de cada semana (viernes)
        semana[f[0].isocalendar()[:2]] = f
    return [semana[k] for k in sorted(semana)]


def regresion(sem):
    p, bcp, bcu, ipsa = (np.array([s[i] for s in sem], float) for i in (1, 2, 3, 4))
    y = np.diff(np.log(p))
    X = np.column_stack([np.ones(len(y)), np.diff(bcu), np.diff(bcp - bcu), np.diff(np.log(ipsa))])
    b = np.linalg.lstsq(X, y, rcond=None)[0]
    e = y - X @ b
    se = np.sqrt(np.diag(e @ e / (len(y) - 4) * np.linalg.inv(X.T @ X)))
    return {"n": len(y), "coef": b.tolist(), "t": (b / se).tolist(), "r2": 1 - e @ e / ((y - y.mean()) @ (y - y.mean()))}


# --- Excel con fórmulas vivas ------------------------------------------------------------------------------
def azul(c):
    c.font = Font(color=AZUL_BI)


def nuevo(c):
    c.fill = AMARILLO_XL


def excel(I, esc, sem, salida):
    bcu10 = sem[-1][3]
    wb = openpyxl.Workbook()
    S = wb.active
    S.title = "Supuestos"
    S["A1"], S["A2"] = "Sensibilidad a tasas largas e inflación — insumos", "Azul = dato pegado del modelo v3 (fuente en col. K); negro = fórmula. Los supuestos de traspaso (B13, B14) son del equipo y se pueden mover."
    esc_ = [("Inflación esperada de largo plazo π (Chile)", I["pi"], "Mallplaza_Valoracion_Consolidada_v3 · Supuestos!B34"),
            ("Crecimiento terminal nominal g", I["g"], "v3 · DCF!C39 (real = 0%)"),
            ("Tasa de impuesto t", I["t"], "v3 · Supuestos!B35"),
            ("DFN al 30-06-2026 (MM CLP)", I["dfn"], "v3 · DCF!C45"),
            ("Minoritarios a libro (MM CLP)", I["min"], "v3 · DCF!C46"),
            ("Acciones (MM)", I["acc"], "v3 · DCF!C48"),
            ("Precio 30-09-2026 (CLP)", I["precio"], "v3 · DCF!C50"),
            ("Valor por múltiplos EV/EBITDA LTM (CLP)", I["mult"], "v3 · Precio_Objetivo!D5"),
            ("Peso del DCF en el precio objetivo", I["peso"], "v3 · Precio_Objetivo!B14"),
            ("θ: parte del flujo que se reajusta con la inflación", THETA, "SUPUESTO del equipo: rentas UF/reajustables y valor terminal; <1 por contratos fijos en CLP y rezago"),
            ("λ: traspaso de la tasa real al Ke (Kd traspasa 1:1)", LAMBDA, "SUPUESTO del equipo: la tasa libre de riesgo mueve el Ke 1:1"),
            ("Escenario único: Δ tasa real (pb)", 100, "Cambiar para ver un caso puntual (resultado en B19)"),
            ("Escenario único: Δ inflación esperada (pb)", 0, "Cambiar para ver un caso puntual")]
    for i, (a, v, f) in enumerate(esc_):
        r = 4 + i
        S.cell(r, 1, a), S.cell(r, 2, v), S.cell(r, 11, f)
        azul(S.cell(r, 2))
    nuevo(S["B13"]), nuevo(S["B14"])
    S["A19"], S["A20"] = "Valor por acción del escenario único (CLP)", "Variación frente a la base"
    S["A22"] = "Período"
    for k, n in enumerate(("2S26E", "2027E", "2028E", "2029E", "2030E", "2031E", "Terminal")):
        S.cell(22, 3 + k, n)
    filas = [("FCFF (MM CLP nominales)", I["fcff"], "v3 · DCF!C22:I22"), ("Fracción de año", I["frac"], "v3 · DCF!C26:H26"),
             ("τ: años a mitad de período desde 30-06-26", list(TAU), "Convención de mitad de período del DCF"),
             ("Ke nominal CLP", I["ke"], "v3 · WACC!C39:I39"), ("Kd real (UF)", I["kd"], "v3 · WACC!C40:I40"), ("D/V", I["dv"], "v3 · WACC!C35:I35")]
    for j, (a, vs, f) in enumerate(filas):
        r = 23 + j
        S.cell(r, 1, a), S.cell(r, 11, f)
        for k, v in enumerate(vs):
            azul(S.cell(r, 3 + k, v))
    # r: FCFF 23, frac 24, tau 25, ke 26, kd 27, dv 28
    C = wb.create_sheet("Calculo")
    C["A1"], C["A2"] = "Cálculo por escenario (réplica del DCF v3 con Δ tasa real y Δ inflación)", "Cada fila es un escenario completo: WACC por período, factores de descuento, flujos ajustados, valor terminal y valor por acción."
    enc = ["Δ tasa real (pb)", "Δ inflación (pb)", "θ", "Escenario"] + [f"WACC {n}" for n in ("2S26E", "2027E", "2028E", "2029E", "2030E", "2031E", "Terminal")] \
        + [f"Factor cierre {k + 1}" for k in range(6)] + [f"Factor medio {k + 1}" for k in range(6)] + [f"FCFF aj. {n}" for n in ("2S26E", "2027E", "2028E", "2029E", "2030E", "2031E", "Terminal")] \
        + ["VP explícito", "g terminal", "Valor terminal", "VP valor terminal", "Valor por acción (CLP)", "Variación vs base"]
    for j, h in enumerate(enc):
        C.cell(4, 1 + j, h)
    cW, cC, cM, cF = 5, 12, 18, 24                      # primera columna de cada bloque (WACC, cierre, medio, flujos)
    cPV, cG, cVT, cPVT, cV, cD = 31, 32, 33, 34, 35, 36
    fila_de = {}
    for n, (dr, dp, th, et) in enumerate(esc):
        r = 5 + n
        fila_de[(round(dr, 6), round(dp, 6), th)] = r
        C.cell(r, 1, dr), C.cell(r, 2, dp), C.cell(r, 4, et)
        C.cell(r, 3, "=Supuestos!$B$13" if th is None else th)
        dr_, dp_, pi = f"$A{r}/10000", f"$B{r}/10000", "Supuestos!$B$4"
        for k in range(7):
            s = L(3 + k)
            C.cell(r, cW + k, f"=(1-Supuestos!{s}$28)*((1+((1+Supuestos!{s}$26)/(1+{pi})-1+Supuestos!$B$14*{dr_}))*(1+{pi}+{dp_})-1)"
                              f"+Supuestos!{s}$28*((1+Supuestos!{s}$27+{dr_})*(1+{pi}+{dp_})-1)*(1-Supuestos!$B$6)")
            C.cell(r, cF + k, f"=Supuestos!{s}$23*(1-$C{r}+$C{r}*((1+{pi}+{dp_})/(1+{pi}))^Supuestos!{s}$25)")
        for k in range(6):
            s = L(3 + k)
            w = f"{L(cW + k)}{r}"
            prev = f"{L(cC + k - 1)}{r}*" if k else ""
            C.cell(r, cC + k, f"={prev}(1+{w})^-Supuestos!{s}$24")
            C.cell(r, cM + k, f"={prev}(1+{w})^-(Supuestos!{s}$24/2)")
        C.cell(r, cPV, f"=SUMPRODUCT({L(cF)}{r}:{L(cF + 5)}{r},{L(cM)}{r}:{L(cM + 5)}{r})")
        C.cell(r, cG, f"=(1+Supuestos!$B$5)/(1+{pi})*(1+{pi}+{dp_})-1")
        C.cell(r, cVT, f"={L(cF + 6)}{r}/({L(cW + 6)}{r}-{L(cG)}{r})")
        C.cell(r, cPVT, f"={L(cVT)}{r}*{L(cC + 5)}{r}")
        C.cell(r, cV, f"=({L(cPV)}{r}+{L(cPVT)}{r}-Supuestos!$B$7-Supuestos!$B$8)/Supuestos!$B$9")
        C.cell(r, cD, f"={L(cV)}{r}/{L(cV)}$5-1")
    # escenario único en Supuestos
    S["B19"] = f"=Calculo!{L(cV)}{5 + len(esc)}"
    S["B20"] = f"=B19/Calculo!{L(cV)}5-1"
    rU = 5 + len(esc)
    C.cell(rU, 1, "=Supuestos!B15"), C.cell(rU, 2, "=Supuestos!B16"), C.cell(rU, 3, "=Supuestos!$B$13"), C.cell(rU, 4, "Escenario único")
    for col in list(range(cW, cF + 7)) + list(range(cPV, cD + 1)):
        src = C.cell(5 + 1, col).value                    # copiar fórmulas de una fila cualquiera al escenario único
        base = C.cell(5, col).value
        from openpyxl.formula.translate import Translator
        C.cell(rU, col, Translator(base, origin=f"{L(col)}5").translate_formula(f"{L(col)}{rU}"))
    # hoja de resultados
    R = wb.create_sheet("Sensibilidad", 0)
    R["A1"], R["A2"] = "Sensibilidad del valor por acción a tasas largas e inflación (NUEVO)", "Valor DCF por acción (CLP) y variación frente a la base (CLP 1.857). Δ tasa real y Δ inflación esperada en puntos base; todo se recalcula desde Calculo."
    R["A4"] = "A. Valor DCF por acción (CLP): filas = Δ inflación esperada (pb), columnas = Δ tasa real (pb)"
    for j, dr in enumerate(DR):
        R.cell(5, 2 + j, dr)
    for i, dp in enumerate(DPI):
        R.cell(6 + i, 1, dp)
        for j, dr in enumerate(DR):
            R.cell(6 + i, 2 + j, f"=Calculo!{L(cV)}{fila_de[(float(dr), float(dp), None)]}")
    R["A12"] = "B. Variación frente a la base"
    for j, dr in enumerate(DR):
        R.cell(13, 2 + j, dr)
    for i, dp in enumerate(DPI):
        R.cell(14 + i, 1, dp)
        for j, dr in enumerate(DR):
            R.cell(14 + i, 2 + j, f"=Calculo!{L(cD)}{fila_de[(float(dr), float(dp), None)]}")
    R["A20"] = "C. Mix: efecto de una misma alza (+100 pb) o baja (−100 pb) de la tasa larga NOMINAL según cuánto venga de inflación (Δi ≈ Δr + Δπ)"
    R["A21"] = "% de la variación que viene de inflación"
    for j, m in enumerate(MIX):
        R.cell(21, 2 + j, m)
    r0 = 22
    for q, (di, th) in enumerate([(di, th) for di in (100, -100) for th in THETAS]):
        R.cell(r0 + q, 1, f"Δi = {di:+d} pb, θ = {th}")
        for j, m in enumerate(MIX):
            R.cell(r0 + q, 2 + j, f"=Calculo!{L(cD)}{fila_de[(round((1 - m) * di, 6), round(m * di, 6), th)]}")
    R["A40"] = "E. Precio objetivo (80% DCF + 20% múltiplos, múltiplos fijos): filas = inflación esperada, columnas = tasa real a 10 años (BCU10, base 30-09-26)"
    R["A41"] = "Tasa real / inflación"
    for j, dr in enumerate((-100, -50, 0, 50, 100)):
        R.cell(41, 2 + j, f"={bcu10}/100+{dr}/10000")
    for i, dp in enumerate((100, 50, 0, -50, -100)):
        R.cell(42 + i, 1, f"=Supuestos!$B$4+{dp}/10000")
        for j, dr in enumerate((-100, -50, 0, 50, 100)):
            R.cell(42 + i, 2 + j, f"=Supuestos!$B$12*Calculo!{L(cV)}{fila_de[(float(dr), float(dp), None)]}+(1-Supuestos!$B$12)*Supuestos!$B$11")
    nuevo(R["A40"]), nuevo(R["D44"])
    R["A29"] = "D. Control empírico (regresión semanal; ver hoja Empirico)"
    R["A30"], R["B30"], R["C30"] = "Variable", "Coef. por +1 pp", "Estadístico t"
    R["A31"], R["A32"], R["A33"], R["A34"], R["A35"] = "Δ tasa real 10 años (BCU10)", "Δ inflación implícita (BCP10 − BCU10)", "Δ ln IPSA (control)", "R²", "N.º de semanas"
    n = len(sem)
    E = wb.create_sheet("Empirico")
    E["A1"], E["A2"] = "Semanal: precio de Mallplaza, tasas a 10 años (BCCh) e IPSA", "Fuente: Mallplaza_Tasas_Historico.xlsx, hoja Datos (último dato de cada semana). Azul = dato; negro = fórmula."
    for j, h in enumerate(("Semana (último día)", "Mallplaza (CLP)", "BCP10 (%)", "BCU10 (%)", "IPSA", "Inflación implícita (%)", "Δ ln precio", "Δ BCU10 (pp)", "Δ infl. implícita (pp)", "Δ ln IPSA")):
        E.cell(4, 1 + j, h)
    for i, s in enumerate(sem):
        r = 5 + i
        for j, v in enumerate(s):
            c = E.cell(r, 1 + j, v)
            azul(c)
        E.cell(r, 1).number_format = "dd-mm-yyyy"
        E.cell(r, 6, f"=C{r}-D{r}")
        if i:
            E.cell(r, 7, f"=LN(B{r}/B{r - 1})"), E.cell(r, 8, f"=D{r}-D{r - 1}"), E.cell(r, 9, f"=F{r}-F{r - 1}"), E.cell(r, 10, f"=LN(E{r}/E{r - 1})")
    y, x = f"Empirico!$G$6:$G${4 + n}", f"Empirico!$H$6:$J${4 + n}"
    lin = f"LINEST({y},{x},TRUE,TRUE)"
    for r, k in ((31, 3), (32, 2), (33, 1)):        # LINEST devuelve los coeficientes en orden inverso de las columnas
        R.cell(r, 2, f"=INDEX({lin},1,{k})"), R.cell(r, 3, f"=INDEX({lin},1,{k})/INDEX({lin},2,{k})")
    R["B34"], R["B35"] = f"=INDEX({lin},3,1)", f"=COUNT({y})"
    R["A37"] = "Lectura: el DCF dice que una alza de 100 pb de tasa real baja el valor ~17%, mientras que la inflación casi no lo mueve; en el dato histórico la sensibilidad es mucho menor y poco significativa porque la acción se re-valorizó con tasas casi planas."
    for h in (R["A1"], R["A4"], R["A12"], R["A20"], R["A29"]):
        h.font = Font(bold=True)
    nuevo(R["A1"])
    wb.save(salida)
    return fila_de


def leer_excel(ruta, fila_de, esc):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    C, R = wb["Calculo"], wb["Sensibilidad"]
    return {"valores": [C.cell(5 + n, 35).value for n in range(len(esc))], "emp": {k: (R.cell(r, 2).value, R.cell(r, 3).value) for k, r in (("real", 31), ("infl", 32), ("ipsa", 33))},
            "r2": R["B34"].value, "n": R["B35"].value}


# --- Figuras -----------------------------------------------------------------------------------------------
T = {
    "es": {"x": "Tasa real a 10 años (BCU)", "y": "Inflación esperada", "base": "Base: CLP {v}",
           "mx": "% de la alza de la tasa larga que viene de inflación", "my": "Δ valor DCF",
           "real": "Todo tasa real", "inf": "Todo inflación", "th": "θ = {t}", "sube": "Tasa larga +100 pb", "xn": "Tasa nominal a 10 años (BCP)", "yn": "% de la alza por inflación", "fila": ("100%", "75%", "50%", "25%", "0%"), "baja": "Tasa larga −100 pb"},
    "en": {"x": "10-year real rate (BCU)", "y": "Expected inflation", "base": "Base: CLP {v}",
           "mx": "% of the long-rate rise that comes from inflation", "my": "Δ DCF value",
           "real": "All real rate", "inf": "All inflation", "th": "θ = {t}", "sube": "Long rate +100 bp", "xn": "10-year nominal rate (BCP)", "yn": "% of the rise from inflation", "fila": ("100%", "75%", "50%", "25%", "0%"), "baja": "Long rate −100 bp"},
}


def pc(x, lang, d=1):
    if d == 0 and 0 < abs(x) < 0.0095:
        d = 1
    return num(x * 100, lang, d).replace("-", "−") + "%"


def fig_grilla(I, lang, salida, bcu10):
    """Mismo estilo que la Figura 25 de Riesgos (WACC vs g): precio objetivo 80/20 por celda, centro = caso base."""
    from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
    t = T[lang]
    drs, dps = (-100, -50, 0, 50, 100), (100, 50, 0, -50, -100)           # inflación mayor arriba
    z = np.array([[I["peso"] * valor(I, dr / 1e4, dp / 1e4) + (1 - I["peso"]) * I["mult"] for dr in drs] for dp in dps])
    ib, jb = dps.index(0), drs.index(0)
    cmap = LinearSegmentedColormap.from_list("mallplaza", [ROJO, "#f2f2f2", CELESTE])
    norm = TwoSlopeNorm(vmin=z.min(), vcenter=z[ib, jb], vmax=z.max())
    fig = plt.figure(figsize=(3550 / 1440, 1.7), dpi=400)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0.2, 0.2, 0.78, 0.78])
    ax.imshow(z, cmap=cmap, norm=norm, aspect="auto")
    for i in range(z.shape[0]):
        for j in range(z.shape[1]):
            r, g, b = cmap(norm(z[i, j]))[:3]
            ax.text(j, i, num(z[i, j], lang), ha="center", va="center", fontsize=5.2, fontweight="bold" if (i, j) == (ib, jb) else "normal",
                    color="white" if 0.299 * r + 0.587 * g + 0.114 * b < 0.55 else NEGRO)
    for k in range(z.shape[0] + 1):
        ax.axhline(k - 0.5, color="white", linewidth=1.1)
    for k in range(z.shape[1] + 1):
        ax.axvline(k - 0.5, color="white", linewidth=1.1)
    ax.add_patch(Rectangle((jb - 0.45, ib - 0.41), 0.9, 0.82, fill=False, edgecolor=AMARILLO, linewidth=1.1, zorder=3))
    ax.set_xticks(range(len(drs)), [num(bcu10 + d / 100, lang, 2) + "%" for d in drs])
    ax.set_yticks(range(len(dps)), [num((I["pi"] + d / 1e4) * 100, lang, 1) + "%" for d in dps])
    ax.tick_params(length=0, colors=NEGRO, labelsize=4.8, pad=2)
    ax.set_xlabel(t["x"], fontsize=5, color=TINTA_SECUNDARIA, labelpad=2)
    ax.set_ylabel(t["y"], fontsize=5, color=TINTA_SECUNDARIA, labelpad=2)
    for sp in ax.spines.values():
        sp.set_visible(False)
    fig.savefig(salida, facecolor="white")
    plt.close(fig)
    return z


def fig_mix(I, lang, salida):
    t = T[lang]
    base = valor(I)
    fig, ax = plt.subplots(figsize=(3550 / 1440, 2.3), dpi=400)
    estilos = {1.0: (CELESTE, (0, (3, 2))), THETA: (ROJO, "-"), 0.70: (AMARILLO, "-")}
    for th, (col, ls) in estilos.items():
        ys = [valor(I, (1 - m) / 100, m / 100, theta=th) / base - 1 for m in MIX]
        ax.plot([m * 100 for m in MIX], [y * 100 for y in ys], color=col, linestyle=ls, linewidth=1.6 if th == THETA else 1.2, label=t["th"].format(t=num(th, lang, 2)))
        if th == THETA:
            ax.annotate(pc(ys[0], lang), (0, ys[0] * 100), textcoords="offset points", xytext=(8, -12), fontsize=6.2, fontweight="bold", color=ROJO)
            ax.annotate(pc(ys[-1], lang), (100, ys[-1] * 100), textcoords="offset points", xytext=(-2, 7), ha="right", fontsize=6.2, fontweight="bold", color=ROJO)
    ax.axhline(0, color=NEGRO, linewidth=0.7)
    ax.set_xlim(-2, 102)
    ax.set_ylim(-20, 3)
    ax.set_xlabel(t["mx"], fontsize=5.6, color=TINTA_SECUNDARIA)
    ax.set_ylabel(t["my"], fontsize=5.4, color=TINTA_SECUNDARIA)
    ax.text(2, 2.0, t["sube"], fontsize=5.4, color=TINTA_SECUNDARIA, va="center")
    ax.tick_params(labelsize=5.6, length=2)
    ax.yaxis.set_major_formatter(lambda v, _: f"{num(v, lang, 0).replace('-', '−')}%")
    ax.xaxis.set_major_formatter(lambda v, _: f"{num(v, lang, 0)}%")
    ax.grid(axis="y", color=GRILLA, linewidth=0.5)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(fontsize=5.4, frameon=False, loc="center right", title=None)
    fig.subplots_adjust(left=0.2, right=0.97, top=0.96, bottom=0.2)
    fig.savefig(salida, facecolor="white")
    plt.close(fig)


def fig_mix_heat(I, lang, salida, bcp10):
    """Mix en estilo heat map: tasa larga nominal (columnas) por parte de la variación que viene de inflación (filas)."""
    from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
    t = T[lang]
    dis, ms = (-100, -50, 0, 50, 100), (1.0, 0.75, 0.5, 0.25, 0.0)       # inflación arriba, tasa real abajo
    z = np.array([[I["peso"] * valor(I, (1 - m) * di / 1e4, m * di / 1e4) + (1 - I["peso"]) * I["mult"] for di in dis] for m in ms])
    ib, jb = 2, dis.index(0)
    cmap = LinearSegmentedColormap.from_list("mallplaza", [ROJO, "#f2f2f2", CELESTE])
    norm = TwoSlopeNorm(vmin=z.min(), vcenter=z[ib, jb], vmax=z.max())
    fig = plt.figure(figsize=(3550 / 1440, 1.7), dpi=400)
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0.2, 0.2, 0.78, 0.78])
    ax.imshow(z, cmap=cmap, norm=norm, aspect="auto")
    for i in range(z.shape[0]):
        for j in range(z.shape[1]):
            r, g, b = cmap(norm(z[i, j]))[:3]
            ax.text(j, i, num(z[i, j], lang), ha="center", va="center", fontsize=5.2, fontweight="bold" if (i, j) == (ib, jb) else "normal",
                    color="white" if 0.299 * r + 0.587 * g + 0.114 * b < 0.55 else NEGRO)
    for k in range(z.shape[0] + 1):
        ax.axhline(k - 0.5, color="white", linewidth=1.1)
    for k in range(z.shape[1] + 1):
        ax.axvline(k - 0.5, color="white", linewidth=1.1)
    ax.add_patch(Rectangle((jb - 0.45, ib - 0.41), 0.9, 0.82, fill=False, edgecolor=AMARILLO, linewidth=1.1, zorder=3))
    ax.set_xticks(range(len(dis)), [num(bcp10 + d / 100, lang, 2) + "%" for d in dis])
    ax.set_yticks(range(len(ms)), [t["fila"][k] for k in range(len(ms))])
    ax.tick_params(length=0, colors=NEGRO, labelsize=4.8, pad=2)
    ax.set_xlabel(t["xn"], fontsize=5, color=TINTA_SECUNDARIA, labelpad=2)
    ax.set_ylabel(t["yn"], fontsize=5, color=TINTA_SECUNDARIA, labelpad=2)
    for sp in ax.spines.values():
        sp.set_visible(False)
    fig.savefig(salida, facecolor="white")
    plt.close(fig)
    return z


if __name__ == "__main__":
    v3, hist = Path(sys.argv[1]), Path(sys.argv[2])
    figs = Path(sys.argv[3]) if len(sys.argv) > 3 else Path(__file__).parent / "figuras"
    I = insumos(v3)
    assert abs(valor(I) - I["dcf"]) < 1e-6, (valor(I), I["dcf"])           # con Δ = 0 reproduce el DCF del modelo
    esc, sem = escenarios(I), semanal(hist)
    xlsx = v3.parent / "Mallplaza_Sensibilidad_Tasas_Inflacion.xlsx"
    fila_de = excel(I, esc, sem, xlsx)
    print("Excel escrito (recalcular con LibreOffice):", xlsx)
    emp = regresion(sem)
    py = [valor(I, dr / 1e4, dp / 1e4, th if th is not None else THETA) for dr, dp, th, _ in esc]
    resumen = {"base": py[0], "theta": THETA, "lambda": LAMBDA, "real_100": valor(I, 0.01) / py[0] - 1, "infl_100": valor(I, 0, 0.01) / py[0] - 1,
               "nominal_100_real": valor(I, 0.01, 0) / py[0] - 1, "nominal_100_infl": valor(I, 0, 0.01) / py[0] - 1, "empirico": emp,
               "precio_objetivo_base": I["peso"] * py[0] + (1 - I["peso"]) * I["mult"],
               "precio_objetivo_real_100": I["peso"] * valor(I, 0.01) + (1 - I["peso"]) * I["mult"]}
    (figs / "datos_sens_tasas.json").write_text(json.dumps(resumen, indent=1, ensure_ascii=False))
    for lang in ("es", "en"):
        fig_grilla(I, lang, figs / f"{lang}_V11_sens_tasas_inflacion.png", sem[-1][3])
        fig_mix(I, lang, figs / f"{lang}_V12_mix_lineas_tasas_inflacion.png")          # versión anterior (líneas)
        fig_mix_heat(I, lang, figs / f"{lang}_V12_mix_tasas_inflacion.png", sem[-1][2])
    print(json.dumps({k: v for k, v in resumen.items() if k != "empirico"}, indent=1), emp)
