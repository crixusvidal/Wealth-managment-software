"""Gráficos de la sección Industry del informe (datos de Excel_DATA_INDUSTRY.xlsx), en español e inglés.

Para cada figura genera:
  - es_<nombre>.png / en_<nombre>.png: versión del informe, SIN título, subtítulo ni fuente (van
    en el Word, ver figuras/pies_de_figura.md). Ancho 4,9 in a 400 dpi: al insertarla a 2,47 in
    (columna de figuras de 3550 twips) la letra queda en ≥ 6 pt efectivos.
  - es_<nombre>_titulado.png / en_<nombre>_titulado.png: copia con título, subtítulo, notas y
    fuente dentro de la imagen, para compartir fuera del Word.
  - altos.json: alto (pulgadas) de cada figura al insertarla con 2,47 in de ancho, para el
    constructor del Word (mismo contenido que ALTOS).

Colores: Mallplaza rojo, Parque Arauco verde, Cencosud morado; el resto con la paleta
institucional de comun/estilo.py.

Uso (desde secciones/):
    python3 industria/graficos/industria.py industria/modelos/Excel_DATA_INDUSTRY.xlsx [carpeta_salida]
"""
import json
import math
import sys
import textwrap
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.offsetbox import AnnotationBbox, OffsetImage
from matplotlib.patches import Patch, Rectangle

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))  # paleta institucional y rutas compartidas
import estilo  # noqa: E402,F401  (registra las carpetas graficos/ de todas las secciones)
from estilo import LOGOS, MORADO_CENCOSUD, VERDE_PARQUE_ARAUCO, num, tinte
from estilo import AMARILLO, AZUL, CELESTE, GRILLA, NEGRO, ROJO, TINTA_SECUNDARIA

COLOR = {"Mallplaza": ROJO, "Parque Arauco": VERDE_PARQUE_ARAUCO, "Cenco Malls": MORADO_CENCOSUD}
TEXTO = {"Mallplaza": "white", "Parque Arauco": NEGRO, "Cenco Malls": "white"}

MODELOS = Path(__file__).resolve().parents[1] / "modelos"
PORTER_XLSX = MODELOS / "Porter_Nuevo.xlsx"                                  # notas por país y finales (equipo, 2Q26)
EFICIENCIA_XLSX = MODELOS / "Eficiencia_EBITDA_GLA_2T26V2.xlsx"             # EBITDA/GLA 2T26 con ajuste multifamily de Parque Arauco
ANCHO = 4.9                  # ancho real de la figura (pulgadas)
ANCHO_WORD = 3550 / 1440     # columna de figuras del Word (≈ 2,47 in)
DPI = 400
FS, FS_M, FS_G = 12, 13, 14.5  # letra mínima / media / destacada (×0,5 al insertar: ≥ 6 pt)
# Alto real de cada figura (pulgadas, con ANCHO de ancho)
ALTO_FIG = {"ind1_ocupacion": 3.3, "ind2_pipeline_gla": 3.3, "ind3_oferta_vs_consumo": 3.1,
            "ind4_vencimiento_contratos": 2.9, "ind5_market_share_paises": 3.0,
            "ind5b_market_share_paises_vertical": 6.0, "ind6_porter": 3.7,
            "ind7_market_share_chile": 3.8, "ind8_ebitda_m2": 2.3}
# Alto al insertar con 2,47 in de ancho (= alto/ancho real × 2,47); se vuelca a altos.json
ALTOS = {k: round(v / ANCHO * ANCHO_WORD, 2) for k, v in ALTO_FIG.items()}

ARCHIVOS_LOGO = {"Mallplaza": "mallplaza.png", "Parque Arauco": "parque_arauco.png",
                 "Cenco Malls": "cencosud_shopping.png"}
ANCHO_LOGO = {"Mallplaza": 0.36, "Parque Arauco": 0.82, "Cenco Malls": 0.85}  # pulgadas


def per(p, lang):
    """'2T26' (es) o '2Q26' (en) a partir de un período del Excel ('2T26', '2Q26', '2026'...)."""
    p = str(p)
    return f"{'2T' if lang == 'es' else '2Q'}{p[-2:]}"


def pc(x, lang, dec=1, signo=False):
    """Porcentaje con formato del idioma; x en unidades (0,123 → 12,3%)."""
    s = num(x * 100, lang, dec) + "%"
    return ("+" + s) if signo and x >= 0 else s


def nueva(clave):
    fig = plt.figure(figsize=(ANCHO, ALTO_FIG[clave]))
    fig.patch.set_facecolor("white")
    return fig


def limpiar(ax, grilla="y"):
    for lado in ax.spines.values():
        lado.set_visible(False)
    ax.tick_params(length=0, colors=TINTA_SECUNDARIA, labelsize=FS)
    if grilla:
        ax.grid(axis=grilla, color=GRILLA, linewidth=0.8)
        ax.set_axisbelow(True)


def guardar(fig, salida):
    fig.savefig(salida, dpi=DPI, facecolor="white")
    plt.close(fig)


def ajustar(fig, textos, margen=0.975):
    """Achica la letra de cada texto hasta que quepa en el ancho de la figura."""
    render = fig.canvas.get_renderer()
    for t in textos:
        while t.get_window_extent(render).x1 > margen * fig.bbox.width and t.get_fontsize() > 7:
            t.set_fontsize(t.get_fontsize() - 0.25)


# 1. Ocupación consolidada -------------------------------------------------------------------------
def ocupacion(wb, lang):
    ws = wb["Ocupacion_consolidada"]
    periodos = [ws.cell(5, c).value for c in range(2, 7)]
    fig = nueva("ind1_ocupacion")
    ax = fig.add_axes([0.135, 0.08, 0.705, 0.76])
    finales = []
    for r in range(6, 9):
        nombre = ws.cell(r, 1).value
        serie = [ws.cell(r, c).value * 100 for c in range(2, 7)]
        ax.plot(range(5), serie, color=COLOR[nombre], linewidth=2.6, marker="o", markersize=6.5,
                markeredgecolor="white", markeredgewidth=1.2, zorder=3, label=nombre)
        finales.append((serie[-1], nombre))
    # Etiquetas finales separadas al menos 0,8 pp para que no se monten
    finales.sort()
    pos = [finales[0][0]]
    for v, _ in finales[1:]:
        pos.append(max(v, pos[-1] + 0.8))
    desplaz = (sum(v for v, _ in finales) - sum(pos)) / len(pos)  # recentra el grupo
    for (v, nombre), y in zip(finales, pos):
        ax.text(4.2, y + desplaz, num(v, lang, 1) + "%",
                va="center", fontsize=FS_M, fontweight="bold", color=COLOR[nombre])
    ax.axvline(3, color=TINTA_SECUNDARIA, linestyle="--", linewidth=1)
    nota = ("Mallplaza incluye\nOpen Plaza desde 2T25" if lang == "es"
            else "Mallplaza includes\nOpen Plaza from 2Q25")
    ax.text(2.92, 90.35, nota, fontsize=FS, color=TINTA_SECUNDARIA, va="bottom", ha="right",
            linespacing=1.15)
    ax.set_xticks(range(5), [per(p, lang) for p in periodos], fontsize=FS, color=NEGRO)
    ax.set_ylim(90, 100)
    ax.set_xlim(-0.2, 4.1)
    ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0f}%")
    limpiar(ax)
    ax.tick_params(axis="x", colors=NEGRO)
    fig.legend(loc="upper center", bbox_to_anchor=(0.5, 0.995), ncol=3, frameon=False, fontsize=FS,
               handlelength=1.4, columnspacing=1.0, handletextpad=0.5)
    p0, p1 = per(periodos[0], lang), per(periodos[-1], lang)
    textos = {
        "es": (f"Tasa de ocupación consolidada, {p0}–{p1}",
               "Mallplaza, Parque Arauco y Cenco Malls · ocupación al cierre de cada 2T (Chile, Perú y Colombia)",
               ["Cada empresa define la ocupación de forma distinta: comparar tendencias, no niveles. "
                "Cenco 2T26 incluye Plaza Central (Bogotá, adquirida en jun-26).",
                "Fuente: reportes de resultados 2T22–2T26 de Mallplaza, Parque Arauco y Cenco Malls; "
                "análisis del equipo."]),
        "en": (f"Consolidated occupancy rate, {p0}–{p1}",
               "Mallplaza, Parque Arauco and Cenco Malls · occupancy at the end of each 2Q (Chile, Peru and Colombia)",
               ["Each company defines occupancy differently: compare trends rather than levels. "
                "Cenco 2Q26 includes Plaza Central (Bogotá, acquired Jun-26).",
                "Source: Mallplaza, Parque Arauco and Cenco Malls 2Q22–2Q26 earnings reports; "
                "Team analysis."]),
    }
    return fig, textos[lang]


# 2. Pipeline de GLA ------------------------------------------------------------------------------
def pipeline(wb, lang):
    ws = wb["PIPELINE"]
    datos = [("Mallplaza", ws["D12"].value, ws["E12"].value),
             ("Parque Arauco", ws["N12"].value, ws["O12"].value),
             ("Cenco Malls", ws["I12"].value, 0)]
    fig = nueva("ind2_pipeline_gla")
    ax = fig.add_axes([0.03, 0.14, 0.94, 0.72])
    blanco = {"facecolor": "white", "edgecolor": "none", "pad": 1.5}
    for x, (nombre, org, inorg) in enumerate(datos):
        ax.bar(x, org, width=0.6, color=COLOR[nombre], linewidth=0)
        ax.text(x, org / 2, num(org, lang), ha="center", va="center", fontsize=FS,
                color=TEXTO[nombre], fontweight="bold")
        if inorg:  # adquisiciones: mismo color del par con trama blanca (sin tintes pastel)
            ax.bar(x, inorg, bottom=org, width=0.6, color=COLOR[nombre], edgecolor="white",
                   linewidth=0, hatch="///")
            ax.plot([x - 0.3, x + 0.3], [org, org], color="white", linewidth=2)
            ax.text(x, org + inorg / 2, num(inorg, lang), ha="center", va="center", fontsize=FS,
                    color=NEGRO, bbox=blanco)
        ax.text(x, org + inorg + 4000, f"{num(org + inorg, lang)} m²", ha="center", va="bottom",
                fontsize=FS_M, fontweight="bold", color=NEGRO)
    ax.set_xticks(range(3), [d[0] for d in datos], fontsize=FS, color=NEGRO)
    ax.set_ylim(0, 200_000)
    ax.set_yticks([])
    limpiar(ax, grilla=None)
    ax.tick_params(axis="x", colors=NEGRO)
    ax.axhline(0, color=GRILLA, linewidth=1)
    ax.set_xlim(-0.5, 2.5)
    org_l, inorg_l = ("Orgánico", "Adquisiciones") if lang == "es" else ("Organic", "Acquisitions")
    fig.legend(handles=[Patch(facecolor=NEGRO, label=org_l),
                        Patch(facecolor=NEGRO, hatch="///", edgecolor="white", label=inorg_l)],
               loc="upper center", bbox_to_anchor=(0.5, 0.995), ncol=2, frameon=False, fontsize=FS,
               handlelength=1.6, handleheight=1.0, columnspacing=1.5)
    textos = {
        "es": ("Pipeline de GLA 2026–2028: orgánico vs adquisiciones",
               "GLA que agregarán Mallplaza, Parque Arauco y Cenco Malls (m²) · solo proyectos con fecha",
               ["GLA inorgánico de Mallplaza en 2S26 no informado (se toma como cero).",
                "Fuente: Plaza S.A., Release 2T26 y Memoria 2025; Parque Arauco, Presentación a "
                "Inversionistas 2T26; Cenco Malls, Presentación Corporativa 2T26; análisis del equipo."]),
        "en": ("GLA pipeline 2026–2028: organic vs acquisitions",
               "GLA to be added by Mallplaza, Parque Arauco and Cenco Malls (m²) · dated projects only",
               ["Mallplaza 2S26 inorganic GLA not disclosed (taken as zero).",
                "Source: Plaza S.A. Release 2Q26 and 2025 Annual Report; Parque Arauco Investor "
                "Presentation 2Q26; Cenco Malls Corporate Presentation 2Q26; Team analysis."]),
    }
    return fig, textos[lang]


# 3. Oferta de GLA vs consumo privado -----------------------------------------------------------
def part_trio_chile(wb):
    """Participación del trío en el GLA de los principales operadores de Chile (slide 7 del Release 2T26): 33% + 29% + 14% = 76%."""
    ws = wb["MS_estatico_2Q26"]
    return sum(r[4] for r in ws.iter_rows(min_row=2, values_only=True)
               if r[0] == "Chile" and r[1] in ("Mallplaza", "Cenco Malls", "Parque Arauco"))


def oferta_demanda(wb, lang):
    """Oferta nueva del trío en Chile como % del GLA de la INDUSTRIA en Chile (no del GLA del propio trío).

    El Excel entrega el crecimiento sobre el GLA del trío (sin oficinas). El GLA de los principales operadores se estima como el GLA del trío dividido
    por su participación en los principales operadores (76%), de modo que el crecimiento sobre la industria es el del trío × 76%.
    Solo mide la oferta del trío (los demás operadores no publican su pipeline): es una cota inferior de la oferta de la industria."""
    ws = wb["OFERTA VS DEMANDA ESPERADA"]
    k = part_trio_chile(wb)
    filas = [(ws.cell(r, 1).value, ws.cell(r, 2).value * k, ws.cell(r, 3).value) for r in range(2, 6)]
    assert all(isinstance(f[1], float) for f in filas), "abrir y guardar el Excel en Excel (vínculo externo)"
    fig = nueva("ind3_oferta_vs_consumo")
    ax = fig.add_axes([0.02, 0.19, 0.96, 0.6])
    xs = [0, 1, 2, 3.4]
    for x, (anio, oferta, consumo) in zip(xs, filas):
        for dx, v, color in ((-0.21, oferta, CELESTE), (0.21, consumo, AMARILLO)):
            ax.bar(x + dx, v * 100, width=0.4, color=color)
            ax.text(x + dx, v * 100 + 0.12, pc(v, lang), ha="center", va="bottom", fontsize=FS,
                    color=NEGRO)
    ax.axvline(2.7, color=GRILLA, linewidth=1.2)
    acum = "Acumulado\n" if lang == "es" else "Cumulative\n"
    ax.set_xticks(xs, [str(f[0]).replace("Acumulado ", acum).replace("-", "–") for f in filas],
                  fontsize=FS, color=NEGRO)
    ax.set_ylim(0, 7.6)
    ax.set_xlim(-0.5, 3.9)
    ax.set_yticks([])
    limpiar(ax, grilla=None)
    ax.tick_params(axis="x", colors=NEGRO)
    ax.axhline(0, color=GRILLA, linewidth=1)
    of_l, co_l = (("Nueva oferta de GLA del trío / GLA de los principales operadores", "Consumo privado real (IPoM)") if lang == "es"
                  else ("New trio GLA supply / main-operators GLA", "Real private consumption (IPoM)"))
    fig.legend(handles=[Patch(color=CELESTE, label=of_l), Patch(color=AMARILLO, label=co_l)],
               loc="upper left", bbox_to_anchor=(0.01, 0.995), ncol=1, frameon=False, fontsize=FS,
               handlelength=1.4)
    textos = {
        "es": ("Nueva oferta de GLA vs consumo privado, Chile 2026–2028",
               "GLA orgánico del trío en Chile como % del GLA de los principales operadores vs crecimiento esperado "
               "del consumo privado real",
               ["Solo GLA orgánico (oferta que depende del consumo). GLA de los principales operadores en Chile ≈ GLA del trío ÷ 76% (participación "
                "del trío en ese universo; no se dispone del GLA total del mercado); no incluye la oferta nueva de otros operadores.",
                "Fuente: reportes 2T26 de Plaza S.A., Parque Arauco y Cenco Malls; Banco Central de "
                "Chile, IPoM sep-2026 (Tabla II.1); análisis del equipo."]),
        "en": ("New GLA supply vs private consumption, Chile 2026–2028",
               "Trio organic GLA in Chile as % of main-operators GLA vs expected real private "
               "consumption growth",
               ["Organic GLA only (consumption-dependent supply). Main-operators GLA in Chile ≈ trio GLA ÷ 76% (trio's share of that "
                "universe; total market GLA is not available); excludes new supply from other operators.",
                "Source: 2Q26 reports of Plaza S.A., Parque Arauco and Cenco Malls; Central Bank "
                "of Chile, IPoM Sep-2026 (Table II.1); Team analysis."]),
    }
    return fig, textos[lang]


# 4. Vencimiento de contratos -------------------------------------------------------------------
def contratos(wb, lang):
    ws = wb["Duración contratos"]
    datos = {}
    for r in range(2, ws.max_row + 1):
        emp, tramo, pct_gla, dur = (ws.cell(r, c).value for c in range(1, 5))
        datos.setdefault(emp, {"tramos": [], "dur": dur})["tramos"].append((tramo, pct_gla))
    empresas = sorted(datos, key=lambda e: datos[e]["dur"])  # mayor duración arriba
    # Vence antes = más riesgo: rojo → amarillo → celeste → azul → negro (solo paleta institucional)
    rampa = [ROJO, AMARILLO, CELESTE, AZUL, NEGRO]
    texto_blanco = {ROJO, AZUL, NEGRO}
    fig = nueva("ind4_vencimiento_contratos")
    ax = fig.add_axes([0.02, 0.12, 0.96, 0.6])
    ax.set_xlim(0, 100)
    ax.set_ylim(-0.4, len(empresas) - 0.2)
    render = fig.canvas.get_renderer()
    anios = "años" if lang == "es" else "yrs"
    for y, emp in enumerate(empresas):
        izq = 0
        for i, (tramo, p) in enumerate(datos[emp]["tramos"]):
            ax.barh(y, p * 100, left=izq, height=0.5, color=rampa[i], edgecolor="white",
                    linewidth=1.2)
            t = ax.text(izq + p * 50, y, num(p * 100, lang, 1), ha="center", va="center",
                        fontsize=FS, color="white" if rampa[i] in texto_blanco else NEGRO)
            ancho_seg = ax.transData.transform((p * 100, 0))[0] - ax.transData.transform((0, 0))[0]
            if t.get_window_extent(render).width > ancho_seg - 2:  # no cabe: va bajo la barra
                t.set(y=y - 0.27, va="top", color=NEGRO)
            izq += p * 100
        ax.text(0, y + 0.3, emp, ha="left", va="bottom", fontsize=FS, fontweight="bold", color=NEGRO)
        prom = "plazo prom." if lang == "es" else "avg. term"
        ax.text(100, y + 0.3, f"{prom} {num(datos[emp]['dur'], lang, 1)} {anios}", ha="right",
                va="bottom", fontsize=FS, fontweight="bold", color=NEGRO)
    ax.set_yticks([])
    ax.set_xticks([])
    limpiar(ax, grilla=None)
    eje = "% del GLA según plazo remanente" if lang == "es" else "% of GLA by remaining term"
    ax.text(0, -0.42, eje, transform=ax.transData, fontsize=FS, color=TINTA_SECUNDARIA, va="top")
    etiquetas = [t.replace("años", anios).replace("-", "–") for t, _ in datos[empresas[0]]["tramos"]]
    fig.legend(handles=[Patch(color=c, label=t) for c, t in zip(rampa, etiquetas)],
               loc="upper left", bbox_to_anchor=(0.0, 0.995), ncol=3, frameon=False, fontsize=FS,
               handlelength=1.3, columnspacing=1.0, handletextpad=0.5)
    textos = {
        "es": ("Perfil de vencimiento de contratos y plazo promedio remanente, 2T26",
               "% del GLA según plazo remanente de los contratos · Cenco Malls, Mallplaza y Parque Arauco",
               ["Cifras consolidadas; la ponderación del plazo promedio puede diferir entre empresas.",
                "Fuente: Cenco Malls, Presentación de Resultados 2T26 (p. 14); Mallplaza, Release 2T26 "
                "(p. 7); Parque Arauco, Reporte 2T26 (p. 7) y Presentación a Inversionistas 2T26 "
                "(lámina 17)."]),
        "en": ("Lease maturity profile and average remaining term, 2Q26",
               "Share of GLA by remaining lease term · Cenco Malls, Mallplaza and Parque Arauco",
               ["Consolidated figures; the weighting of the average term may differ across companies.",
                "Source: Cenco Malls 2Q26 Results Presentation (p.14); Mallplaza Release 2Q26 (p.7); "
                "Parque Arauco 2Q26 Report (p.7) and Investor Presentation 2Q26 (slide 17)."]),
    }
    return fig, textos[lang]


# 5. Market share estático por país -------------------------------------------------------------
def market_share_paises(wb, lang, vertical=False):
    ws = wb["MS_estatico_2Q26"]
    otros = tinte(CELESTE, 0.35)  # "Otros": celeste claro (ya existía así)
    roles = {"Líder fuera del trío": AMARILLO, "Otros": otros}
    paises = {}
    for r in range(2, ws.max_row + 1):
        pais, emp, rol = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value
        paises.setdefault(pais, []).append((emp, rol, ws.cell(r, 5).value, ws.cell(r, 7).value))
    for pais, filas in paises.items():  # "Otros" = 1 − resto (fórmula en el Excel)
        resto = 1 - sum(f[2] for f in filas if f[1] != "Otros")
        paises[pais] = [(e, ro, resto if ro == "Otros" else v, n) for e, ro, v, n in filas]

    clave = "ind5b_market_share_paises_vertical" if vertical else "ind5_market_share_paises"
    fig = nueva(clave)
    es = lang == "es"
    nombres_pais = {"Chile": "Chile", "Perú": "Perú" if es else "Peru", "Colombia": "Colombia"}
    for k, (pais, filas) in enumerate(paises.items()):
        if vertical:  # donas apiladas, una por fila; participación del trío a la derecha
            ax = fig.add_axes([0.0, 0.59 - k * 0.293, 1.0, 0.29])
        else:
            ax = fig.add_axes([0.005 + k * 0.333, 0.0, 0.323, 0.74])
        valores = [f[2] for f in filas]
        colores = [COLOR.get(f[0], roles.get(f[1], otros)) for f in filas]
        ancho_anillo = 0.6 if vertical else 0.62
        cuñas, _ = ax.pie(valores, colors=colores, startangle=90, counterclock=False,
                          wedgeprops={"width": ancho_anillo, "edgecolor": "white", "linewidth": 1.5})
        angulos = [(c.theta1 + c.theta2) / 2 for c in cuñas]
        fuera = [v < 0.08 for v in valores]
        # Etiquetas fuera de la dona: separadas al menos 32° (abanico centrado en su grupo)
        idx = [i for i, f in enumerate(fuera) if f]
        if len(idx) > 1:
            orig = [angulos[i] for i in idx]  # sentido horario: ángulos decrecientes
            pos = [orig[0]]
            for a in orig[1:]:
                pos.append(min(a, pos[-1] - 32))
            corr = (sum(orig) - sum(pos)) / len(pos)
            for i, a in zip(idx, pos):
                angulos[i] = a + corr
        for (emp, rol, v, nota), color, a, f in zip(filas, colores, angulos, fuera):
            ang = math.radians(a)
            marca = "*" if nota and "pro forma" in nota else ""
            radio = 1.27 if f else 1 - ancho_anillo / 2
            claro = color in (otros, VERDE_PARQUE_ARAUCO, AMARILLO)
            ax.text(radio * math.cos(ang), radio * math.sin(ang), f"{v * 100:.0f}%{marca}",
                    ha="center", va="center", fontsize=FS, fontweight="bold",
                    color=NEGRO if f or claro else "white")
        trio = sum(f[2] for f in filas if f[0] in COLOR)
        if vertical:
            ax.text(1.6, 0.62, nombres_pais[pais], ha="left", va="center", fontsize=FS_G,
                    fontweight="bold", color=NEGRO)
            ax.text(1.6, 0.05, "Participación del trío" if es else "Trio share", ha="left",
                    va="center", fontsize=FS, color=TINTA_SECUNDARIA)
            ax.text(1.6, -0.55, f"{trio * 100:.0f}%", ha="left", va="center", fontsize=20,
                    fontweight="bold", color=NEGRO)
            ax.set_xlim(-2.75, 5.0)
            ax.set_ylim(-1.38, 1.38)
        else:
            ax.text(0, 1.78, nombres_pais[pais], ha="center", va="center", fontsize=FS_M,
                    fontweight="bold", color=NEGRO)
            ax.text(0, -1.68, f"{'Trío' if es else 'Trio'}: {trio * 100:.0f}%", ha="center",
                    va="center", fontsize=FS, fontweight="bold", color=NEGRO)
            ax.set_xlim(-1.45, 1.45)
            ax.set_ylim(-1.9, 1.98)
        ax.set_aspect("equal")
    lider, resto_l = ("Líder fuera del trío", "Otros") if es else ("Leader outside the trio", "Others")
    fig.legend(handles=[Patch(color=COLOR["Mallplaza"], label="Mallplaza"),
                        Patch(color=COLOR["Cenco Malls"], label="Cenco Malls"),
                        Patch(color=COLOR["Parque Arauco"], label="Parque Arauco"),
                        Patch(color=AMARILLO, label=lider),
                        Patch(color=otros, label=resto_l)],
               loc="upper center", bbox_to_anchor=(0.5, 0.997), ncol=3 if not vertical else 2,
               frameon=False, fontsize=FS, handlelength=1.2, columnspacing=1.0, handletextpad=0.5)
    textos = {
        "es": ("Participación de mercado por GLA y país, 2T26",
               "Participación en el GLA de los principales operadores de centros comerciales de cada país",
               ["* Pro forma, incluye Gran Plaza (4,3% sin ella). Participación del trío = Mallplaza + "
                "Cenco Malls + Parque Arauco.",
                "Fuente: Plaza S.A., Presentación Institucional 2T26 (lámina 7); análisis del equipo."]),
        "en": ("GLA market share by country, 2Q26",
               "Share of GLA among the main shopping center operators of each country",
               ["* Pro forma, includes Gran Plaza (4.3% without it). Trio share = Mallplaza + "
                "Cenco Malls + Parque Arauco.",
                "Source: Plaza S.A., Institutional Presentation 2Q26 (slide 7); Team analysis."]),
    }
    return fig, textos[lang]


# 6. Cinco fuerzas de Porter --------------------------------------------------------------------
# Notas por país (juicio del equipo), pesos de GLA de Mallplaza y notas finales: Porter_Nuevo.xlsx (2Q26).
# Nota final = promedio ponderado redondeado al entero; en nuevos entrantes el ponderado queda en 1,5 y el
# equipo lo fija en 2 (columna «Nota final (gráfico)» del Excel).
def leer_porter():
    ws = openpyxl.load_workbook(PORTER_XLSX, data_only=True)["Sheet1"]
    pesos = tuple(ws.cell(r, 4).value for r in (7, 8, 9))
    notas = {ws.cell(r, 2).value: tuple(ws.cell(r, c).value for c in (3, 4, 5)) for r in range(16, 21)}
    finales = {ws.cell(r, 2).value: ws.cell(r, 8).value for r in range(16, 21)}
    ponderados = {ws.cell(r, 2).value: ws.cell(r, 6).value for r in range(16, 21)}
    for k, n in notas.items():
        assert abs(sum(x * w for x, w in zip(n, pesos)) - ponderados[k]) < 1e-6, k
    return pesos, notas, finales, ponderados


def porter(wb, lang):
    pesos, notas, finales, ponderados = leer_porter()
    for k in notas:
        if lang == "es":
            print(f"  Porter · {k}: ponderado {ponderados[k]:.2f} → {finales[k]}")
    orden = [("Poder de negociación de los locatarios", "Poder de\nnegociación de\nlos locatarios",
              "Bargaining\npower of\ntenants"),
             ("Poder de negociación de los proveedores", "Poder de\nnegociación de\nlos proveedores",
              "Bargaining\npower of\nsuppliers"),
             ("Amenaza de sustitutos (e-commerce)", "Amenaza de\nsustitutos\n(e-commerce)",
              "Threat of\nsubstitutes\n(e-commerce)"),
             ("Amenaza de nuevos entrantes", "Amenaza de\nnuevos\nentrantes", "Threat of\nnew\nentrants"),
             ("Rivalidad entre competidores", "Rivalidad entre\ncompetidores",
              "Rivalry among\ncompetitors")]
    valores = [finales[k] for k, _, _ in orden]
    n = len(orden)
    ang = np.linspace(0, 2 * np.pi, n, endpoint=False)
    fig = nueva("ind6_porter")
    lado = 1.9  # diámetro del radar (pulgadas)
    ax = fig.add_axes([0.5 - lado / ANCHO / 2, 0.51 - lado / ALTO_FIG["ind6_porter"] / 2,
                       lado / ANCHO, lado / ALTO_FIG["ind6_porter"]], polar=True)
    ax.set_theta_zero_location("N")
    ax.set_theta_direction(-1)
    for nivel in range(1, 6):  # grilla pentagonal
        ax.plot(np.append(ang, ang[0]), [nivel] * (n + 1), color=GRILLA, linewidth=1, zorder=1)
    for a in ang:
        ax.plot([a, a], [0, 5], color=GRILLA, linewidth=1, zorder=1)
    cerrado = np.append(valores, valores[0])
    ax.fill(np.append(ang, ang[0]), cerrado, color=AZUL, alpha=0.12, linewidth=0, zorder=2)
    ax.plot(np.append(ang, ang[0]), cerrado, color=AZUL, linewidth=2.6, zorder=3)
    ax.scatter(ang, valores, s=60, color=AZUL, edgecolor="white", linewidth=1.2, zorder=4)
    for a, v in zip(ang, valores):
        ax.text(a, v + 0.6, f"{v}", ha="center", va="center", fontsize=FS_G, fontweight="bold",
                color=AZUL, zorder=5)
    for nivel in range(2, 6):  # el 1 se omite: chocaría con las notas 1
        ax.text(np.radians(36), nivel, str(nivel), fontsize=FS, color=TINTA_SECUNDARIA, zorder=6,
                ha="center", va="center", bbox={"facecolor": "white", "edgecolor": "none", "pad": 0.3})
    for a, (_, es_l, en_l) in zip(ang, orden):  # etiquetas alineadas hacia afuera
        s, c = math.sin(a), math.cos(a)
        ha = "left" if s > 0.2 else "right" if s < -0.2 else "center"
        va = "bottom" if c > 0.5 else "top" if c < -0.5 else "center"
        ax.text(a, 5.6, es_l if lang == "es" else en_l, ha=ha, va=va, fontsize=FS, color=NEGRO,
                linespacing=1.1, clip_on=False)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_ylim(0, 5)
    ax.grid(False)
    ax.spines["polar"].set_visible(False)
    escala = ("1 = fuerza débil (industria atractiva) · 5 = fuerza fuerte" if lang == "es"
              else "1 = weak force (attractive industry) · 5 = strong force")
    fig.text(0.5, 0.02, escala, ha="center", va="bottom", fontsize=FS, color=TINTA_SECUNDARIA)
    textos = {
        "es": ("Cinco fuerzas de Porter, 2T26",
               "Centros comerciales en Chile, Perú y Colombia · 1 = fuerza débil (industria "
               "atractiva), 5 = fuerza fuerte",
               ["Nota = promedio de Chile, Perú y Colombia ponderado por el GLA de Mallplaza "
                "(61,9% / 26,2% / 11,9%), redondeado al entero; en nuevos entrantes el ponderado "
                "(1,5) se fija en 2.",
                "Fuente: análisis del equipo (Anexo 6)."]),
        "en": ("Porter's five forces, 2Q26",
               "Shopping centers in Chile, Peru and Colombia · 1 = weak force (attractive "
               "industry), 5 = strong force",
               ["Score = average of Chile, Peru and Colombia weighted by Mallplaza's GLA "
                "(61.9% / 26.2% / 11.9%), rounded to the nearest integer; threat of new entrants "
                "(1.5) set at 2.",
                "Source: Team analysis (Appendix 6)."]),
    }
    return fig, textos[lang]


# 7. Market share evolutivo Chile + crecimiento de la industria (base 100) ----------------------
def market_share_chile(wb, lang):
    """Área apilada de GLA en Chile (mismo formato que cuota_mercado.png) con índice 2T22 = 100."""
    ws = wb["MS_evolutivo_Chile"]
    periodos = [per(ws.cell(5, c).value, lang) for c in range(2, 7)]
    filas = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, 7)] for r in (6, 7, 8)}
    empresas = ["Mallplaza", "Parque Arauco", "Cenco Malls"]  # mismo orden que cuota_mercado.png
    gla = np.array([filas[e] for e in empresas], dtype=float) / 1000  # miles de m²
    total = gla.sum(axis=0)
    indice = total / total[0] * 100
    cuota = gla / total
    x = np.arange(len(periodos))
    cagr = (total[-1] / total[0]) ** (1 / (len(x) - 1)) - 1
    es = lang == "es"

    fig = nueva("ind7_market_share_chile")
    ax = fig.add_axes([0.03, 0.27, 0.7, 0.65])
    ax.stackplot(x, gla, colors=[COLOR[e] for e in empresas], edgecolor="white", linewidth=1.5)
    # Columnas de etiquetas: extremos hacia adentro (no se salen del área) y vecinas algo abiertas
    cx = [x[i] + (0.36 if i == 0 else -0.36 if i == len(x) - 1 else 0.1 if i == 1
                  else -0.1 if i == len(x) - 2 else 0) for i in range(len(x))]
    base = np.zeros(len(x))
    for emp, capa, c in zip(empresas, gla, cuota):
        for i in range(len(x)):
            ax.text(cx[i], base[i] + capa[i] / 2, num(c[i] * 100, lang, 1) + "%", ha="center",
                    va="center", fontsize=FS, color=TEXTO[emp])
        centro = (len(x) - 1, base[-1] + capa[-1] / 2)
        logo = LOGOS / ARCHIVOS_LOGO[emp]
        if logo.exists():
            img = plt.imread(logo)
            caja_logo = OffsetImage(img, zoom=ANCHO_LOGO[emp] * 72 / img.shape[1])
            ax.add_artist(AnnotationBbox(caja_logo, centro, xybox=(7, 0), xycoords="data",
                                         boxcoords="offset points", box_alignment=(0, 0.5),
                                         frameon=False, annotation_clip=False))
        else:
            ax.text(centro[0] + 0.08, centro[1], emp, ha="left", va="center", fontsize=FS,
                    fontweight="bold", color=NEGRO)
        base = base + capa
    for i, t in enumerate(total):  # total sobre cada columna
        ax.text(cx[i], t + total.max() * 0.03, num(t, lang), ha="center", va="bottom",
                fontsize=FS, fontweight="bold", color=NEGRO)
    ax.set_xlim(x[0], x[-1])
    ax.set_ylim(0, total.max() * 1.16)
    ax.set_xticks(cx, periodos, fontsize=FS, color=NEGRO)
    ax.set_yticks([])
    limpiar(ax, grilla=None)
    ax.tick_params(axis="x", colors=NEGRO, pad=4)
    # Filas bajo el eje: variación anual e índice (base 2T22 = 100)
    filas_tabla = [("Var. a/a" if es else "YoY", ["–"] + [pc(total[i] / total[i - 1] - 1, lang, 1, True)
                                                          for i in range(1, len(x))]),
                   (f"Índice\n({periodos[0]} = 100)" if es else f"Index\n({periodos[0]} = 100)",
                    [num(v, lang, 1) for v in indice])]
    alto_ax = 0.65 * ALTO_FIG["ind7_market_share_chile"]
    der = 1.025  # rótulos de fila a la derecha, bajo los logos
    for f, (rotulo, vals) in enumerate(filas_tabla):
        y = -(0.45 + f * 0.31) / alto_ax
        ax.text(der, y, rotulo, transform=ax.transAxes, ha="left", va="center",
                fontsize=FS, color=TINTA_SECUNDARIA, linespacing=1.05, clip_on=False)
        for xi, v in zip(cx, vals):
            ax.text(xi, y, v, transform=ax.get_xaxis_transform(), ha="center", va="center",
                    fontsize=FS, color=NEGRO, clip_on=False)
    ax.text(der, total.max() * 1.07, "Total (miles\nde m²)" if es else "Total\n(thousand m²)",
            transform=ax.get_yaxis_transform(), ha="left", va="center", fontsize=FS,
            color=TINTA_SECUNDARIA, linespacing=1.05, clip_on=False)
    textos = {
        "es": (f"GLA en Chile de los operadores listados, {periodos[0]}–{periodos[-1]}",
               f"Mallplaza, Parque Arauco y Cenco Malls · GLA en Chile (miles de m²) y participación "
               f"dentro del trío · índice {periodos[0]} = 100 · {pc(cagr, lang, 1, True)} anual",
               ["Parque Arauco 2T25 incluye la compra de Kennedy Oriente (70.000 m²). El GLA de "
                "Cenco incluye las oficinas de Costanera.",
                "Participación dentro del trío, no del mercado total.",
                "Fuente: reportes de resultados 2T22–2T26 de Mallplaza, Cenco Malls y Parque Arauco; "
                "análisis del equipo."]),
        "en": (f"Chilean GLA of listed mall operators, {periodos[0]}–{periodos[-1]}",
               f"Mallplaza, Parque Arauco and Cenco Malls · GLA in Chile (thousand m²) and share "
               f"of the trio · index {periodos[0]} = 100 · {pc(cagr, lang, 1, True)}/yr",
               ["Parque Arauco 2Q25 includes the Kennedy Oriente acquisition (70,000 m²). Cenco "
                "GLA includes Costanera office space.",
                "Share within the trio, not of the total market.",
                "Source: Mallplaza, Cenco Malls and Parque Arauco 2Q22–2Q26 earnings reports; "
                "Team analysis."]),
    }
    return fig, textos[lang]


# 8. Eficiencia operativa: EBITDA por m² -------------------------------------------------------
def eficiencia(wb, lang):
    ws = openpyxl.load_workbook(EFICIENCIA_XLSX, data_only=True)["Eficiencia_EBITDA_GLA_2T26"]
    filas = [(ws.cell(r, 1).value, ws.cell(r, 4).value * 1e6 / ws.cell(r, 5).value)  # trimestral, EBITDA usado / GLA
             for r in range(6, 9)]
    promedio = sum(f[1] for f in filas) / len(filas)
    brecha = max(f[1] for f in filas) / min(f[1] for f in filas) - 1
    assert abs(promedio - ws["F9"].value) < 1e-6 and abs(brecha - ws["F10"].value) < 1e-9
    filas.sort(key=lambda f: f[1])
    es = lang == "es"
    mil = " mil" if es else "k"
    fig = nueva("ind8_ebitda_m2")
    ax = fig.add_axes([0.29, 0.25, 0.68, 0.72])
    for y, (emp, v) in enumerate(filas):
        ax.barh(y, v / 1000, height=0.6, color=COLOR[emp])
        ax.text(v / 1000 + 1, y, num(v / 1000, lang, 1) + mil, va="center", fontsize=FS_M,
                fontweight="bold", color=NEGRO)
    ax.set_yticks(range(len(filas)), [f[0] for f in filas], fontsize=FS, color=NEGRO)
    ax.set_xlim(0, 82)
    ax.set_xticks([0, 20, 40, 60])
    ax.xaxis.set_major_formatter(lambda v, _: num(v, lang))
    fig.text(0.5, 0.02, "EBITDA trimestral por m² de GLA, 2T26 (miles de CLP)" if es
             else "Quarterly EBITDA per m² of GLA, 2Q26 (CLP thousand)",
             ha="center", va="bottom", color=TINTA_SECUNDARIA, fontsize=FS)
    limpiar(ax, grilla="x")
    ax.tick_params(axis="y", colors=NEGRO)
    textos = {
        "es": ("Eficiencia operativa: EBITDA por m² de GLA, 2T26",
               "EBITDA trimestral 2T26 / GLA al 30-jun-2026 · CLP",
               [f"Promedio simple: CLP {num(promedio, lang)}. Brecha máx.–mín.: {pc(brecha, lang)}. Parque Arauco sin el NOI multifamily (CLP 503 MM).",
                "Fuente: reportes de resultados 2T26 de Mallplaza, Cenco Malls y Parque Arauco; "
                "análisis del equipo."]),
        "en": ("Operating efficiency: EBITDA per m² of GLA, 2Q26",
               "2Q26 quarterly EBITDA / GLA at 30-Jun-2026 · CLP",
               [f"Simple average: CLP {num(promedio, lang)}. Max–min gap: {pc(brecha, lang)}. Parque Arauco excludes multifamily NOI (CLP 503 million).",
                "Source: Mallplaza, Cenco Malls and Parque Arauco 2Q26 earnings reports; "
                "Team analysis."]),
    }
    return fig, textos[lang]


# Copia titulada ---------------------------------------------------------------------------------
def titulado(png, textos, salida):
    """Copia de la figura con título, subtítulo (arriba) y notas + fuente (abajo) en la imagen."""
    titulo, subtitulo, notas = textos
    img = plt.imread(png)
    alto_img = ANCHO * img.shape[0] / img.shape[1]
    lin_t = textwrap.wrap(titulo.upper(), 40)
    lin_s = textwrap.wrap(subtitulo, 70)
    lin_n = [l for nota in notas for l in textwrap.wrap(nota, 84)]
    cab = 0.2 + 0.24 * len(lin_t) + 0.17 * len(lin_s) + 0.12
    pie = 0.14 + 0.15 * len(lin_n)
    h = alto_img + cab + pie
    fig = plt.figure(figsize=(ANCHO, h))
    fig.patch.set_facecolor("white")
    ax = fig.add_axes([0, pie / h, 1, alto_img / h])
    ax.imshow(img)
    ax.axis("off")
    y = h - 0.16
    fig.add_artist(Rectangle((0.03, (y - 0.24 * len(lin_t) + 0.02) / h), 0.05 / ANCHO,
                             (0.24 * len(lin_t) - 0.02) / h, color=ROJO, transform=fig.transFigure))
    for linea in lin_t:
        fig.text(0.03 + 0.1 / ANCHO, y / h, linea, fontsize=12, fontweight="bold", color=NEGRO,
                 va="top")
        y -= 0.24
    y -= 0.04
    for linea in lin_s:
        fig.text(0.03, y / h, linea, fontsize=8.5, color=TINTA_SECUNDARIA, va="top")
        y -= 0.17
    for k, linea in enumerate(lin_n):
        fig.text(0.03, (pie - 0.08 - k * 0.15) / h, linea, fontsize=7.5, color=TINTA_SECUNDARIA,
                 va="top")
    ajustar(fig, fig.texts)
    guardar(fig, salida)


FIGURAS = [("ind1_ocupacion", ocupacion), ("ind2_pipeline_gla", pipeline),
           ("ind3_oferta_vs_consumo", oferta_demanda), ("ind4_vencimiento_contratos", contratos),
           ("ind5_market_share_paises", market_share_paises),
           ("ind5b_market_share_paises_vertical", lambda wb, lang: market_share_paises(wb, lang, True)),
           ("ind6_porter", porter), ("ind7_market_share_chile", market_share_chile),
           ("ind8_ebitda_m2", eficiencia)]


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    salida = Path(sys.argv[2] if len(sys.argv) > 2 else Path(__file__).parent / "figuras")
    salida.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans"})
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    for nombre, funcion in FIGURAS:
        for lang in ("es", "en"):
            fig, textos = funcion(wb, lang)
            png = salida / f"{lang}_{nombre}.png"
            guardar(fig, png)
            titulado(png, textos, salida / f"{lang}_{nombre}_titulado.png")
    (salida / "altos.json").write_text(json.dumps(ALTOS, indent=2, ensure_ascii=False) + "\n")
    print(f"Figuras guardadas en {salida}")
