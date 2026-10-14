"""Motor del tornado de riesgo operacional de Mallplaza (biblioteca).

Replica el motor Forecast del archivo de exhibits y el DCF del modelo consolidado
(expansión continua, capital de trabajo, valor terminal normalizado) y mueve cada variable
operacional a su mínimo y máximo histórico (hoja Supuestos del archivo de tornado), dejando
todo lo demás en el caso base. La usa pagina_riesgos.py (Figura 24).
"""
import sys
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))  # paleta institucional y rutas compartidas
import estilo  # noqa: E402,F401  (registra las carpetas graficos/ de todas las secciones)

AÑOS = "EFGHI"  # 2027E–2031E en la hoja Forecast de Exhibits
CORTO = [0, 1]  # 2027–2028: horizonte del tornado (2026 ya está anclado al 1S26 real)
LARGO = [2, 3, 4]  # 2029–2031
PAISES = {"cl": (55, 56, 58, 59, 46, 49), "pe": (65, 66, 68, 69, 47, 50),
          "co": (75, 76, 78, 79, 48, None)}
# Rango de crecimiento real de largo plazo que usa la compañía en su tasación (Nota 11)
NOTA11 = {"cl": (0.001, 0.011), "pe": (-0.004, 0.012), "co": (0.008, 0.034)}
# Mantención + TI / ingresos: 2024 (~4,8%) y UDM 2T26 (~6,6%); base 5,5% (Supuestos!B47)
CAPEX_MANT = (0.048, 0.066)


def fila(ws, r, cols=AÑOS):
    return np.array([ws[f"{c}{r}"].value or 0.0 for c in cols], dtype=float)


def leer_modelo(ruta_fa, ruta_vf):
    """Motor operativo (Forecast de Exhibits v6) + DCF del modelo consolidado."""
    wb = openpyxl.load_workbook(ruta_fa, data_only=True)
    f, i = wb["Forecast"], wb["Inputs"]
    assert f["D84"].value == 0 and f["D43"].value == 0, "escenarios fuera del caso base"
    m = {"ipc": {"cl": fila(f, 6), "pe": fila(f, 7), "co": fila(f, 8)},
         "fx": {"cl": np.ones(5), "pe": fila(f, 15), "co": fila(f, 16)},
         "fx_pe26": f["D15"].value, "parking_real": fila(f, 53), "paises": {}}
    for p, (r_occ, r_ss, r_park, r_rent, r_gla, r_nuevo) in PAISES.items():
        m["paises"][p] = {
            "occ0": f[f"D{r_occ}"].value, "occ": fila(f, r_occ), "ss": fila(f, r_ss),
            "park0": f[f"D{r_park}"].value, "rent0": f[f"D{r_rent}"].value,
            "gla": f[f"D{r_gla}"].value,
            "nuevo": fila(f, r_nuevo) if r_nuevo else np.zeros(5)}
    m["margen"] = {"cl": fila(f, 101), "pe": fila(f, 102), "co": fila(f, 103)}
    m["margen26"] = {"cl": f["D101"].value, "pe": f["D102"].value, "co": f["D103"].value}
    m["ing26"] = {"cl": f["D91"].value, "pe": f["D92"].value, "co": f["D93"].value}
    m["elim26"] = f["D94"].value
    m["share_var"] = ((i["C43"].value - i["C44"].value + i["C45"].value)
                      / i["C82"].value)  # rentas variables LTM / ingresos LTM

    wb = openpyxl.load_workbook(ruta_vf, data_only=True)
    s, fb, ex, d = wb["Supuestos"], wb["Forecast_Base"], wb["Expansion"], wb["DCF"]
    v = lambda c: s[c].value
    m.update({
        "acciones": v("B9"), "dfn": v("B10"), "nci": v("B11"),
        "ing_1s26": v("B54"), "ebitda_1s26": v("B55"), "ing25": v("B56"),
        "da": v("B57"), "mant": v("B58"), "tax_lp": v("B59"),
        "pago_t2": v("B68"), "mix": v("B75"), "margen_cl": v("B76"), "margen_pe": v("B77"),
        "rampa": (v("B78"), v("B79"), v("B80")), "yield_input": v("B81"),
        "k_ct": v("B84") * v("B85"), "g": v("B88"), "inflacion": v("B34"),
        "tax": fila(fb, 11, "CDEFGH"), "capex_proy": fila(fb, 14, "CDEFGH"),
        "costo_m2": fila(ex, 7, "CDEFGH"), "año": fila(ex, 5, "CDEFGH"),
        "cohortes": [(ex[f"A{r}"].value, ex[f"B{r}"].value) for r in range(12, 17)],
        # WACC por período (v2): cada período se descuenta con su propia WACC
        "waccs": fila(d, 25, "CDEFGH"), "wacc_t": d["I25"].value, "fracciones": fila(d, 26, "CDEFGH")})
    m["wacc"] = m["wacc_t"]
    control = {"ingresos": fila(fb, 5, "CDEFGH"), "fcff": fila(d, 22, "CDEFGHI"),
               "valor": d["C49"].value}
    return m, control


def proyectar(m, d_occ=0.0, d_ss_corto=0.0, d_park_corto=0.0, d_margen=0.0,
              mant=None, ss_largo=None, d_wacc=0.0, g=None):
    """Devuelve ingresos base 2026E–2031E, FCFF 2S26E–terminal y valor por acción."""
    ingresos = np.zeros(5)
    ebitda = np.zeros(5)
    por_m2 = {}
    for p, x in m["paises"].items():
        occ = x["occ"] + d_occ
        ss = x["ss"].copy()
        ss[CORTO] += d_ss_corto
        if ss_largo:
            ss[LARGO] = ss_largo[p]
        park_g = m["parking_real"].copy()
        park_g[CORTO] += d_park_corto
        park, rent, occ_prev = x["park0"], x["rent0"], x["occ0"]
        local = np.zeros(5)
        por_m2[p] = [(rent + park) / x["gla"]]
        for t in range(5):
            ipc = m["ipc"][p][t]
            park = park * (1 + ipc) * (1 + park_g[t])
            rent = rent * (1 + ipc) * (1 + ss[t]) * occ[t] / occ_prev
            occ_prev = occ[t]
            por_m2[p].append((rent + park) / x["gla"])
            local[t] = rent + park + por_m2[p][-1] * x["nuevo"][t]
        clp = local * m["fx"][p]
        ingresos += clp
        ebitda += clp * (m["margen"][p] + d_margen)

    ing26 = sum(m["ing26"].values()) * (1 + m["elim26"])
    ebitda26 = sum(m["ing26"][p] * m["margen26"][p] for p in m["ing26"]) * (1 + m["elim26"])
    base_cal = np.concatenate([[ing26], ingresos])           # Forecast_Base fila 5
    base = np.concatenate([[ing26 - m["ing_1s26"]], ingresos])  # fila 6 (2S26E)
    ebitda_base = np.concatenate([[ebitda26 - m["ebitda_1s26"] + d_margen * (ing26 - m["ing_1s26"])],
                                  ebitda])

    # Expansión continua (hoja Expansion): ingreso por m² del forecast, mix Chile/Perú
    mix = m["mix"]
    rev_m2 = 1e6 * (mix * np.array(por_m2["cl"]) + (1 - mix) * np.array(por_m2["pe"])
                    * np.concatenate([[m["fx_pe26"]], m["fx"]["pe"]]))  # CLP por m²
    margen_exp = mix * m["margen_cl"] + (1 - mix) * m["margen_pe"] + d_margen
    yield_derivado = rev_m2[0] * margen_exp / 1e6 / m["costo_m2"][0]
    yield_usado = m["yield_input"] if isinstance(m["yield_input"], (int, float)) else yield_derivado
    ajuste = yield_usado / yield_derivado
    ronic = yield_usado * (1 - m["tax_lp"])
    ing_exp, capex_exp = np.zeros(6), np.zeros(6)
    for c, m2 in m["cohortes"]:
        for t, a in enumerate(m["año"]):
            factor = 0 if a < c else m["rampa"][min(int(a - c), 2)]
            ing_exp[t] += m2 * rev_m2[t] / 1e6 * ajuste * factor
            capex_exp[t] += m2 * m["costo_m2"][t] * (m["pago_t2"] * (a == c - 2)
                                                      + (1 - m["pago_t2"]) * (a == c - 1))

    # DCF (hoja DCF)
    mant = m["mant"] if mant is None else mant
    g = m["g"] if g is None else g
    rev = base + ing_exp
    ebd = ebitda_base + ing_exp * margen_exp
    da = m["da"] * rev
    nopat = (ebd - da) * (1 - m["tax"])
    capex_base = m["capex_proy"] + mant * np.concatenate([[ing26 / 2], ingresos])
    total_cal = np.concatenate([[ing26 + ing_exp[0]], rev[1:]])
    ct = m["k_ct"] * np.concatenate([[(total_cal[0] - m["ing25"]) / 2],
                                     np.diff(total_cal)])
    fcff = nopat + da - capex_base - capex_exp - ct

    m2_maduro = sum(m2 for c, m2 in m["cohortes"] if c <= 2032)
    ing_norm = ingresos[-1] + m2_maduro * rev_m2[-1] * ajuste / 1e6
    ebitda_norm = ebitda[-1] + m2_maduro * rev_m2[-1] * ajuste / 1e6 * margen_exp
    ebit_norm = ebitda_norm - m["da"] * ing_norm
    g_real = (1 + g) / (1 + m["inflacion"]) - 1
    nopat_t = ebit_norm * (1 + g) * (1 - m["tax_lp"])
    fcff_t = (nopat_t + m["da"] * ing_norm * (1 + g) - mant * ing_norm * (1 + g)
              - m["k_ct"] * ing_norm * g - g_real / ronic * nopat_t)

    w = m["waccs"] + d_wacc                       # desplazamiento paralelo de la WACC (todos los períodos)
    cierre = np.cumprod((1 + w) ** -m["fracciones"])     # factor de descuento al cierre de cada período
    mitad = np.concatenate([[1.0], cierre[:-1]]) * (1 + w) ** -(m["fracciones"] / 2)
    vp = np.sum(fcff * mitad)
    vt = fcff_t / (m["wacc_t"] + d_wacc - g) * cierre[-1]
    valor = (vp + vt - m["dfn"] - m["nci"]) / m["acciones"]
    return base_cal, np.append(fcff, fcff_t), valor


def leer_rangos(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True)["Supuestos"]
    rangos = {}
    for r in range(7, 13):
        rangos[ws[f"B{r}"].value] = tuple(ws[f"{c}{r}"].value for c in "DEF")
    return rangos


def escenarios(m, rangos):
    """(etiqueta, rango mostrado, kwargs bajo, kwargs alto) para cada variable."""
    def delta(nombre):
        base, lo, hi = rangos[nombre]
        return lo - base, hi - base, lo, hi

    out = []
    d_lo, d_hi, lo, hi = delta("Margen EBITDA")
    out.append(("EBITDA margin", f"{lo:.1%} – {hi:.1%}",
                {"d_margen": d_lo}, {"d_margen": d_hi}))
    d_lo, d_hi, lo, hi = delta("Ocupación")
    out.append(("Occupancy", f"{lo:.1%} – {hi:.1%}", {"d_occ": d_lo}, {"d_occ": d_hi}))
    d_lo, d_hi, lo, hi = delta("Same-store rent growth (SSR)")
    out.append(("Same-store rent growth 2027–28", f"{lo:.1%} – {hi:.1%}",
                {"d_ss_corto": d_lo}, {"d_ss_corto": d_hi}))
    out.append(("Long-term real rent growth 2029–31", "appraisal range (Note 11)",
                {"ss_largo": {p: v[0] for p, v in NOTA11.items()}},
                {"ss_largo": {p: v[1] for p, v in NOTA11.items()}}))
    lo_c, hi_c = CAPEX_MANT
    out.append(("Maintenance + IT capex", f"{lo_c:.1%} – {hi_c:.1%} of revenue",
                {"mant": hi_c}, {"mant": lo_c}))
    d_lo, d_hi, lo, hi = delta("Visitas — crecimiento orgánico")
    out.append(("Visitor traffic 2027–28 (parking)", f"{lo:.1%} – {hi:.1%}",
                {"d_park_corto": d_lo}, {"d_park_corto": d_hi}))
    d_lo, d_hi, lo, hi = delta("Same-store sales growth (SSS)")
    s = m["share_var"]
    out.append(("Tenant sales growth 2027–28", f"{lo:.1%} – {hi:.1%} (variable rent)",
                {"d_ss_corto": s * d_lo}, {"d_ss_corto": s * d_hi}))
    return out
