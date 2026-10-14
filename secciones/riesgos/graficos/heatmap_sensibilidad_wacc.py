"""Mapa de calor de la sensibilidad del valor por acción de Mallplaza (WACC vs g).

Muestra el precio objetivo (80% DCF + 20% múltiplos, hoja Precio_Objetivo) en una grilla
5 x 5 con el caso base al centro (modelo v2: WACC por período, desplazamiento paralelo de ±1 pp
vs g de 2,0% a 4,0%). El DCF replica la lógica de la hoja Sensibilidad y se verifica contra su
tabla 1; la columna g = 2,0% no existe en el Excel y se calcula con la misma fórmula.
Rojo = valor bajo el caso base, celeste = sobre el caso base.

Uso:
    python graficos/heatmap_sensibilidad_wacc.py <Valoracion_Consolidada.xlsx> [salida.png]
"""
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
from matplotlib.patches import Rectangle

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "comun"))  # paleta institucional y rutas compartidas
import estilo  # noqa: E402,F401  (registra las carpetas graficos/ de todas las secciones)
from estilo import AMARILLO, CELESTE, NEGRO, ROJO, TINTA_SECUNDARIA

NEUTRO = "#f2f2f2"


PASOS_G = [-0.01, -0.005, 0.0, 0.005, 0.01]   # g alrededor del caso base (5 columnas)
PASOS_W = [-0.01, -0.005, 0.0, 0.005, 0.01]   # desplazamiento paralelo de la WACC (5 filas)


def leer_insumos(ruta):
    """Modelo v2: cada período se descuenta con su propia WACC (hoja DCF, fila 25)."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    d, s, ex = wb["DCF"], wb["Supuestos"], wb["Expansion"]
    v = lambda c: s[c].value
    return {"fcff": np.array([d[f"{c}22"].value for c in "CDEFGH"]),
            "waccs": np.array([d[f"{c}25"].value for c in "CDEFGH"]), "wacc_t": d["I25"].value,
            "frac": np.array([d[f"{c}26"].value for c in "CDEFGH"]),
            "ing_norm": d["C34"].value, "ebitda_norm": d["C35"].value,
            "tax": v("B59"), "da": v("B57"), "mant": v("B58"), "k_ct": v("B84") * v("B85"),
            "inflacion": v("B34"), "ronic": ex["C38"].value,
            "ajustes": v("B10") + v("B11"), "acciones": v("B9")}


def dcf_por_accion(d_wacc, g, x):
    """Valor DCF por acción con la WACC de todos los períodos desplazada en d_wacc (hoja Sensibilidad)."""
    w = x["waccs"] + d_wacc
    cierre = np.cumprod((1 + w) ** -x["frac"])
    mitad = np.concatenate([[1.0], cierre[:-1]]) * (1 + w) ** -(x["frac"] / 2)
    ing, ebitda = x["ing_norm"] * (1 + g), x["ebitda_norm"] * (1 + g)
    ebit = ebitda - x["da"] * ing
    nopat = ebit * (1 - x["tax"])
    g_real = (1 + g) / (1 + x["inflacion"]) - 1
    fcff_t = (nopat + x["da"] * ing - x["mant"] * ing - x["k_ct"] * (ing - x["ing_norm"])
              - g_real / x["ronic"] * nopat)
    vt = fcff_t / (x["wacc_t"] + d_wacc - g) * cierre[-1]
    return (np.sum(x["fcff"] * mitad) + vt - x["ajustes"]) / x["acciones"]


def grilla(ruta):
    """Precio objetivo (peso DCF × DCF + resto × múltiplos) por WACC (filas, mayor a menor) y g (columnas).

    El DCF de cada celda se recalcula con la lógica de la hoja Sensibilidad y se verifica contra ella;
    el valor por múltiplos queda fijo (no depende de la WACC ni de g).
    """
    wb = openpyxl.load_workbook(ruta, data_only=True)
    x = leer_insumos(ruta)
    po, sens, s = wb["Precio_Objetivo"], wb["Sensibilidad"], wb["Supuestos"]
    g_base, peso, multiplos = s["B88"].value, po["B14"].value, po["B13"].value
    for r in range(6, 11):      # verificación contra la tabla 1 del Excel
        for c in range(2, 6):
            esperado = sens.cell(r, c).value
            calculado = dcf_por_accion(sens.cell(r, 1).value, sens.cell(5, c).value, x)
            assert abs(calculado - esperado) < 1e-6, (r, c, calculado, esperado)
    gs = [round(g_base + p, 6) for p in PASOS_G]
    deltas = sorted(PASOS_W, reverse=True)       # +1 pp arriba, −1 pp abajo
    waccs = [x["wacc_t"] + dw for dw in deltas]
    z = np.array([[peso * dcf_por_accion(dw, g, x) + (1 - peso) * multiplos for g in gs] for dw in deltas])
    i_base, j_base = deltas.index(0.0), gs.index(round(g_base, 6))
    assert abs(z[i_base, j_base] - po["B15"].value) < 1e-6, "el centro debe ser el precio objetivo"
    return waccs, gs, z, i_base, j_base, po["B16"].value


def pct(x, decimales=1):
    return f"{x * 100:.{decimales}f}%"


def miles(x):
    return f"{x:,.0f}"


def color_texto(rgba):
    r, g, b = rgba[:3]
    return "white" if 0.299 * r + 0.587 * g + 0.114 * b < 0.55 else NEGRO


def graficar(ruta_xlsx, salida):
    waccs, GS, z, i_base, j_base, precio = grilla(ruta_xlsx)
    wacc_base, g_base = waccs[i_base], GS[j_base]
    base = z[i_base, j_base]
    d = openpyxl.load_workbook(ruta_xlsx, data_only=True)["DCF"]
    ronic, wacc_real = d["C59"].value, d["C58"].value
    multiplos = openpyxl.load_workbook(ruta_xlsx, data_only=True)["Precio_Objetivo"]["B13"].value

    cmap = LinearSegmentedColormap.from_list("mallplaza", [ROJO, NEUTRO, CELESTE])
    norm = TwoSlopeNorm(vmin=z.min(), vcenter=base, vmax=z.max())

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig = plt.figure(figsize=(12, 7.5), dpi=200)
    fig.patch.set_facecolor("white")

    fig.add_artist(Rectangle((0.05, 0.905), 0.008, 0.055, color=ROJO,
                             transform=fig.transFigure))
    fig.text(0.065, 0.915, "TARGET PRICE SENSITIVITY", fontsize=22,
             fontweight="bold", color=NEGRO)
    fig.text(0.05, 0.87, "Mallplaza · target price (CLP) = 80% DCF + 20% EV/EBITDA multiples, by nominal WACC and terminal growth",
             fontsize=11, color=TINTA_SECUNDARIA)

    ax = fig.add_axes([0.1, 0.24, 0.87, 0.58])
    im = ax.imshow(z, cmap=cmap, norm=norm, aspect="auto")
    for i in range(z.shape[0]):
        for j in range(z.shape[1]):
            ax.text(j, i, miles(z[i, j]), ha="center", va="center", fontsize=14,
                    fontweight="bold" if (i, j) == (i_base, j_base) else "normal",
                    color=color_texto(cmap(norm(z[i, j]))))
    for i in range(z.shape[0] + 1):
        ax.axhline(i - 0.5, color="white", linewidth=3)
    for j in range(z.shape[1] + 1):
        ax.axvline(j - 0.5, color="white", linewidth=3)
    ax.add_patch(Rectangle((j_base - 0.46, i_base - 0.44), 0.92, 0.88, fill=False,
                           edgecolor=AMARILLO, linewidth=3, zorder=3))

    ax.set_xticks(range(len(GS)), [pct(g, 2) for g in GS])
    ax.set_yticks(range(len(waccs)), [pct(w, 2) for w in waccs])
    ax.set_xlabel("Terminal growth rate (g, nominal)", color=TINTA_SECUNDARIA,
                  labelpad=10)
    ax.set_ylabel("Terminal WACC (all periods shifted in parallel, ±1 pp)", color=TINTA_SECUNDARIA, labelpad=10)
    ax.tick_params(length=0, colors=NEGRO, labelsize=12)
    for lado in ax.spines.values():
        lado.set_visible(False)

    cax = fig.add_axes([0.1, 0.09, 0.87, 0.02])
    cb = fig.colorbar(im, cax=cax, orientation="horizontal")
    cb.set_ticks([z.min(), base, z.max()])
    cb.set_ticklabels([miles(z.min()),
                       f"{miles(base)}\ntarget price (terminal WACC {pct(wacc_base, 2)} · g {pct(g_base)})",
                       miles(z.max())])
    cb.outline.set_visible(False)
    cax.tick_params(length=0, colors=TINTA_SECUNDARIA, labelsize=9)
    cax.set_title("Target price (CLP)", fontsize=9, color=TINTA_SECUNDARIA,
                  loc="left")

    fig.text(0.05, 0.015,
             "Multiples value held at CLP " + miles(multiplos) + ". Higher g lowers value: terminal growth needs reinvestment at a " + pct(ronic) + " real RONIC, below the "
             + pct(wacc_real) + " real WACC. Share price CLP " + miles(precio) + ". Source: Team analysis.",
             fontsize=8, color=TINTA_SECUNDARIA)

    fig.savefig(salida, facecolor="white")
    print(f"Gráfico guardado en {salida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    graficar(sys.argv[1],
             sys.argv[2] if len(sys.argv) > 2 else Path(__file__).parent / "figuras" / "heatmap_sensibilidad_wacc.png")
