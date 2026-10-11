"""Crea data/macro/macro_2022_2026.xlsx: datos del Anexo macro con fuente y estado de cada cifra.

Estado: "Verificado" = cifra citada por la fuente indicada (búsqueda oct-2026);
"Aproximado" = valor de la serie pública según el equipo, a reemplazar con Bloomberg antes de enviar;
"Interpolado" = mes sin dato publicado, interpolado entre meses verificados.
Uso: python data/macro/crear_macro_xlsx.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

MESES = [f"{a}-{m:02d}" for a in range(2022, 2027) for m in range(1, 13)][:57]  # ene-22 a sep-26

BRENT = [86.5, 97.1, 117.3, 104.6, 113.3, 122.7, 111.9, 100.5, 89.8, 93.3, 91.4, 80.9,
         82.5, 82.6, 78.4, 84.6, 75.5, 74.8, 80.1, 86.2, 94.0, 90.6, 82.9, 77.6,
         80.1, 83.5, 85.4, 90.1, 81.8, 82.6, 85.3, 80.9, 74.3, 75.6, 74.4, 73.9,
         79.0, 75.0, 72.0, 67.0, 64.0, 71.0, 71.0, 68.0, 68.0, 64.0, 63.0, 63.0,
         69.0, 71.0, 103.0, 117.0, 107.0, 85.0, 84.0, 91.0, 114.0]
BRENT_EST = {**{m: ("Aproximado", "EIA, Europe Brent spot (RBRTE), promedio mensual") for m in MESES[:36]},
             **{m: ("Aproximado", "EIA; total anual 2025 = USD 69 (EIA)") for m in MESES[36:48]},
             "2025-01": ("Verificado", "EIA Today in Energy: USD 79 en enero 2025"),
             "2025-07": ("Verificado", "EIA STEO: USD 71 en julio 2025"),
             "2025-08": ("Verificado", "EIA STEO: USD 68 en agosto 2025"),
             "2025-09": ("Verificado", "EIA STEO oct-25 (Rigzone): USD 68 en septiembre 2025"),
             "2025-12": ("Verificado", "EIA: USD 63 en diciembre 2025, mínimo desde 2021"),
             "2026-01": ("Interpolado", "Promedio 1T26 USD 81 (EIA STEO abr-26) menos feb y mar"),
             "2026-02": ("Verificado", "EIA STEO abr-26: marzo USD 103, USD 32 sobre febrero"),
             "2026-03": ("Verificado", "EIA STEO abr-26: USD 103"),
             "2026-04": ("Verificado", "Statista: ~USD 117 en abril 2026"),
             "2026-05": ("Verificado", "Statista: ~USD 107 en mayo 2026"),
             "2026-06": ("Verificado", "Statista: ~USD 85 en junio 2026"),
             "2026-07": ("Interpolado", "EIA STEO sep-26: agosto USD 91, USD 7 sobre julio"),
             "2026-08": ("Verificado", "EIA STEO sep-26: USD 91"),
             "2026-09": ("Verificado", "EIA STEO oct-26: USD 114")}

UST = [1.76, 1.93, 2.13, 2.75, 2.90, 3.14, 2.90, 2.90, 3.52, 3.98, 3.89, 3.62,
       3.53, 3.75, 3.66, 3.46, 3.57, 3.75, 3.90, 4.17, 4.38, 4.80, 4.50, 4.02,
       4.06, 4.21, 4.21, 4.54, 4.48, 4.31, 4.25, 3.87, 3.72, 4.10, 4.36, 4.39,
       4.63, 4.45, 4.28, 4.28, 4.42, 4.38, 4.39, 4.26, 4.12, 4.06, 4.09, 4.14,
       4.05, 3.95, 4.25, 4.40, 4.48, 4.47, 4.60, 4.68, 4.99]
BUND = [-0.06, 0.16, 0.39, 0.82, 0.97, 1.45, 1.00, 1.00, 1.79, 2.17, 2.00, 2.06,
        2.25, 2.39, 2.32, 2.36, 2.33, 2.39, 2.47, 2.53, 2.71, 2.83, 2.61, 2.12,
        2.18, 2.34, 2.35, 2.48, 2.53, 2.49, 2.42, 2.22, 2.17, 2.23, 2.29, 2.18,
        2.51, 2.45, 2.74, 2.55, 2.58, 2.54, 2.64, 2.71, 2.70, 2.62, 2.67, 2.94,
        2.95, 2.85, 3.10, 3.15, 3.00, 2.93, 3.10, 3.30, 3.64]
OAT = [0.30, 0.62, 0.79, 1.34, 1.48, 2.00, 1.60, 1.60, 2.37, 2.78, 2.53, 2.62,
       2.82, 2.88, 2.85, 2.89, 2.92, 2.94, 3.04, 3.08, 3.27, 3.47, 3.25, 2.69,
       2.74, 2.85, 2.83, 2.97, 3.03, 3.20, 3.14, 2.99, 2.94, 3.02, 3.10, 3.10,
       3.30, 3.17, 3.45, 3.30, 3.25, 3.24, 3.33, 3.42, 3.50, 3.40, 3.45, 3.55,
       3.55, 3.40, 3.65, 3.80, 3.80, 3.77, 3.85, 4.00, 4.45]
CHILE = [5.90, 6.20, 6.60, 6.70, 6.60, 6.90, 6.50, 6.30, 6.60, 6.90, 6.00, 5.40,
         5.30, 5.60, 5.40, 5.30, 5.20, 5.40, 5.50, 5.80, 6.20, 6.60, 6.00, 5.40,
         5.40, 5.60, 5.80, 6.10, 6.20, 6.10, 5.90, 5.60, 5.40, 5.70, 5.90, 5.90,
         5.90, 5.90, 6.00, 5.90, 5.80, 5.70, 5.60, 5.60, 5.60, 5.50, 5.50, 5.40,
         5.35, 5.30, 5.40, 5.46, 5.57, 5.50, 5.60, 5.80, 6.10]
TASAS_EST = {
    ("UST", "2026-05"): ("Verificado", "FRED GS10: 4,48"), ("UST", "2026-06"): ("Verificado", "FRED GS10: 4,47"),
    ("UST", "2026-07"): ("Verificado", "FRED GS10: 4,60"), ("UST", "2026-08"): ("Verificado", "FRED GS10: 4,68"),
    ("UST", "2026-09"): ("Verificado", "FRED GS10: 4,99 (posible promedio parcial; diario 5,26 el 29-sep)"),
    ("BUND", "2025-12"): ("Verificado", "Bundesbank: 2,94"), ("BUND", "2026-03"): ("Verificado", "Bundesbank: 3,10"),
    ("BUND", "2026-06"): ("Verificado", "Bundesbank: 2,93"), ("BUND", "2026-09"): ("Verificado", "Bundesbank: 3,64"),
    ("OAT", "2026-08"): ("Verificado", "FRED (OECD): 4,00, máximo desde oct-2008"),
    ("CHILE", "2026-04"): ("Verificado", "OECD vía GuruFocus: 5,46"),
    ("CHILE", "2026-05"): ("Verificado", "Cbonds 5,57 (5-jun); Investing 5,52-5,62"),
}

MUNDO = [("2022", 3.6, 8.7), ("2023", 3.5, 6.7), ("2024", 3.3, 5.8), ("2025", 3.3, 4.1), ("2026P", 3.0, 4.7)]
CHILE_MACRO = [("2022", 2.2, 12.8, 11.25), ("2023", 0.5, 3.9, 8.25), ("2024", 2.6, 4.5, 5.00),
               ("2025", 2.4, 3.5, 4.50), ("2026P", 0.5, 3.9, 4.50)]

AZUL = Font(color="0000FF")
NEGRITA = Font(bold=True)
ESTADO_FILL = {"Verificado": "E2EFDA", "Aproximado": "FCE4D6", "Interpolado": "FFF2CC"}


def hoja(wb, nombre, encabezado, filas, nota):
    ws = wb.create_sheet(nombre)
    ws.append([nota])
    ws["A1"].font = Font(italic=True, color="5F5F5F")
    ws.append(encabezado)
    for c in ws[2]:
        c.font = NEGRITA
    for f in filas:
        ws.append(list(f))
        r = ws.max_row
        for c in ws[r][1:]:
            if isinstance(c.value, (int, float)):
                c.font = AZUL
        estado = next((v for v in f if v in ESTADO_FILL), None)
        if estado:
            for c in ws[r]:
                c.fill = PatternFill("solid", fgColor=ESTADO_FILL[estado])
    for col, ancho in zip("ABCDEFGHIJ", [10, 12, 12, 12, 12, 14, 70, 14, 60, 10]):
        ws.column_dimensions[col].width = ancho
    return ws


wb = Workbook()
wb.remove(wb.active)
ley = wb.create_sheet("Leyenda")
for fila in [["Datos del Anexo macro (2022–2026). Azul = dato ingresado. Reemplazar las filas Aproximado/Interpolado con Bloomberg antes de enviar y volver a correr graficos/anexo_macro.py."],
             ["Verificado", "Cifra citada por la fuente de la columna Fuente (búsqueda oct-2026)."],
             ["Aproximado", "Valor de la serie pública (EIA/FRED/Bundesbank/BCCh) según el equipo; verificar con Bloomberg."],
             ["Interpolado", "Mes sin dato publicado; derivado de meses verificados."]]:
    ley.append(fila)
for r, est in ((2, "Verificado"), (3, "Aproximado"), (4, "Interpolado")):
    ley.cell(r, 1).fill = PatternFill("solid", fgColor=ESTADO_FILL[est])
ley.column_dimensions["A"].width = 14
ley.column_dimensions["B"].width = 100

hoja(wb, "Brent", ["Mes", "USD/barril", "Estado", "Fuente"],
     [(m, v, *BRENT_EST.get(m, ("Aproximado", "EIA"))) for m, v in zip(MESES, BRENT)],
     "Brent spot, promedio mensual (USD/barril).")
filas = []
for m, u, b, o, c in zip(MESES, UST, BUND, OAT, CHILE):
    est = [TASAS_EST.get((k, m)) for k in ("UST", "BUND", "OAT", "CHILE")]
    verif = [f"{k}: {e[1]}" for k, e in zip(("UST", "BUND", "OAT", "CHILE"), est) if e]
    filas.append((m, u, b, o, c, "Verificado" if len(verif) == 4 else "Aproximado",
                  "; ".join(verif) or "FRED GS10 / Bundesbank / Banque de France / BCCh (BCP 10Y)"))
hoja(wb, "Tasas10Y", ["Mes", "EE.UU.", "Alemania", "Francia", "Chile", "Estado", "Verificados en el mes"],
     filas, "Rendimiento soberano a 10 años, promedio mensual (%). Chile = BCP/BTP nominal en pesos.")
hoja(wb, "Mundo", ["Año", "PIB mundial (%)", "Inflación mundial (%)", "Estado", "Fuente"],
     [(a, g, i, "Verificado" if a in ("2025", "2026P") else "Aproximado",
       "FMI WEO jul-26: crecimiento 2026 3,0%; inflación 2025 4,1% y 2026 4,7%" if a in ("2025", "2026P")
       else "FMI WEO (datos efectivos)") for a, g, i in MUNDO],
     "FMI, World Economic Outlook. 2026P = proyección (actualización jul-2026).")
hoja(wb, "Chile", ["Año", "PIB (%)", "IPC dic (%)", "TPM fin de año (%)", "Estado", "Fuente"],
     [(a, g, i, t, "Verificado" if a == "2026P" else "Aproximado",
       "IPoM: PIB 2026 0,25–0,75% (punto medio), inflación 3,9% (IPoM jun-26), TPM 4,5%" if a == "2026P"
       else "BCCh / INE (PIB 2025 = estimación IPoM dic-25)") for a, g, i, t in CHILE_MACRO],
     "Chile: crecimiento del PIB, IPC a diciembre y TPM al cierre de cada año. 2026P = proyección / último dato.")
salida = Path(__file__).with_name("macro_2022_2026.xlsx")
wb.save(salida)
print(salida)
