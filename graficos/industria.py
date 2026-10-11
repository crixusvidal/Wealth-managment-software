"""Gráficos de la sección Industry del informe (datos de Excel_DATA_INDUSTRY.xlsx).

Sin título ni fuente dentro de la imagen: ambos van en el pie de figura del Word
(ver graficos/industria/pies_de_figura.md). Colores: Mallplaza rojo, Parque Arauco verde,
Cencosud morado; el resto con la paleta corporativa de Mallplaza.

Uso:
    python graficos/industria.py <Excel_DATA_INDUSTRY.xlsx> [carpeta_salida]
"""
import math
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.colors import to_rgb
from matplotlib.patches import Patch

from cuota_mercado import MORADO_CENCOSUD, VERDE_PARQUE_ARAUCO
from sensibilidad_wacc import AMARILLO, CELESTE, GRILLA, NEGRO, ROJO, TINTA_SECUNDARIA

COLOR = {"Mallplaza": ROJO, "Parque Arauco": VERDE_PARQUE_ARAUCO, "Cenco Malls": MORADO_CENCOSUD}
TEXTO = {"Mallplaza": "white", "Parque Arauco": NEGRO, "Cenco Malls": "white"}
GRIS_OSCURO, GRIS_CLARO = "#8c8c8c", "#d9d9d9"  # solo grilla y líneas auxiliares
DPI = 300


def tinte(color, fuerza):
    return tuple(1 - fuerza * (1 - c) for c in to_rgb(color))


def limpiar(ax, grilla="y"):
    for lado in ax.spines.values():
        lado.set_visible(False)
    ax.tick_params(length=0, colors=TINTA_SECUNDARIA, labelsize=9)
    if grilla:
        ax.grid(axis=grilla, color=GRILLA, linewidth=0.8)
        ax.set_axisbelow(True)


def guardar(fig, salida):
    fig.savefig(salida, dpi=DPI, facecolor="white")
    plt.close(fig)
    print(f"Gráfico guardado en {salida}")


def figura(ancho=7, alto=4.2):
    fig = plt.figure(figsize=(ancho, alto))
    fig.patch.set_facecolor("white")
    return fig


def miles(x):
    return f"{x:,.0f}"


# 1. Ocupación consolidada -------------------------------------------------------------------------
def ocupacion(wb, salida):
    ws = wb["Ocupacion_consolidada"]
    periodos = [ws.cell(5, c).value for c in range(2, 7)]
    fig = figura()
    ax = fig.add_axes([0.08, 0.1, 0.66, 0.85])
    for r in range(6, 9):
        nombre = ws.cell(r, 1).value
        serie = [ws.cell(r, c).value * 100 for c in range(2, 7)]
        ax.plot(range(5), serie, color=COLOR[nombre], linewidth=2.2, marker="o", markersize=5,
                markeredgecolor="white", markeredgewidth=1, zorder=3)
        ax.text(4.15, serie[-1], f"{nombre} {serie[-1]:.1f}%", va="center", fontsize=9.5,
                fontweight="bold", color=COLOR[nombre])
    ax.axvline(3, color=GRIS_OSCURO, linestyle="--", linewidth=0.9)
    ax.text(3.05, 90.3, "Mallplaza includes\nOpen Plaza from 2Q25", fontsize=7.5,
            color=TINTA_SECUNDARIA, va="bottom")
    ax.set_xticks(range(5), [f"2Q{p[-2:]}" for p in periodos], fontsize=10, color=NEGRO)
    ax.set_ylim(90, 100)
    ax.set_xlim(-0.2, 4.1)
    ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0f}%")
    limpiar(ax)
    guardar(fig, salida)


# 2. Pipeline de GLA ------------------------------------------------------------------------------
def pipeline(wb, salida):
    ws = wb["PIPELINE"]
    datos = [("Mallplaza", ws["D12"].value, ws["E12"].value),
             ("Parque Arauco", ws["N12"].value, ws["O12"].value),
             ("Cenco Malls", ws["I12"].value, 0)]
    fig = figura()
    ax = fig.add_axes([0.11, 0.1, 0.86, 0.85])
    for x, (nombre, org, inorg) in enumerate(datos):
        ax.bar(x, org, width=0.55, color=COLOR[nombre], linewidth=0)
        ax.text(x, org / 2, miles(org), ha="center", va="center", fontsize=9.5,
                color=TEXTO[nombre])
        if inorg:
            ax.bar(x, inorg, bottom=org, width=0.55, color=tinte(COLOR[nombre], 0.35),
                   edgecolor="white", linewidth=0, hatch="///")
            ax.plot([x - 0.275, x + 0.275], [org, org], color="white", linewidth=1.5)
            ax.text(x, org + inorg / 2, miles(inorg), ha="center", va="center", fontsize=9.5,
                    color=NEGRO, bbox={"facecolor": "white", "edgecolor": "none", "pad": 1.5})
        ax.text(x, org + inorg + 3000, f"{miles(org + inorg)} m²", ha="center", va="bottom",
                fontsize=10.5, fontweight="bold", color=NEGRO)
    ax.set_xticks(range(3), [d[0] for d in datos], fontsize=10, color=NEGRO)
    ax.set_ylim(0, 200_000)
    ax.yaxis.set_major_formatter(lambda y, _: miles(y))
    ax.set_ylabel("GLA (m²)", color=TINTA_SECUNDARIA)
    limpiar(ax)
    ax.legend(handles=[Patch(facecolor=TINTA_SECUNDARIA, label="Organic"),
                       Patch(facecolor=tinte(TINTA_SECUNDARIA, 0.35), hatch="///", edgecolor="white",
                             label="Acquisitions")],
              loc="upper right", frameon=False, fontsize=9)
    guardar(fig, salida)


# 3. Oferta de GLA vs consumo privado -----------------------------------------------------------
def oferta_demanda(wb, salida):
    ws = wb["OFERTA VS DEMANDA ESPERADA"]
    filas = [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(2, 6)]
    assert all(isinstance(f[1], float) for f in filas), "abrir y guardar el Excel en Excel (vínculo externo)"
    fig = figura()
    ax = fig.add_axes([0.08, 0.13, 0.9, 0.82])
    xs = [0, 1, 2, 3.4]
    for x, (anio, oferta, consumo) in zip(xs, filas):
        for dx, v, color in ((-0.19, oferta, CELESTE), (0.19, consumo, AMARILLO)):
            ax.bar(x + dx, v * 100, width=0.36, color=color)
            ax.text(x + dx, v * 100 + 0.1, f"{v * 100:.1f}%", ha="center", va="bottom", fontsize=9,
                    color=NEGRO)
    ax.axvline(2.7, color=GRIS_CLARO, linewidth=1)
    ax.set_xticks(xs, [str(f[0]).replace("Acumulado", "Cumulative\n").replace(" 2026", "2026")
                       for f in filas], fontsize=10, color=NEGRO)
    ax.set_ylim(0, 7.8)
    ax.yaxis.set_major_formatter(lambda y, _: f"{y:.0f}%")
    limpiar(ax)
    ax.legend(handles=[Patch(color=CELESTE, label="New GLA supply (trio, Chile)"),
                       Patch(color=AMARILLO, label="Real private consumption (IPoM)")],
              loc="upper left", frameon=False, fontsize=9)
    guardar(fig, salida)


# 4. Vencimiento de contratos -------------------------------------------------------------------
def contratos(wb, salida):
    ws = wb["Duración contratos"]
    datos = {}
    for r in range(2, ws.max_row + 1):
        emp, tramo, pct, dur = (ws.cell(r, c).value for c in range(1, 5))
        datos.setdefault(emp, {"tramos": [], "dur": dur})["tramos"].append((tramo, pct))
    empresas = sorted(datos, key=lambda e: datos[e]["dur"])  # mayor duración arriba
    # Vence antes = más riesgo: rojo → amarillo → celeste (paleta Mallplaza)
    rampa = [ROJO, tinte(ROJO, 0.45), AMARILLO, tinte(CELESTE, 0.55), CELESTE]
    fig = figura(7, 3.2)
    ax = fig.add_axes([0.17, 0.08, 0.66, 0.76])
    for y, emp in enumerate(empresas):
        izq = 0
        for i, (tramo, pct) in enumerate(datos[emp]["tramos"]):
            ax.barh(y, pct * 100, left=izq, height=0.6, color=rampa[i], edgecolor="white",
                    linewidth=1.2)
            if pct >= 0.05:
                ax.text(izq + pct * 50, y, f"{pct * 100:.1f}%", ha="center", va="center",
                        fontsize=8, color="white" if i == 0 else NEGRO)
            izq += pct * 100
        ax.text(103, y, f"{datos[emp]['dur']:.1f} yrs", va="center", fontsize=11,
                fontweight="bold", color=NEGRO)
        ax.text(-2, y, emp, ha="right", va="center", fontsize=10, color=NEGRO)
    ax.text(103, len(empresas) - 0.45, "Avg. remaining\nterm", fontsize=7.5, color=TINTA_SECUNDARIA,
            va="bottom")
    ax.set_xlim(0, 100)
    ax.set_yticks([])
    ax.set_xticks([0, 25, 50, 75, 100], ["0%", "25%", "50%", "75%", "100%"])
    limpiar(ax, grilla=None)
    etiquetas = [t.replace("años", "yrs") for t, _ in datos[empresas[0]]["tramos"]]
    fig.legend(handles=[Patch(color=c, label=t) for c, t in zip(rampa, etiquetas)],
               loc="upper left", bbox_to_anchor=(0.16, 1.0), ncol=5, frameon=False, fontsize=8.5,
               handlelength=1.4, columnspacing=1.2)
    guardar(fig, salida)


# 5. Market share estático por país -------------------------------------------------------------
def market_share_paises(wb, salida):
    ws = wb["MS_estatico_2Q26"]
    otros = tinte(CELESTE, 0.35)
    roles = {"Líder fuera del trío": ("Leader outside the trio", AMARILLO),
             "Otros": ("Others", otros)}
    paises = {}
    for r in range(2, ws.max_row + 1):
        pais, emp, rol = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 4).value
        paises.setdefault(pais, []).append((emp, rol, ws.cell(r, 5).value, ws.cell(r, 7).value))
    for pais, filas in paises.items():  # "Otros" = 1 − resto (fórmula en el Excel)
        resto = 1 - sum(f[2] for f in filas if f[1] != "Otros")
        paises[pais] = [(e, ro, resto if ro == "Otros" else v, n) for e, ro, v, n in filas]

    fig = figura(9, 3.9)
    nombres_pais = {"Chile": "Chile", "Perú": "Peru", "Colombia": "Colombia"}
    for k, (pais, filas) in enumerate(paises.items()):
        ax = fig.add_axes([0.02 + k * 0.33, 0.14, 0.29, 0.68])
        valores = [f[2] for f in filas]
        colores = [COLOR.get(f[0], roles.get(f[1], (None, otros))[1]) for f in filas]
        cuñas, _ = ax.pie(valores, colors=colores, startangle=90, counterclock=False,
                          wedgeprops={"width": 0.38, "edgecolor": "white", "linewidth": 1.5})
        for cuña, (emp, rol, v, nota) in zip(cuñas, filas):
            ang = math.radians((cuña.theta1 + cuña.theta2) / 2)
            marca = "*" if nota and "pro forma" in nota else ""
            fuera = v < 0.05
            radio = 1.22 if fuera else 0.81
            claro = colores[filas.index((emp, rol, v, nota))] in (otros, VERDE_PARQUE_ARAUCO, AMARILLO)
            ax.text(radio * math.cos(ang), radio * math.sin(ang), f"{v:.0%}{marca}",
                    ha="center", va="center", fontsize=9, fontweight="bold",
                    color=NEGRO if fuera or claro else "white")
        ax.text(0, 0, nombres_pais[pais], ha="center", va="center", fontsize=12,
                fontweight="bold", color=NEGRO)
        trio = sum(f[2] for f in filas if f[0] in COLOR)
        ax.text(0, -1.42, f"Trio share: {trio:.0%}", ha="center", va="center", fontsize=10,
                color=NEGRO)
        ax.set_xlim(-1.3, 1.3)
        ax.set_ylim(-1.55, 1.3)
    fig.legend(handles=[Patch(color=COLOR["Mallplaza"], label="Mallplaza"),
                        Patch(color=COLOR["Cenco Malls"], label="Cenco Malls"),
                        Patch(color=COLOR["Parque Arauco"], label="Parque Arauco"),
                        Patch(color=AMARILLO, label="Leader outside the trio"),
                        Patch(color=otros, label="Others")],
               loc="upper center", bbox_to_anchor=(0.5, 0.99), ncol=5, frameon=False, fontsize=9)
    fig.text(0.02, 0.02, "* Pro forma, includes Gran Plaza. Share of GLA among the main operators "
             "of each country.", fontsize=7.5, color=TINTA_SECUNDARIA)
    guardar(fig, salida)


# 6. Cinco fuerzas de Porter --------------------------------------------------------------------
def porter(wb, salida):
    ws = wb["PORTER"]
    notas = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(2, ws.max_row + 1)}
    orden = [("Poder de negociación de los locatarios", "Bargaining power\nof tenants"),
             ("Poder de negociación de los proveedores", "Bargaining power\nof suppliers"),
             ("Amenaza de sustitutos (e-commerce)", "Threat of substitutes\n(e-commerce)"),
             ("Amenaza de nuevos entrantes", "Threat of\nnew entrants"),
             ("Rivalidad entre competidores", "Rivalry among\ncompetitors")]
    valores = [math.ceil(notas[k]) for k, _ in orden]  # redondeo hacia arriba (instrucción del equipo)
    n = len(orden)
    ang = np.linspace(0, 2 * np.pi, n, endpoint=False)
    fig = figura(6, 5.4)
    ax = fig.add_axes([0.2, 0.17, 0.6, 0.66], polar=True)
    ax.set_theta_zero_location("N")
    ax.set_theta_direction(-1)
    for nivel in range(1, 6):  # grilla pentagonal
        ax.plot(np.append(ang, ang[0]), [nivel] * (n + 1), color=GRIS_CLARO, linewidth=0.8)
    for a in ang:
        ax.plot([a, a], [0, 5], color=GRIS_CLARO, linewidth=0.8)
    cerrado = np.append(valores, valores[0])
    ax.fill(np.append(ang, ang[0]), cerrado, color=ROJO, alpha=0.15)
    ax.plot(np.append(ang, ang[0]), cerrado, color=ROJO, linewidth=2.4)
    ax.scatter(ang, valores, s=45, color=ROJO, edgecolor="white", linewidth=1.2, zorder=3)
    for a, v in zip(ang, valores):
        ax.text(a, v + 0.55, f"{v}", ha="center", va="center", fontsize=12, fontweight="bold",
                color=NEGRO)
    for nivel in range(1, 6):
        ax.text(np.radians(36), nivel, str(nivel), fontsize=7.5, color=GRIS_OSCURO,
                ha="center", va="center", bbox={"facecolor": "white", "edgecolor": "none", "pad": 0.5})
    ax.set_xticks(ang, [e for _, e in orden], fontsize=10, color=NEGRO)
    ax.tick_params(axis="x", pad=22)
    ax.set_yticks([])
    ax.set_ylim(0, 5)
    ax.grid(False)
    ax.spines["polar"].set_visible(False)
    fig.text(0.5, 0.02, "Scale 1–5: 1 = weak force (attractive industry), 5 = strong force.\n"
             "Scores rounded up; country scores weighted by Mallplaza GLA (2Q26).",
             ha="center", fontsize=7.5, color=TINTA_SECUNDARIA)
    guardar(fig, salida)


# 7. Market share evolutivo Chile + crecimiento de la industria (base 100) ----------------------
def market_share_chile(wb, salida):
    ws = wb["MS_evolutivo_Chile"]
    periodos = [f"2Q{ws.cell(5, c).value[-2:]}" for c in range(2, 7)]
    empresas = [ws.cell(r, 1).value for r in (6, 7, 8)]
    gla = np.array([[ws.cell(r, c).value for c in range(2, 7)] for r in (6, 7, 8)], dtype=float)
    total = gla.sum(axis=0)
    indice = gla / total[0] * 100  # capas apiladas: el borde superior es el índice de la industria
    cuota = gla / total
    x = np.arange(len(periodos))

    fig = figura()
    ax = fig.add_axes([0.09, 0.1, 0.86, 0.86])
    ax.stackplot(x, indice, colors=[COLOR[e] for e in empresas], edgecolor="white", linewidth=1.5)
    base = np.zeros(len(x))
    for emp, capa, c in zip(empresas, indice, cuota):
        for i in (0, len(x) - 1):
            dx = 0.22 if i == 0 else -0.22
            ax.text(x[i] + dx, base[i] + capa[i] / 2, f"{c[i]:.1%}", ha="center", va="center",
                    fontsize=9.5, fontweight="bold", color=TEXTO[emp])
        ax.text(x[2], base[2] + capa[2] / 2, emp, ha="center", va="center", fontsize=10,
                color=TEXTO[emp])
        base = base + capa
    tope = indice.sum(axis=0)
    ax.plot(x, tope, color=NEGRO, linewidth=1.2, marker="o", markersize=4, zorder=3)
    for i, v in enumerate(tope):
        ha = "left" if i == 0 else "right" if i == len(x) - 1 else "center"
        ax.text(x[i], v + 2.5, f"{v:.1f}", ha=ha, va="bottom", fontsize=9.5, fontweight="bold",
                color=NEGRO)
    ax.text(x[0], 124, "Trio GLA index (2Q22 = 100)", ha="left", va="center", fontsize=9,
            color=TINTA_SECUNDARIA)
    ax.set_xlim(0, len(x) - 1)
    ax.set_ylim(0, 128)
    ax.set_yticks([0, 20, 40, 60, 80, 100])
    ax.set_xticks(x, periodos, fontsize=10, color=NEGRO)
    ax.set_ylabel("GLA index, 2Q22 = 100", color=TINTA_SECUNDARIA)
    limpiar(ax)
    guardar(fig, salida)


# 8. Eficiencia operativa: EBITDA por m² -------------------------------------------------------
def eficiencia(wb, salida):
    ws = wb["EF. OPERATIVA (CLPM2)"]
    filas = [(ws.cell(r, 1).value, ws.cell(r, 2).value * 1e6 / ws.cell(r, 3).value * 4)
             for r in range(6, 9)]
    filas.sort(key=lambda f: f[1])
    fig = figura(7, 2.8)
    ax = fig.add_axes([0.17, 0.16, 0.76, 0.8])
    for y, (emp, v) in enumerate(filas):
        ax.barh(y, v / 1000, height=0.58, color=COLOR[emp])
        ax.text(v / 1000 + 3, y, f"{v / 1000:,.0f}k", va="center", fontsize=10.5,
                fontweight="bold", color=NEGRO)
    ax.set_yticks(range(len(filas)), [f[0] for f in filas], fontsize=10, color=NEGRO)
    ax.set_xlim(0, 270)
    ax.xaxis.set_major_formatter(lambda v, _: f"{v:,.0f}k")
    ax.set_xlabel("EBITDA per m² of GLA, 2Q26 annualized (CLP)", color=TINTA_SECUNDARIA)
    limpiar(ax, grilla="x")
    guardar(fig, salida)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    salida = Path(sys.argv[2] if len(sys.argv) > 2 else Path(__file__).parent / "industria")
    salida.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans"})
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ocupacion(wb, salida / "ind1_ocupacion.png")
    pipeline(wb, salida / "ind2_pipeline_gla.png")
    oferta_demanda(wb, salida / "ind3_oferta_vs_consumo.png")
    contratos(wb, salida / "ind4_vencimiento_contratos.png")
    market_share_paises(wb, salida / "ind5_market_share_paises.png")
    porter(wb, salida / "ind6_porter.png")
    market_share_chile(wb, salida / "ind7_market_share_chile.png")
    eficiencia(wb, salida / "ind8_ebitda_m2.png")
