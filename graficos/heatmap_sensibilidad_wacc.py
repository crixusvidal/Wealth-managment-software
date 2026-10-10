"""Mapa de calor de la sensibilidad del valor por acción de Mallplaza (WACC vs g).

Muestra el precio objetivo (80% DCF + 20% múltiplos, hoja Precio_Objetivo) en una grilla
5 x 5 con el caso base al centro. El DCF replica la fórmula de la grilla de la hoja DCF
(celda C61: valor terminal normalizado con reinversión g real / RONIC) y se verifica contra
la grilla original del Excel (B61:F65).
Rojo = valor bajo el caso base, celeste = sobre el caso base.

Uso:
    python graficos/heatmap_sensibilidad_wacc.py <Valoracion_Consolidada.xlsx> [salida.png]
"""
import sys

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
from matplotlib.patches import Rectangle

from sensibilidad_wacc import AMARILLO, CELESTE, NEGRO, ROJO, TINTA_SECUNDARIA

NEUTRO = "#f2f2f2"


def leer_insumos(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    d, s, ex = wb["DCF"], wb["Supuestos"], wb["Expansion"]
    v = lambda c: s[c].value
    return {"fcff": np.array([d[f"{c}22"].value for c in "CDEFGH"]),
            "periodos": np.array([d[f"{c}24"].value for c in "CDEFGH"]),
            "ing_norm": d["I29"].value, "ebit_norm": d["I31"].value,
            "tax": v("B48"), "da": v("B46"), "mant": v("B47"), "k_ct": v("B73") * v("B74"),
            "inflacion": v("B35"), "ronic": ex["B44"].value,
            "ajustes": v("B11") + v("B12"), "acciones": v("B10")}


def leer_grilla(ruta):
    d = openpyxl.load_workbook(ruta, data_only=True)["DCF"]
    gs = [d.cell(60, c).value for c in range(3, 7)]
    waccs = [d.cell(r, 2).value for r in range(61, 66)]
    z = np.array([[d.cell(r, c).value for c in range(3, 7)] for r in range(61, 66)])
    return waccs, gs, z, d["B34"].value, d["B35"].value, d["B46"].value


def valor_por_accion(wacc, g, x):
    vp = np.sum(x["fcff"] / (1 + wacc) ** x["periodos"])
    nopat = x["ebit_norm"] * (1 + g) * (1 - x["tax"])
    g_real = (1 + g) / (1 + x["inflacion"]) - 1
    fcff_t = (nopat + (x["da"] - x["mant"]) * x["ing_norm"] * (1 + g)
              - x["k_ct"] * x["ing_norm"] * g - g_real / x["ronic"] * nopat)
    vt = fcff_t / (wacc - g) / (1 + wacc) ** 5.5
    return (vp + vt - x["ajustes"]) / x["acciones"]


PASOS_G = [-0.01, -0.005, 0.0, 0.005, 0.01]  # g alrededor del caso base: 5 x 5 con la base al centro


def grilla(ruta):
    """Precio objetivo (80% DCF + 20% múltiplos) por WACC (filas, de mayor a menor) y g.

    El DCF de cada celda se recalcula con la fórmula de la grilla del Excel y se verifica contra
    ella en las g que el Excel trae; el valor por múltiplos queda fijo (no depende de WACC ni g).
    """
    waccs, gs_xl, z_xl, wacc_base, g_base, precio = leer_grilla(ruta)
    x = leer_insumos(ruta)
    for i, w in enumerate(waccs):
        for j, g in enumerate(gs_xl):
            assert abs(valor_por_accion(w, g, x) - z_xl[i, j]) < 1e-6, "la grilla no coincide con el Excel"
    po = openpyxl.load_workbook(ruta, data_only=True)["Precio_Objetivo"]
    peso, multiplos = po["B14"].value, po["B13"].value
    gs = [round(g_base + p, 6) for p in PASOS_G]
    waccs = sorted(waccs, reverse=True)
    z = np.array([[peso * valor_por_accion(w, g, x) + (1 - peso) * multiplos for g in gs]
                  for w in waccs])
    i_base = int(np.argmin([abs(w - wacc_base) for w in waccs]))
    j_base = gs.index(round(g_base, 6))
    assert abs(z[i_base, j_base] - po["B15"].value) < 1e-6, "el centro debe ser el precio objetivo"
    return waccs, gs, z, i_base, j_base, precio


def pct(x, decimales=1):
    return f"{x * 100:.{decimales}f}%"


def miles(x):
    return f"{x:,.0f}"


def color_texto(rgba):
    r, g, b = rgba[:3]
    return "white" if 0.299 * r + 0.587 * g + 0.114 * b < 0.55 else NEGRO


def graficar(ruta_xlsx, salida):
    waccs, GS, z, i_base, j_base, precio = grilla(ruta_xlsx)
    wacc_base, g_base = waccs[i_base], GS[j_base]
    base = z[i_base, j_base]
    d = openpyxl.load_workbook(ruta_xlsx, data_only=True)["DCF"]
    ronic, wacc_real = d["B55"].value, d["B54"].value
    multiplos = openpyxl.load_workbook(ruta_xlsx, data_only=True)["Precio_Objetivo"]["B13"].value

    cmap = LinearSegmentedColormap.from_list("mallplaza", [ROJO, NEUTRO, CELESTE])
    norm = TwoSlopeNorm(vmin=z.min(), vcenter=base, vmax=z.max())

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig = plt.figure(figsize=(12, 7.5), dpi=200)
    fig.patch.set_facecolor("white")

    fig.add_artist(Rectangle((0.05, 0.905), 0.008, 0.055, color=ROJO,
                             transform=fig.transFigure))
    fig.text(0.065, 0.915, "TARGET PRICE SENSITIVITY", fontsize=22,
             fontweight="bold", color=NEGRO)
    fig.text(0.05, 0.87, "Mallplaza · target price (CLP) = 80% DCF + 20% EV/EBITDA multiples, by nominal WACC and terminal growth",
             fontsize=11, color=TINTA_SECUNDARIA)

    ax = fig.add_axes([0.1, 0.24, 0.87, 0.58])
    im = ax.imshow(z, cmap=cmap, norm=norm, aspect="auto")
    for i in range(z.shape[0]):
        for j in range(z.shape[1]):
            ax.text(j, i, miles(z[i, j]), ha="center", va="center", fontsize=14,
                    fontweight="bold" if (i, j) == (i_base, j_base) else "normal",
                    color=color_texto(cmap(norm(z[i, j]))))
    for i in range(z.shape[0] + 1):
        ax.axhline(i - 0.5, color="white", linewidth=3)
    for j in range(z.shape[1] + 1):
        ax.axvline(j - 0.5, color="white", linewidth=3)
    ax.add_patch(Rectangle((j_base - 0.46, i_base - 0.44), 0.92, 0.88, fill=False,
                           edgecolor=AMARILLO, linewidth=3, zorder=3))

    ax.set_xticks(range(len(GS)), [pct(g, 2) for g in GS])
    ax.set_yticks(range(len(waccs)), [pct(w, 2) for w in waccs])
    ax.set_xlabel("Terminal growth rate (g, nominal)", color=TINTA_SECUNDARIA,
                  labelpad=10)
    ax.set_ylabel("WACC (nominal, CLP)", color=TINTA_SECUNDARIA, labelpad=10)
    ax.tick_params(length=0, colors=NEGRO, labelsize=12)
    for lado in ax.spines.values():
        lado.set_visible(False)

    cax = fig.add_axes([0.1, 0.09, 0.87, 0.02])
    cb = fig.colorbar(im, cax=cax, orientation="horizontal")
    cb.set_ticks([z.min(), base, z.max()])
    cb.set_ticklabels([miles(z.min()),
                       f"{miles(base)}\ntarget price (WACC {pct(wacc_base, 2)} · g {pct(g_base)})",
                       miles(z.max())])
    cb.outline.set_visible(False)
    cax.tick_params(length=0, colors=TINTA_SECUNDARIA, labelsize=9)
    cax.set_title("Target price (CLP)", fontsize=9, color=TINTA_SECUNDARIA,
                  loc="left")

    fig.text(0.05, 0.015,
             "Multiples value held at CLP " + miles(multiplos) + ". Higher g lowers value: terminal growth needs reinvestment at a " + pct(ronic) + " real RONIC, below the "
             + pct(wacc_real) + " real WACC. Share price CLP " + miles(precio) + ". Source: Team analysis.",
             fontsize=8, color=TINTA_SECUNDARIA)

    fig.savefig(salida, facecolor="white")
    print(f"Gráfico guardado en {salida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    graficar(sys.argv[1],
             sys.argv[2] if len(sys.argv) > 2 else "heatmap_sensibilidad_wacc.png")
