"""Gráfico de sensibilidad del valor por acción de Mallplaza a la WACC y a g.

Lee la grilla "Sensibilidad: valor por acción (CLP) — WACC (filas) vs g (columnas)"
de la hoja DCF del modelo consolidado y guarda un PNG.

Uso:
    python graficos/sensibilidad_wacc.py <Valoracion_Consolidada.xlsx> [salida.png]
"""
import sys

import matplotlib.pyplot as plt
import openpyxl
from matplotlib.ticker import FuncFormatter

# Colores corporativos Mallplaza
CELESTE = "#64b0e1"
AMARILLO = "#ecb100"
ROJO = "#e9004b"
NEGRO = "#000000"
TINTA_SECUNDARIA = "#5f5f5f"
GRILLA = "#e6e6e6"

def leer_grilla(ruta):
    """Grilla WACC (filas) vs g (columnas) de la hoja DCF del modelo consolidado (B60:F65)."""
    ws = openpyxl.load_workbook(ruta, data_only=True)["DCF"]
    gs = [ws.cell(60, c).value for c in range(3, 7)]
    waccs = [ws.cell(r, 2).value for r in range(61, 66)]
    valores = {g: [ws.cell(r, 3 + i).value for r in range(61, 66)] for i, g in enumerate(gs)}
    return waccs, valores, ws["B34"].value, ws["B35"].value, ws["B46"].value


def pct(x, decimales=1):
    return f"{x * 100:.{decimales}f}%".replace(".", ",")


def miles(x):
    return f"{x:,.0f}".replace(",", ".")


def graficar(ruta_xlsx, salida):
    waccs, valores, wacc_base, g_base, precio = leer_grilla(ruta_xlsx)

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=200)
    fig.patch.set_facecolor("white")

    estilos = [(CELESTE, "o"), (ROJO, "s"), (AMARILLO, "^"), (TINTA_SECUNDARIA, "D")]
    x = [w * 100 for w in waccs]
    for (g, serie), (color, marcador) in zip(valores.items(), estilos):
        es_base = abs(g - g_base) < 1e-9
        ax.plot(
            x, serie, color=color, linewidth=2.5 if es_base else 2,
            marker=marcador, markersize=8, markeredgecolor=NEGRO,
            markeredgewidth=0.8, zorder=3,
            label=f"g = {pct(g)}" + (" (caso base)" if es_base else ""),
        )
        ax.annotate(
            f"g = {pct(g)}", (x[-1], serie[-1]), xytext=(10, 0),
            textcoords="offset points", va="center", fontsize=9, color=NEGRO,
        )

    # Caso base
    base = valores[g_base][[round(w, 6) for w in waccs].index(round(wacc_base, 6))]
    ax.scatter([wacc_base * 100], [base], s=200, facecolor="none",
               edgecolor=NEGRO, linewidth=1.5, zorder=4)
    ax.annotate(
        f"Caso base\nWACC {pct(wacc_base, 2)} · g {pct(g_base)}\nCLP {miles(base)}",
        (wacc_base * 100, base), xytext=(40, 55), textcoords="offset points",
        fontsize=9, color=NEGRO,
        arrowprops=dict(arrowstyle="-", color=NEGRO, linewidth=0.8),
    )

    # Precio de referencia
    ax.axhline(precio, color=NEGRO, linewidth=1, linestyle=(0, (4, 3)), zorder=2)
    ax.text(x[0], precio, f"Precio de mercado CLP {miles(precio)}",
            va="bottom", ha="left", fontsize=9, color=NEGRO)

    ax.set_xticks(x)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: pct(v / 100, 2)))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: miles(v)))
    ax.set_ylim(min(min(v) for v in valores.values()) * 0.92, precio * 1.06)
    ax.set_xlim(x[0] - 0.15, x[-1] + 0.35)
    ax.set_xlabel("WACC nominal CLP", color=TINTA_SECUNDARIA)
    ax.set_ylabel("Valor por acción (CLP)", color=TINTA_SECUNDARIA)

    ax.grid(axis="y", color=GRILLA, linewidth=0.8)
    ax.set_axisbelow(True)
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    for lado in ("left", "bottom"):
        ax.spines[lado].set_color("#bdbdbd")
    ax.tick_params(colors=TINTA_SECUNDARIA)

    fig.text(0.07, 0.955, "Mallplaza — Sensibilidad del valor por acción a la WACC",
             fontsize=14, fontweight="bold", color=NEGRO)
    fig.text(0.07, 0.915,
             "DCF (FCFF), valor por acción en CLP según WACC y crecimiento perpetuo g",
             fontsize=10, color=TINTA_SECUNDARIA)
    ax.legend(loc="upper right", bbox_to_anchor=(1, 0.86), frameon=False, fontsize=9)
    fig.text(0.07, 0.01, "Fuente: modelo de valoración consolidado del equipo (hoja DCF). Precio al 30-09-2026.",
             fontsize=8, color=TINTA_SECUNDARIA)

    fig.tight_layout(rect=(0.02, 0.03, 1, 0.9))
    fig.savefig(salida, facecolor="white")
    print(f"Gráfico guardado en {salida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    graficar(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "sensibilidad_wacc.png")
