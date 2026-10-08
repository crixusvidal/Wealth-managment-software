"""Agrega al libro de comparables (Bloomberg) de Mallplaza las hojas Valor_Implicito y DCF_WACC_Nominal.

DCF_WACC_Nominal reconstruye el FCFF del modelo del equipo (Mallplaza_FA_Exhibits) y lo descuenta
con una WACC nominal consistente (Ke nominal, Kd nominal después de impuestos, pesos de mercado).
Valor_Implicito aplica los múltiplos de los peers (P25 / mediana / P75) a las métricas de Mallplaza
y combina el resultado con ese DCF (80/20). Las hojas originales no se modifican.

Uso:
    python valoracion/valoracion_multiplos.py <Comparables_Mallplaza_CFA_BBG.xlsx> <Mallplaza_FA_Exhibits.xlsx> <salida.xlsx>
"""
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NAVY, SALMON, MEDIAN_FILL, YELLOW = "1F3864", "FCE4D6", "D9E1F2", "FFFF00"
BLUE, GREEN, BLACK = "0000FF", "008000", "000000"
NUM = r"#,##0;\(#,##0\);\-"
MULT = r'0.0\x;\(0.0"x)";\-'
PCT = r"0.0%;\(0.0%\);\-"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

FA = "Mallplaza_FA_Exhibits_v6.xlsx"

# (método, base, grupo, hoja, columna del múltiplo, filas de peers, celda de la métrica, tipo)
METODOS = [
    ("EV/EBITDA", "LTM", "Core", "Valoracion_LTM", "O", "8:12", "C5", "EV"),
    ("EV/EBITDA", "NTM", "Core", "Valoracion_NTM", "O", "8:12", "C6", "EV"),
    ("EV/EBIT", "LTM", "Core", "Valoracion_LTM", "P", "8:12", "C7", "EV"),
    ("EV/EBITDA", "LTM", "Chile (Parque Arauco, Cencosud Shopping)", "Valoracion_LTM", "O", "8:9", "C5", "EV"),
    ("EV/EBITDA", "NTM", "Chile (Parque Arauco, Cencosud Shopping)", "Valoracion_NTM", "O", "8:9", "C6", "EV"),
    ("P/E", "NTM", "Core", "Valoracion_NTM", "N", "8:12", "C8", "Equity"),
    ("P/BV", "30-jun-2026", "Core", "Valoracion_LTM", "R", "8:12", "C9", "Equity"),
    ("P/BV", "30-jun-2026", "Chile (Parque Arauco, Cencosud Shopping)", "Valoracion_LTM", "R", "8:9", "C9", "Equity"),
]


def estilo(c, bold=False, color=BLACK, fill=None, fmt=None, size=10, wrap=False, border=True):
    c.font = Font(name="Arial", size=size, bold=bold, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BOX
    c.alignment = Alignment(vertical="center", wrap_text=wrap)


def encabezado(ws, fila, textos):
    for col, t in enumerate(textos, start=1):
        c = ws.cell(fila, col, t)
        estilo(c, bold=True, color="FFFFFF", fill=NAVY, wrap=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def entrada(ws, celda, valor, fmt, fuente):
    ws[celda] = valor
    estilo(ws[celda], color=BLUE, fmt=fmt)
    ws[celda].comment = Comment(fuente, "Team")


def construir_dcf(wb, ruta_fa):
    fa = openpyxl.load_workbook(ruta_fa, data_only=True)
    d = fa["DCF_Valuation"]
    ws = wb.create_sheet("DCF_WACC_Nominal", index=0)
    ws.sheet_view.showGridLines = False
    estilo(ws["A1"], bold=True, size=13, border=False)
    ws["A1"] = "DCF de Mallplaza con WACC nominal consistente (fecha de valorización 30-jun-2026)"
    estilo(ws["A2"], size=8, border=False)
    ws["A2"] = ("FCFF en CLP nominales del modelo del equipo (hoja DCF_Valuation) descontado con una WACC "
                "nominal: Ke nominal, Kd nominal después de impuestos y pesos a valor de mercado. MM CLP.")

    estilo(ws["A4"], bold=True, size=11, border=False)
    ws["A4"] = "1. WACC nominal"
    wacc = [
        ("Ke nominal CLP", "C5", 0.12, PCT, "Supuesto del equipo (12%), tratado como nominal en CLP. Si fuera real: (1+Ke)×(1+inflación)−1."),
        ("Kd real en UF (última emisión)", "C6", 0.036, PCT, "Modelo de valuación del equipo (FFCF), según la auditoría de la hoja DCF_Valuation del modelo."),
        ("Inflación de largo plazo Chile", "C7", 0.03, PCT, "Meta de inflación del BCCh (3%)."),
        ("Kd nominal CLP", "C8", "=(1+C6)*(1+C7)-1", PCT, None),
        ("Tasa de impuesto", "C9", 0.27, PCT, "Tasa legal de primera categoría, Chile."),
        ("Kd después de impuestos", "C10", "=C8*(1-C9)", PCT, None),
        ("Market cap (MM CLP)", "C11", "=Valoracion_LTM!E7", NUM, None),
        ("Caja 30-jun-2026 (MM CLP)", "C12", 347621, NUM, f"{FA}, hoja Forecast, celda D117 (caja jun-26, release 2T26)."),
        ("Deuda financiera bruta (MM CLP)", "C13", "=Valoracion_LTM!F7+C12", NUM, None),
        ("E/V", "C14", "=C11/(C11+C13)", PCT, None),
        ("D/V", "C15", "=1-C14", PCT, None),
        ("WACC nominal", "C16", "=C14*C5+C15*C10", PCT, None),
        ("Crecimiento terminal nominal (g)", "C17", 0.03, PCT, f"{FA}, DCF_Valuation!J37: meta de inflación, crecimiento real 0%."),
    ]
    for label, celda, val, fmt, fuente in wacc:
        r = int(celda[1:])
        estilo(ws.cell(r, 1, label))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        if fuente:
            entrada(ws, celda, val, fmt, fuente)
        else:
            ws[celda] = val
            estilo(ws[celda], fmt=fmt, color=GREEN if str(val).startswith("=Valoracion") else BLACK,
                   bold=celda == "C16", fill=SALMON if celda == "C16" else None)
    ws["C5"].fill = PatternFill("solid", fgColor=YELLOW)

    estilo(ws["A19"], bold=True, size=11, border=False)
    ws["A19"] = "2. Flujo de caja libre de la firma (FCFF)"
    cols = "JKLMNO"
    encabezado(ws, 20, ["Concepto", ""] + [d[c + "5"].value for c in cols])
    ws.merge_cells("A20:B20")
    filas = [
        (21, "EBITDA ajustado", "15", "EBITDA ajustado (definición compañía)"),
        (22, "(−) D&A", "11", "D&A"),
        (24, "Tasa de impuesto operacional", "22", "Tasa de impuesto operacional"),
        (28, "(−) Capex", "26", "Capex"),
        (29, "(−) Variación de capital de trabajo", "27", "Variación de capital de trabajo"),
        (31, "Período de descuento (años, mitad de período)", "32", "Período de descuento"),
    ]
    for r, label, fila_fa, nombre in filas:
        estilo(ws.cell(r, 1, label))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        for j, c in enumerate(cols):
            entrada(ws, ws.cell(r, 3 + j).coordinate, d[f"{c}{fila_fa}"].value,
                    PCT if r == 24 else "0.00" if r == 31 else NUM,
                    f"{FA}, DCF_Valuation!{c}{fila_fa} ({nombre}).")
    calc = [
        (23, "EBIT", "={c}21-{c}22"),
        (25, "(−) Impuestos operacionales", "=-{c}23*{c}24"),
        (26, "NOPAT", "={c}23+{c}25"),
        (27, "(+) D&A", "={c}22"),
        (30, "FCFF", "={c}26+{c}27+{c}28-{c}29"),
        (32, "Factor de descuento", "=1/(1+$C$16)^{c}31"),
        (33, "Valor presente del FCFF", "={c}30*{c}32"),
    ]
    for r, label, f in calc:
        estilo(ws.cell(r, 1, label), bold=r == 30)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        for j in range(6):
            c = openpyxl.utils.get_column_letter(3 + j)
            estilo(ws.cell(r, 3 + j, f.format(c=c)), fmt="0.0000" if r == 32 else NUM, bold=r == 30,
                   fill=MEDIAN_FILL if r == 30 else None)

    estilo(ws["A35"], bold=True, size=11, border=False)
    ws["A35"] = "3. Valor de la empresa y valor por acción"
    valor = [
        (36, "Suma VP del FCFF explícito", "=SUM(C33:H33)", NUM, None),
        (37, "FCFF 2032E = FCFF 2031E × (1 + g)", "=H30*(1+C17)", NUM, None),
        (38, "Valor terminal al 31-12-2031 (Gordon)", "=C37/(C16-C17)", NUM, None),
        (39, "VP del valor terminal (5,5 años)", "=C38/(1+C16)^5.5", NUM, None),
        (40, "Enterprise value", "=C36+C39", NUM, None),
        (41, "(−) Deuda financiera neta 30-jun-2026", 1285740, NUM, f"{FA}, DCF_Valuation!J45 (DFN definición compañía, release 2T26)."),
        (42, "(−) Minoritarios (libro) 30-jun-2026", 50841, NUM, f"{FA}, DCF_Valuation!J46 (release 2T26)."),
        (43, "Equity value", "=C40-C41-C42", NUM, None),
        (44, "Acciones (MM)", "=Datos_BBG!O8/Datos_BBG!P8", "#,##0.0", None),
        (45, "Valor por acción DCF (CLP)", "=C43/C44", NUM, None),
        (46, "Precio al 30-sep-2026 (CLP)", "=Datos_BBG!P8", NUM, None),
        (47, "Upside / (downside) vs precio", "=C45/C46-1", PCT, None),
        (48, "EV / EBITDA LTM implícito", "=C40/Valoracion_LTM!J7", MULT, None),
        (49, "Cap rate implícito (EBITDA LTM / EV)", "=Valoracion_LTM!J7/C40", PCT, None),
        (50, "VP del valor terminal / EV", "=C39/C40", PCT, None),
    ]
    for r, label, val, fmt, fuente in valor:
        estilo(ws.cell(r, 1, label), bold=r == 45)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        if fuente:
            entrada(ws, f"C{r}", val, fmt, fuente)
        else:
            ws[f"C{r}"] = val
            estilo(ws[f"C{r}"], fmt=fmt, bold=r == 45, fill=SALMON if r == 45 else None,
                   color=GREEN if "Datos_BBG!P8" == str(val)[1:] else BLACK)

    estilo(ws["A52"], bold=True, size=11, border=False)
    ws["A52"] = "4. Control: WACC original del modelo"
    estilo(ws.cell(53, 1, "WACC original (FFCF!B120)"))
    ws.merge_cells("A53:B53")
    entrada(ws, "C53", 0.1003, PCT, f"{FA}, DCF_Valuation!J31 (WACC del modelo de valuación, 10,03%).")
    estilo(ws.cell(54, 1, "Valor por acción con WACC original (CLP)"))
    ws.merge_cells("A54:B54")
    ws["C54"] = ("=(SUMPRODUCT(C30:H30,1/(1+C53)^C31:H31)+H30*(1+C17)/(C53-C17)/(1+C53)^5.5"
                 "-C41-C42)/C44")
    estilo(ws["C54"], fmt=NUM)
    estilo(ws.cell(54, 4, f"Debe dar {d['J49'].value:,.1f} (DCF_Valuation!J49): confirma que el FCFF se reconstruyó igual."),
           size=8, border=False)

    for k, w in {"A": 30, "B": 14, "C": 14, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13}.items():
        ws.column_dimensions[k].width = w
    return d["J49"].value


def construir(ruta_in, ruta_fa, ruta_out):
    wb = openpyxl.load_workbook(ruta_in)
    control = construir_dcf(wb, ruta_fa)
    ws = wb.create_sheet("Valor_Implicito", index=0)
    ws.sheet_view.showGridLines = False

    estilo(ws["A1"], bold=True, size=13, border=False)
    ws["A1"] = "Valor implícito de Mallplaza por múltiplos comparables (precio al 30-sep-2026)"
    estilo(ws["A2"], size=8, border=False)
    ws["A2"] = ("Aplica los múltiplos de los peers (hojas Valoracion_LTM y Valoracion_NTM) a las métricas de "
                "Mallplaza. Montos en MM CLP; valor por acción en CLP. Mallplaza excluido de los estadísticos.")

    # 1. Datos de Mallplaza
    estilo(ws["A4"], bold=True, size=11, border=False)
    ws["A4"] = "1. Datos de Mallplaza"
    datos = [
        ("EBITDA LTM", "=Valoracion_LTM!J7", NUM, GREEN, "Bloomberg, LTM 3T25–2T26"),
        ("EBITDA NTM (consenso)", "=Valoracion_NTM!J7", NUM, GREEN, "Bloomberg BEst, 13 analistas"),
        ("EBIT LTM", "=Valoracion_LTM!K7", NUM, GREEN, "EBITDA − D&A"),
        ("Utilidad neta NTM (consenso)", "=Valoracion_NTM!L7", NUM, GREEN, "Bloomberg BEst; puede incluir revalorización"),
        ("Patrimonio controlador (book value)", "=Valoracion_LTM!M7", NUM, GREEN, "30-jun-2026"),
        ("Deuda neta", "=Valoracion_LTM!F7", NUM, GREEN, "Bloomberg NET_DEBT, 30-jun-2026 (misma definición que los peers)"),
        ("Minoritarios", "=Valoracion_LTM!G7", NUM, GREEN, "30-jun-2026"),
        ("Acciones (MM)", "=Datos_BBG!O8/Datos_BBG!P8", "#,##0.0", BLACK, "Market cap / precio"),
        ("Precio al 30-sep-2026 (CLP)", "=Datos_BBG!P8", NUM, GREEN, "Bloomberg PX_LAST"),
        ("Valor por acción DCF (CLP)", "=DCF_WACC_Nominal!C45", NUM, GREEN, "Hoja DCF_WACC_Nominal (WACC nominal)"),
        ("Peso DCF en el precio objetivo", 0.8, PCT, BLUE, "Supuesto del equipo"),
        ("Peso múltiplos en el precio objetivo", "=1-C15", PCT, BLACK, ""),
    ]
    for i, (label, val, fmt, color, nota) in enumerate(datos):
        r = 5 + i
        a = ws.cell(r, 1, label)
        estilo(a)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 3, val)
        estilo(c, color=color, fmt=fmt, fill=YELLOW if r == 15 else None)
        n = ws.cell(r, 4, nota)
        estilo(n, size=8, border=False)
    ws["C15"].comment = Comment("Ponderación definida por el equipo: 80% DCF / 20% múltiplos.", "Team")

    # 2. Valor implícito por acción
    estilo(ws["A18"], bold=True, size=11, border=False)
    ws["A18"] = "2. Valor implícito por acción"
    encabezado(ws, 19, ["Múltiplo", "Base", "Peers", "Múltiplo P25", "Múltiplo mediana", "Múltiplo P75",
                        "Métrica Mallplaza (MM CLP)", "Valor/acción P25 (CLP)", "Valor/acción mediana (CLP)",
                        "Valor/acción P75 (CLP)", "Mediana vs precio"])
    for i, (mult, base, grupo, hoja, col, filas, metrica, tipo) in enumerate(METODOS):
        r = 20 + i
        a, b = filas.split(":")
        rango = f"{hoja}!{col}{a}:{col}{b}"
        valores = [mult, base, grupo,
                   f"=QUARTILE({rango},1)", f"=MEDIAN({rango})", f"=QUARTILE({rango},3)", f"={metrica}"]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(r, j, v)
            estilo(c, fmt=MULT if j in (4, 5, 6) else NUM if j == 7 else None,
                   color=GREEN if j in (4, 5, 6) else BLACK, wrap=j == 3,
                   fill=MEDIAN_FILL if j == 5 else None)
        for j, m in zip((8, 9, 10), "DEF"):
            if tipo == "EV":
                f = f"=({m}{r}*$G{r}-$C$10-$C$11)/$C$12"
            else:
                f = f"={m}{r}*$G{r}/$C$12"
            estilo(ws.cell(r, j, f), fmt=NUM, bold=j == 9, fill=MEDIAN_FILL if j == 9 else None)
        estilo(ws.cell(r, 11, f"=I{r}/$C$13-1"), fmt=PCT)

    # 3. Precio objetivo combinado
    fila = 20 + len(METODOS) + 1
    estilo(ws.cell(fila, 1), bold=True, size=11, border=False)
    ws.cell(fila, 1, "3. Precio objetivo combinado (80% DCF / 20% múltiplos)")
    combinado = [
        ("Valor por múltiplos: EV/EBITDA LTM, mediana Core (CLP)", "=I20", SALMON),
        ("Valor DCF (CLP)", "=C14", None),
        ("Precio objetivo ponderado (CLP)", f"=C15*C{fila + 2}+C16*C{fila + 1}", SALMON),
        ("Precio al 30-sep-2026 (CLP)", "=C13", None),
        ("Upside / (downside)", f"=C{fila + 3}/C{fila + 4}-1", SALMON),
    ]
    for i, (label, f, fill) in enumerate(combinado):
        r = fila + 1 + i
        estilo(ws.cell(r, 1, label), bold=fill is not None, fill=fill)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        estilo(ws.cell(r, 3, f), bold=fill is not None, fmt=PCT if "Upside" in label else NUM, fill=fill)
    ws.cell(fila + 1, 4, "Método principal: EV/EBITDA (neutral a la contabilidad de valor razonable vs costo). "
                         "Cambiar la referencia de esta celda para usar otro método de la tabla 2.")
    estilo(ws.cell(fila + 1, 4), size=8, border=False)

    notas_ini = fila + 8
    estilo(ws.cell(notas_ini, 1), bold=True, size=11, border=False)
    ws.cell(notas_ini, 1, "Notas")
    notas = [
        "EV implícito = múltiplo × métrica de Mallplaza; patrimonio = EV − deuda neta − minoritarios; valor por acción = patrimonio / acciones.",
        "P/E y P/BV se aplican directo sobre patrimonio. P/BV mezcla peers a valor razonable (Chile) con peers a costo (Brasil): usar solo como referencia.",
        "P/E NTM: la utilidad de consenso de Mallplaza puede incluir revalorización de propiedades, lo que infla el valor implícito; usar solo como referencia.",
        "Peers Core = Parque Arauco, Cencosud Shopping, Multiplan, Allos, Iguatemi. IRSA y Fibra Danhos (Sensibilidad) quedan fuera de los estadísticos.",
        "Los peers brasileños transan a múltiplos menores por tasas de interés y riesgo país más altos; por eso se muestra también el grupo Chile.",
        "P25 y P75 con QUARTILE sobre los peers del grupo (con dos peers, entre el mínimo y el máximo).",
    ]
    for i, t in enumerate(notas):
        c = ws.cell(notas_ini + 1 + i, 1, f"• {t}")
        estilo(c, size=9, border=False)

    widths = {"A": 16, "B": 22, "C": 24, "D": 12, "E": 12, "F": 12, "G": 15, "H": 14, "I": 15, "J": 14, "K": 11}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    for r in range(20, 20 + len(METODOS)):
        ws.row_dimensions[r].height = 26
    ws.row_dimensions[19].height = 40
    wb.save(ruta_out)
    print(f"Guardado {ruta_out} (control DCF original: {control:,.1f})")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    construir(sys.argv[1], sys.argv[2], sys.argv[3])
