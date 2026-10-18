"""Regenera el contexto del LLM del proyecto: las cifras clave del informe, leídas de los modelos VIGENTES.

Qué hace:
  1. Elige el archivo más reciente de cada familia de modelos (versión más alta; ver FAMILIAS).
  2. BORRA lo generado antes en esta carpeta (cifras_clave.*, archivos_vigentes.md) y escribe lo nuevo: nunca se acumulan
     versiones viejas del contexto.
  3. Escribe cifras_clave.json (máquina), cifras_clave.md (lo que lee el LLM) y archivos_vigentes.md (qué archivos son
     los vigentes y cuáles quedaron viejos).

Cada cifra lleva su fuente (archivo!hoja!celda), así el LLM puede citarla y el equipo verificarla. Nada se escribe a mano.

Uso (desde la raíz del repo):   python contexto_llm/generar.py
Lo ejecuta el agente `contexto-llm` (.claude/agents/contexto-llm.md) y generar_todo.sh.
"""
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
SEC = RAIZ / "secciones"
SALIDA = Path(__file__).parent
GENERADOS = ("cifras_clave.json", "cifras_clave.md", "archivos_vigentes.md")

# Familias de modelos: patrón (la versión más alta gana; a igual versión, el más reciente por fecha) -------
FAMILIAS = {
    "valoracion": ("valuacion/modelos", r"Mallplaza_Valoracion_Consolidada_v(\d+)\.xlsx"),
    "fa_exhibits": ("analisis_financiero/modelos", r"Mallplaza_FA_Exhibits_v(\d+)\.xlsx"),
    "figuras_valuacion": ("valuacion/modelos", r"Mallplaza_Figuras_Valuacion_recalc\.xlsx"),
    "industria": ("industria/modelos", r"Excel_DATA_INDUSTRY\.xlsx"),
    "eficiencia": ("industria/modelos", r"Eficiencia_EBITDA_GLA_2T26V(\d+)\.xlsx"),
    "negocio": ("negocio/modelos", r"Excel_bussines_para_armar_graficos\.xlsx"),
    "tornado": ("riesgos/modelos", r"Mallplaza_Tornado_Sensibilidad_Operacional_[\d_]+\.xlsx"),
    "esg": ("esg/modelos", r"ESG_sanciones_y_datos\.xlsx"),
}
DOCUMENTOS = [("informe_final/documento", r"Mallplaza_Informe_CFA_(ES|EN)\.(pdf|docx)")]


def vigente(carpeta, patron):
    rx = re.compile(patron)
    cands = [p for p in (SEC / carpeta).glob("*") if rx.fullmatch(p.name)]
    if not cands:
        return None, []
    clave = lambda p: (int((rx.fullmatch(p.name).groups() or ("0",))[0]) if (rx.fullmatch(p.name).groups() or ("0",))[0].isdigit() else 0,
                       p.stat().st_mtime)
    cands.sort(key=clave, reverse=True)
    return cands[0], cands[1:]


def sha(p, n=10):
    return hashlib.sha256(p.read_bytes()).hexdigest()[:n]


def abrir(p):
    return openpyxl.load_workbook(p, data_only=True)


class Cifras:
    def __init__(self):
        self.filas = []     # (grupo, clave, etiqueta, valor, formato, fuente)

    def add(self, grupo, clave, etiqueta, valor, fmt, fuente):
        self.filas.append({"grupo": grupo, "clave": clave, "etiqueta": etiqueta, "valor": valor, "formato": fmt, "fuente": fuente})


def fila_de(ws, texto, col=1, desde=1):
    for r in range(desde, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip().startswith(texto):
            return r
    raise KeyError(f"{ws.title}: no hay fila que empiece con «{texto}»")


def main():
    arch = {}
    viejos = {}
    for fam, (carpeta, patron) in FAMILIAS.items():
        p, otros = vigente(carpeta, patron)
        if p is None:
            print(f"AVISO: falta la familia {fam} ({carpeta})")
            continue
        arch[fam], viejos[fam] = p, otros
    ref = lambda fam, hoja, celda: f"{arch[fam].name}!{hoja}!{celda}"

    c = Cifras()
    v = abrir(arch["valoracion"])
    S, PO, W, D, E, DB = v["Supuestos"], v["Precio_Objetivo"], v["WACC"], v["DCF"], v["Expansion"], v["Deuda"]
    c.add("Precio y valor", "precio", "Precio de referencia (30-sep-26), CLP", S["B8"].value, "clp", ref("valoracion", "Supuestos", "B8"))
    c.add("Precio y valor", "acciones_mm", "Acciones en circulación (millones)", S["B9"].value, "n0", ref("valoracion", "Supuestos", "B9"))
    c.add("Precio y valor", "dcf", "Valor DCF por acción, CLP", PO["B12"].value, "clp", ref("valoracion", "Precio_Objetivo", "B12"))
    c.add("Precio y valor", "multiplos", "Valor por múltiplos (EV/EBITDA LTM ponderado) por acción, CLP", PO["B13"].value, "clp", ref("valoracion", "Precio_Objetivo", "B13"))
    c.add("Precio y valor", "peso_dcf", "Peso del DCF en el precio objetivo", PO["B14"].value, "pct0", ref("valoracion", "Precio_Objetivo", "B14"))
    c.add("Precio y valor", "precio_objetivo", "PRECIO OBJETIVO (80% DCF / 20% múltiplos), CLP", PO["B15"].value, "clp", ref("valoracion", "Precio_Objetivo", "B15"))
    c.add("Precio y valor", "potencial", "Potencial vs precio", PO["B17"].value, "pct1", ref("valoracion", "Precio_Objetivo", "B17"))
    c.add("Precio y valor", "multiplo_ev_ebitda", "Múltiplo EV/EBITDA LTM ponderado (70% Chile / 30% set común)", PO["B5"].value, "x1", ref("valoracion", "Precio_Objetivo", "B5"))
    c.add("Precio y valor", "vt_ev", "Valor terminal / EV (DCF)", D["C54"].value, "pct1", ref("valoracion", "DCF", "C54"))
    c.add("Precio y valor", "cap_rate_dcf", "Cap rate implícito del DCF (EBITDA LTM / EV)", D["C56"].value, "pct1", ref("valoracion", "DCF", "C56"))
    c.add("Balance y operación", "dfn_mm", "Deuda financiera neta 30-jun-26 (CLP MM)", S["B10"].value, "n0", ref("valoracion", "Supuestos", "B10"))
    c.add("Balance y operación", "ebitda_ltm_mm", "EBITDA LTM 2T26 (CLP MM)", S["B13"].value, "n0", ref("valoracion", "Supuestos", "B13"))
    c.add("Balance y operación", "dfn_ebitda", "DFN / EBITDA LTM", S["B10"].value / S["B13"].value, "x2", "calculado: Supuestos!B10 / B13")
    c.add("Balance y operación", "impuesto_lp", "Tasa de impuesto operacional de largo plazo", S["B59"].value, "pct0", ref("valoracion", "Supuestos", "B59"))
    c.add("Balance y operación", "g", "Crecimiento terminal nominal (g)", S["B88"].value, "pct1", ref("valoracion", "Supuestos", "B88"))
    c.add("Costo de capital", "wacc_inicial", "WACC nominal CLP 2S26E", W["C43"].value, "pct2", ref("valoracion", "WACC", "C43"))
    c.add("Costo de capital", "wacc_terminal", "WACC nominal CLP terminal", W["I43"].value, "pct2", ref("valoracion", "WACC", "I43"))
    c.add("Costo de capital", "wacc_real_terminal", "WACC real terminal", W["I44"].value, "pct2", ref("valoracion", "WACC", "I44"))
    c.add("Costo de capital", "ke_terminal", "Ke nominal CLP terminal", W["I39"].value, "pct2", ref("valoracion", "WACC", "I39"))
    c.add("Costo de capital", "kd_dt_terminal", "Kd nominal después de impuestos terminal", W["I42"].value, "pct2", ref("valoracion", "WACC", "I42"))
    c.add("Costo de capital", "beta_terminal", "Beta re-apalancada terminal", W["I37"].value, "n2", ref("valoracion", "WACC", "I37"))
    c.add("Costo de capital", "crp", "CRP ponderado por ingresos", W["C21"].value, "pct2", ref("valoracion", "WACC", "C21"))
    c.add("Costo de capital", "rf", "Tasa libre de riesgo usada (UST 10 años)", W["C22"].value, "pct2", ref("valoracion", "WACC", "C22"))
    c.add("Costo de capital", "wacc_real_tasacion", "WACC real implícita en la tasación (Nota 11, Chile)", W["C47"].value, "pct2", ref("valoracion", "WACC", "C47"))
    c.add("Expansión", "yield_on_cost", "Yield on cost de la expansión", E["C37"].value, "pct2", ref("valoracion", "Expansion", "C37"))
    c.add("Expansión", "ronic", "RONIC real después de impuestos (yield × (1 − t))", E["C38"].value, "pct2", ref("valoracion", "Expansion", "C38"))

    fa = abrir(arch["fa_exhibits"])
    if "E8_ROIC_RONIC" in fa.sheetnames:
        e8 = fa["E8_ROIC_RONIC"]
        c.add("Retornos", "roic_2025", "ROIC 2025 (NOPAT / capital promedio a valor razonable)", e8["C21"].value, "pct2", ref("fa_exhibits", "E8_ROIC_RONIC", "C21"))
        c.add("Retornos", "roic_menos_wacc_real", "ROIC 2025 − WACC real (pp)", e8["C33"].value, "n2", ref("fa_exhibits", "E8_ROIC_RONIC", "C33"))
        c.add("Retornos", "ronic_menos_wacc_real", "RONIC − WACC real (pp)", e8["C41"].value, "n2", ref("fa_exhibits", "E8_ROIC_RONIC", "C41"))
    d1 = fa["E1_DuPont"]
    c.add("Retornos", "roe_normalizado_2025", "ROE normalizado 2025", d1["H28"].value, "pct1", ref("fa_exhibits", "E1_DuPont", "H28"))
    c.add("Retornos", "roe_reportado_2025", "ROE reportado 2025", d1["H27"].value, "pct1", ref("fa_exhibits", "E1_DuPont", "H27"))

    # Escenarios (réplica del DCF v4 → escenarios_dcf.py / anexos_valuacion.py) y datos del resumen
    da = SEC / "valuacion/graficos/figuras/datos_anexos.json"
    if da.exists():
        esc = json.loads(da.read_text())["F"]["esc"]
        for k, nombre in (("bear", "Bear"), ("base", "Base"), ("bull", "Bull"), ("tas", "Tasadores")):
            c.add("Escenarios", f"esc_{k}", f"Precio objetivo escenario {nombre}, CLP", esc[k]["po"], "clp", "datos_anexos.json (réplica del DCF v4, Anexo 18)")
            c.add("Escenarios", f"esc_{k}_vs", f"{nombre}: variación vs precio", esc[k]["vs"], "pct1", "datos_anexos.json")
    dr = RAIZ / "secciones/informe_final/graficos/datos_resumen.json"
    if dr.exists():
        r = json.loads(dr.read_text())
        c.add("Mercado", "rango_52s_min", "Mínimo de 52 semanas, CLP", r["rango_52s"][0], "clp", "datos_resumen.json (Bloomberg)")
        c.add("Mercado", "rango_52s_max", "Máximo de 52 semanas, CLP", r["rango_52s"][1], "clp", "datos_resumen.json (Bloomberg)")
        c.add("Mercado", "dividendo_2027", "Dividendo esperado 2027E, CLP por acción", r["dividendo"], "n1", "datos_resumen.json (modelo vigente, hoja Deuda)")
        c.add("Mercado", "capitalizacion_mm", "Capitalización bursátil (CLP MM)", r["capitalizacion_mm"], "n0", "datos_resumen.json")

    ci = SEC / "informe_final/graficos/cifras.json"
    if ci.exists():                       # cifras que el informe calcula con la réplica del DCF y los datos del equipo (cifras_informe.py)
        k = json.loads(ci.read_text())
        fuente = "informe_final/graficos/cifras.json (cifras_informe.py; réplica del DCF v4)"
        c.add("Escenarios", "esc_ponderado_po", "Precio objetivo ponderado con las probabilidades de la matriz de riesgo, CLP", k["ponderado"]["po"], "clp", fuente)
        c.add("Pipeline", "pipeline_m2", "Pipeline con fecha 2026–2028 (m²)", k["pipeline_m2"], "n0", fuente)
        c.add("Pipeline", "pipeline_usd_mm", "Pipeline con fecha, inversión (USD MM)", k["pipeline_usd_mm"], "n0", fuente)
        c.add("Mercado", "consenso_po_12m", "Precio objetivo de consenso a 12 meses (Bloomberg, 14 fuentes), CLP", k["consenso"]["po_12m"], "clp", fuente)
        c.add("Sensibilidades", "wacc_menos_100pb", "Δ valor DCF por acción con WACC −100 pb (CLP)", k["sens"]["wacc_m100pb"], "n0", fuente)
        c.add("Sensibilidades", "wacc_mas_100pb", "Δ valor DCF por acción con WACC +100 pb (CLP)", k["sens"]["wacc_100pb"], "n0", fuente)
        c.add("Sensibilidades", "ssr_1pp", "Δ valor DCF por acción con +1 pp de arriendo mismas tiendas 2027–28 (CLP)", k["sens"]["ssr_1pp"], "n1", fuente)
        c.add("Sensibilidades", "sss_1pp", "Δ valor DCF por acción con +1 pp de ventas de locatarios 2027–28 (CLP)", k["sens"]["sss_1pp"], "n1", fuente)

    ind = abrir(arch["industria"])
    ef = ind["EF. OPERATIVA (CLPM2)"]
    for rr in range(6, 9):
        c.add("Industria", f"gla_{ef.cell(rr,1).value}".lower().replace(" ", "_"), f"GLA {ef.cell(rr,1).value} (m²)", ef.cell(rr, 3).value, "n0", ref("industria", "EF. OPERATIVA (CLPM2)", f"C{rr}"))
    ee = abrir(arch["eficiencia"])["Eficiencia_EBITDA_GLA_2T26"]      # EBITDA/m² (Parque Arauco ajustado por multifamily)
    filas_ee = [rr for rr in range(1, 12) if ee.cell(rr, 1).value in ("Mallplaza", "Cenco Malls", "Parque Arauco")]
    for rr in filas_ee:
        c.add("Industria", f"ebitda_m2_{ee.cell(rr,1).value}".lower().replace(" ", "_"), f"EBITDA por m² trimestral 2T26 {ee.cell(rr,1).value} (CLP)", ee.cell(rr, 6).value, "n0", ref("eficiencia", "Eficiencia_EBITDA_GLA_2T26", f"F{rr}"))
    neg = abrir(arch["negocio"])
    ws = neg["Grafico 1, 2 tortas"]
    for row in ws.iter_rows(min_row=1, values_only=True):
        if row[1] and isinstance(row[3], (int, float)) and row[2] in ("Chile", "Perú", "Colombia", "Arriendo", "Parking",
                                                                       "Servicios comerciales y otros", "Publicidad"):
            c.add("Negocio", f"{'gla' if row[1].startswith('A') else 'ingreso'}_{row[2]}".lower().replace(" ", "_"), f"{row[1]} — {row[2]} (%)",
                  row[3], "n1", ref("negocio", "Grafico 1, 2 tortas", "B:D"))

    # ---- salidas ------------------------------------------------------------------------------------------
    for nombre in GENERADOS:                                        # borra lo viejo
        (SALIDA / nombre).unlink(missing_ok=True)
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M")
    fmt = {"clp": lambda x: f"CLP {x:,.0f}".replace(",", "."), "n0": lambda x: f"{x:,.0f}".replace(",", "."),
           "n1": lambda x: f"{x:,.1f}".replace(".", ",").replace("-", "−"), "n2": lambda x: f"{x:,.2f}".replace(".", ",").replace("-", "−"),
           "x1": lambda x: f"{x:,.1f}x".replace(".", ","), "x2": lambda x: f"{x:,.2f}x".replace(".", ","),
           "pct0": lambda x: f"{x*100:.0f}%", "pct1": lambda x: f"{x*100:.1f}%".replace(".", ",").replace("-", "−"),
           "pct2": lambda x: f"{x*100:.2f}%".replace(".", ",")}
    (SALIDA / "cifras_clave.json").write_text(json.dumps(
        {"generado": ahora, "modelos": {k: {"archivo": p.name, "sha256": sha(p)} for k, p in arch.items()}, "cifras": c.filas},
        ensure_ascii=False, indent=1), encoding="utf-8")
    L = ["# Cifras clave del informe (contexto para el LLM)", "",
         f"Generado {ahora} por `contexto_llm/generar.py` a partir de los modelos **vigentes**. No editar a mano: se reemplaza completo en cada ejecución.",
         "Si una cifra de un texto contradice esta tabla, está mal el texto. Cada fila indica su fuente (archivo!hoja!celda).", "",
         "Reglas del informe: ver `CLAUDE.md` y `narrativa_y_coherencia.md`. Recomendación: VENDER.", ""]
    grupo = None
    for f in c.filas:
        if f["grupo"] != grupo:
            grupo = f["grupo"]
            L += ["", f"## {grupo}", "", "| Cifra | Valor | Fuente |", "|---|---|---|"]
        L.append(f"| {f['etiqueta']} | {fmt[f['formato']](f['valor'])} | `{f['fuente']}` |")
    (SALIDA / "cifras_clave.md").write_text("\n".join(L) + "\n", encoding="utf-8")

    A = ["# Archivos vigentes", "", f"Generado {ahora}. Vigente = versión más alta de cada familia; los demás son viejos.", "",
         "## Modelos", "", "| Familia | Vigente | Más viejos (candidatos a eliminar si la regla de dos versiones lo pide) |", "|---|---|---|"]
    for fam, p in arch.items():
        A.append(f"| {fam} | `{p.relative_to(SEC)}` (`{sha(p)}`) | {', '.join('`' + o.name + '`' for o in viejos[fam]) or '—'} |")
    A += ["", "## Documentos finales", ""]
    for carpeta, patron in DOCUMENTOS:
        for p in sorted((SEC / carpeta).glob("*")):
            if re.fullmatch(patron, p.name):
                A.append(f"- `{p.relative_to(SEC)}` — {datetime.fromtimestamp(p.stat().st_mtime):%Y-%m-%d %H:%M}")
    (SALIDA / "archivos_vigentes.md").write_text("\n".join(A) + "\n", encoding="utf-8")
    print(f"{len(c.filas)} cifras de {len(arch)} modelos → {SALIDA.relative_to(RAIZ)}/")
    return c


if __name__ == "__main__":
    main()
